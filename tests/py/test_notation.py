import copy
import io
import json
import subprocess

from conftest import UNC

import indices_depot as idx
import rendre_notation as rn

EXEMPLE = UNC / "notation-etudiants" / "exemples" / "notation-exemple.json"


def exemple():
    return json.loads(EXEMPLE.read_text(encoding="utf-8"))


def test_exemple_valide():
    n, rapport = rn.normaliser_notation(exemple())
    assert rapport.erreurs == []
    assert rapport.avertissements == []
    assert n["_sur"] == 20
    totaux = {e["_qui"]: e["_total"] for e in n["_etudiants"]}
    assert totaux == {"MARTIN Léa": 19, "BERNARD Hugo": 15, "PETIT Chloé": 18, "ROBERT Lucas": None}


def test_erreurs():
    donnees = exemple()
    lea, hugo, _, lucas = donnees["etudiants"]
    lea["notes"]["mail"]["note"] = 2
    lea["notes"]["inconnu"] = {"note": 1}
    hugo["ia"]["probabilite"] = "65 %"
    hugo["ia"]["part"] = "80 à 100 %"
    lucas["notes"] = {"mail": {"note": 1}}
    _, rapport = rn.normaliser_notation(donnees)
    textes = " ".join(rapport.erreurs)
    assert "maximum de 1" in textes
    assert "« inconnu » absent du barème" in textes
    assert "fourchette" in textes
    assert "100 %" in textes
    assert "ROBERT Lucas (non évaluable)" in textes


def test_avertissements():
    donnees = exemple()
    hugo = donnees["etudiants"][1]
    hugo["aTrancher"] = []
    hugo["notes"]["interface"]["justification"] = ""
    hugo["commentaire"] = "Bon travail — à revoir"
    _, rapport = rn.normaliser_notation(donnees)
    textes = " ".join(rapport.avertissements)
    assert "échange oral" in textes
    assert "« interface » sans justification" in textes
    assert "tiret cadratin" in textes


def test_noms_onglets():
    pris = {"synthèse", "barème"}
    assert rn.nom_onglet({"nom": "o'neil", "prenom": "Sam/Lee"}, pris) == "O'NEIL Sam Lee"
    long = {"nom": "X" * 40, "prenom": "Y"}
    assert len(rn.nom_onglet(long, pris)) <= 31
    assert rn.nom_onglet(long, pris).endswith(" 2")


def test_xlsx():
    from openpyxl import load_workbook

    n, _ = rn.normaliser_notation(exemple())
    wb = load_workbook(io.BytesIO(rn.rendre_xlsx(n)))
    assert wb.sheetnames == ["Synthèse", "Barème", "BERNARD Hugo", "MARTIN Léa", "PETIT Chloé", "ROBERT Lucas"]
    ws = wb["MARTIN Léa"]
    assert ws["A1"].value == "Critère" and ws["D1"].value.startswith("Justification")
    assert ws["A14"].value == "Total" and ws["B14"].value == "=SUM(B2:B13)" and ws["C14"].value == 20
    ia = [r for r in ws.iter_rows(min_row=15) if r[0].value and r[0].value.startswith("Probabilité")][0]
    assert ia[1].value == "20 à 50 %" and ia[1].number_format == "@"
    # Synthèse : meilleure note d'abord, formule liée à l'onglet, non évaluable à la fin.
    synthese = wb["Synthèse"]
    assert [synthese.cell(row=k, column=1).value for k in range(5, 9)] == ["MARTIN Léa", "PETIT Chloé", "BERNARD Hugo", "ROBERT Lucas"]
    assert synthese["B5"].value == "='MARTIN Léa'!B14"
    assert synthese["B8"].value is None
    # Projet libre : lignes renommées.
    assert wb["PETIT Chloé"]["A6"].value == "1.1 Récupération des données"


def git(depot, *args):
    subprocess.run(["git", "-C", str(depot), "-c", "user.name=Etudiant Test", "-c", "user.email=etu@exemple.org", *args], check=True, capture_output=True)


def test_indices_depot(tmp_path):
    a = tmp_path / "a"
    a.mkdir()
    git(a, "init", "-q")
    (a / "server.mjs").write_text("// Serveur\nconsole.log('ok');\nif (x) {\n  y();\n}\n", encoding="utf-8")
    (a / ".env").write_text("SECRET=1\n", encoding="utf-8")
    git(a, "add", ".")
    git(a, "commit", "-q", "-m", "feat: serveur")
    (a / "app.js").write_text("/** Client */\nfunction f() {\n    return 1; // ✅ fait\n}\n", encoding="utf-8")
    (a / "package-lock.json").write_text("{}\n" * 500, encoding="utf-8")
    git(a, "add", ".")
    git(a, "commit", "-q", "-m", "Add client")
    git(a, "tag", "v1")
    git(a, "tag", "v2")

    x = idx.analyser_depot(a)
    assert x["nb_commits"] == 2 and x["tags"] == ["v1", "v2"] and x["tags_meme_commit"]
    assert x["lignes_ajoutees"] == 10  # package-lock.json exclu
    assert x["messages_anglais"] == ["Add client"]
    assert x["a_ne_pas_versionner"] == [".env"]
    assert x["console"] == [{"fichier": "server.mjs", "nombre": 1}]
    assert [e["fichier"] for e in x["emojis"]] == ["app.js"]
    assert x["jsdoc"] == 1
    assert x["indentation"] == {"server.mjs": "2 espaces", "app.js": "4 espaces"}
    assert x["melange_indentation"] is False  # extensions différentes

    b = tmp_path / "b"
    git(tmp_path, "clone", "-q", str(a), str(b))
    paire = idx.comparer_depots([a, b])[0]
    assert paire["similarite"] == 1 and paire["fichiers_identiques"] >= 2
    assert paire["auteurs_communs"] == ["etu@exemple.org"]
    assert "Tags : v1, v2 (tous sur le même commit)" in idx.resume(copy.deepcopy(x))
