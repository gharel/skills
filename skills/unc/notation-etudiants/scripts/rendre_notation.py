#!/usr/bin/env python3
"""Notation de rendus étudiants : notation.json → classeur .xlsx (synthèse, barème, un onglet par étudiant).

Usage :
    python scripts/rendre_notation.py notation.json [--sortie DOSSIER] [--verifier]

Le JSON suit references/format-notation.md. Le classeur s'ouvre dans Excel, LibreOffice ou Google Sheets
(Fichier > Importer). Les totaux sont des formules : l'enseignant peut corriger une note, tout se recalcule.
Les estimations de l'usage de l'IA restent du texte, hors du total.
"""

from __future__ import annotations

import argparse
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import charte as ch  # noqa: E402

FORMAT = "notation-etudiants"
STATUTS = {"evalue": "Évalué", "non-evaluable": "Non évaluable", "sans-rendu": "Sans rendu"}
FOURCHETTE = re.compile(r"^(\d{1,3})\s*à\s*(\d{1,3})\s*%$")
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿⭐⬆⬇]")
INTERDITS_ONGLET = re.compile(r"[\[\]:*?/\\]")
LIBELLE_PROBABILITE = "Probabilité d'usage de l'IA (indicatif, hors note)"
LIBELLE_PART = "Part estimée du travail produite avec l'IA"
RAPPEL_IA = (
    "Estimations de l'usage de l'IA : indicatives, fondées sur des indices, sans valeur de preuve. "
    "Elles n'entrent pas dans la note. En cas de doute, seul un échange avec l'étudiant permet de trancher."
)


# ---------------------------------------------------------------- contrôle


def nombre(valeur):
    """Nombre JSON ou texte « 0,5 » ; None si la valeur n'est pas un nombre."""
    if isinstance(valeur, bool):
        return None
    if isinstance(valeur, (int, float)):
        return float(valeur)
    if isinstance(valeur, str):
        try:
            return float(valeur.strip().replace(",", ".").replace(ch.NBSP, ""))
        except ValueError:
            return None
    return None


def affichage(x: float) -> str:
    return ("{:g}".format(x)).replace(".", ",")


def fourchette(texte, champ: str, qui: str, declare: bool, r: ch.Rapport):
    """Contrôle une estimation « 50 à 80 % » ; renvoie (bas, haut) ou None."""
    t = str(texte or "").replace(ch.NBSP, " ").replace(ch.NNBSP, " ").strip()
    m = FOURCHETTE.match(t)
    if not m:
        r.erreur("{} : « {} » doit être une fourchette (par exemple « 50 à 80 % »), jamais une valeur précise.".format(qui, champ))
        return None
    bas, haut = int(m.group(1)), int(m.group(2))
    if bas >= haut or haut > 100:
        r.erreur("{} : fourchette « {} » incohérente ({}).".format(qui, champ, t))
        return None
    if haut == 100 and not declare:
        r.erreur("{} : « {} » ne monte à 100 % que pour un usage de l'IA déclaré sur tout le rendu.".format(qui, champ))
    return bas, haut


def textes(valeur, chemin=""):
    """Tous les textes d'une structure JSON, avec leur chemin."""
    if isinstance(valeur, str):
        yield chemin, valeur
    elif isinstance(valeur, dict):
        for k, v in valeur.items():
            yield from textes(v, "{}.{}".format(chemin, k) if chemin else k)
    elif isinstance(valeur, list):
        for k, v in enumerate(valeur):
            yield from textes(v, "{}[{}]".format(chemin, k))


def nom_onglet(e: dict, pris: set[str]) -> str:
    base = INTERDITS_ONGLET.sub(" ", "{} {}".format(e["nom"].upper(), e.get("prenom", "")).strip())
    base = re.sub(r"\s+", " ", base)[:31].strip().strip("'")
    nom, k = base, 2
    while nom.lower() in pris:
        suffixe = " {}".format(k)
        nom, k = base[: 31 - len(suffixe)] + suffixe, k + 1
    pris.add(nom.lower())
    return nom


def normaliser_notation(brut: dict) -> tuple[dict, ch.Rapport]:
    """Contrôle le JSON de notation et calcule les totaux. Les notes ne sont jamais corrigées en silence."""
    r = ch.Rapport()
    n = dict(brut)
    if n.get("format", FORMAT) != FORMAT:
        r.erreur("« format » doit valoir « {} ».".format(FORMAT))
    controle = dict(n.get("controle") or {})
    if not str(controle.get("intitule") or "").strip():
        r.erreur("« controle.intitule » (nom du TP ou du contrôle) est obligatoire.")
    bareme = n.get("bareme")
    if not isinstance(bareme, list) or not bareme:
        r.erreur("« bareme » est vide ou absent.")
    etudiants = n.get("etudiants")
    if not isinstance(etudiants, list) or not etudiants:
        r.erreur("« etudiants » est vide ou absent.")
    if not r.ok:
        return n, r

    for chemin, t in textes(brut):
        if "—" in t:
            r.avertir("{} : tiret cadratin à remplacer (virgule, deux-points ou parenthèses).".format(chemin))
        if EMOJI.search(t):
            r.avertir("{} : émoji à retirer.".format(chemin))

    # Barème : règles fixées avant de noter, identiques pour tous.
    criteres, ids = [], set()
    for i, c in enumerate(bareme, 1):
        c = dict(c) if isinstance(c, dict) else {}
        cid = str(c.get("id") or "").strip()
        if not cid:
            r.erreur("barème, ligne {} : « id » manquant.".format(i))
            continue
        if cid in ids:
            r.erreur("barème : identifiant « {} » en double.".format(cid))
            continue
        ids.add(cid)
        if not str(c.get("libelle") or "").strip():
            r.erreur("barème « {} » : « libelle » manquant.".format(cid))
        maxi, penalite = nombre(c.get("max")), nombre(c.get("penalite"))
        if (maxi is None) == (penalite is None):
            r.erreur("barème « {} » : indiquez soit « max » (points), soit « penalite » (nombre négatif).".format(cid))
            continue
        if maxi is not None and maxi <= 0:
            r.erreur("barème « {} » : « max » doit être positif.".format(cid))
        if penalite is not None and penalite >= 0:
            r.erreur("barème « {} » : « penalite » doit être négative (par exemple -10).".format(cid))
        if not str(c.get("regle") or "").strip():
            r.avertir("barème « {} » : « regle » absente. Écrivez la règle avant de noter, pour l'appliquer à tous.".format(cid))
        c.update(id=cid, max=maxi, penalite=penalite)
        criteres.append(c)
    if not r.ok:
        return n, r
    sur = sum(c["max"] for c in criteres if c["max"] is not None)

    vus, eleves = set(), []
    for i, e in enumerate(etudiants, 1):
        e = dict(e) if isinstance(e, dict) else {}
        if not str(e.get("nom") or "").strip():
            r.erreur("étudiant {} : « nom » manquant.".format(i))
            continue
        qui = "{} {}".format(e["nom"].upper(), e.get("prenom", "")).strip()
        if qui.lower() in vus:
            r.erreur("{} apparaît deux fois.".format(qui))
        vus.add(qui.lower())
        statut = e.get("statut") or "evalue"
        if statut not in STATUTS:
            r.erreur("{} : statut « {} » inconnu ({}).".format(qui, statut, ", ".join(STATUTS)))
            continue
        notes_brutes = e.get("notes") or {}
        for cle in notes_brutes:
            if cle not in ids:
                r.erreur("{} : critère « {} » absent du barème.".format(qui, cle))
        notes, total = {}, 0.0
        if statut != "evalue":
            donnees = [v.get("note") if isinstance(v, dict) else v for v in notes_brutes.values()]
            if any(x not in (None, "") for x in donnees):
                r.erreur("{} ({}) : ne donnez aucune note, expliquez la situation dans « commentaire ».".format(qui, STATUTS[statut].lower()))
            if not str(e.get("commentaire") or "").strip():
                r.avertir("{} ({}) : « commentaire » vide, dites pourquoi.".format(qui, STATUTS[statut].lower()))
        else:
            for c in criteres:
                v = notes_brutes.get(c["id"])
                d = v if isinstance(v, dict) else {"note": v}
                valeur = nombre(d.get("note"))
                if d.get("note") in (None, ""):
                    if c["penalite"] is not None:
                        valeur = 0.0
                    else:
                        r.erreur("{} : note manquante pour « {} ».".format(qui, c["id"]))
                        continue
                elif valeur is None:
                    r.erreur("{} : note « {} » illisible pour « {} ».".format(qui, d.get("note"), c["id"]))
                    continue
                if c["max"] is not None and not 0 <= valeur <= c["max"]:
                    r.erreur("{} : « {} » vaut {} pour un maximum de {}.".format(qui, c["id"], affichage(valeur), affichage(c["max"])))
                if c["penalite"] is not None and not c["penalite"] <= valeur <= 0:
                    r.erreur("{} : pénalité « {} » de {} hors de [{} ; 0].".format(qui, c["id"], affichage(valeur), affichage(c["penalite"])))
                justification = str(d.get("justification") or "").strip()
                if not justification and (c["max"] is not None or valeur):
                    r.avertir("{} : « {} » sans justification. Chaque note s'appuie sur un fait observé.".format(qui, c["id"]))
                notes[c["id"]] = {"note": valeur, "justification": justification}
                total += valeur
            if total < 0:
                r.avertir("{} : total négatif ({}), pénalités à trancher par l'enseignant.".format(qui, affichage(total)))

        ia = dict(e.get("ia") or {})
        declare = bool(str(ia.get("declare") or "").strip())
        if statut == "evalue":
            if not ia:
                r.avertir("{} : estimation de l'usage de l'IA absente (« ia »).".format(qui))
            else:
                proba = fourchette(ia.get("probabilite"), "ia.probabilite", qui, declare, r)
                fourchette(ia.get("part"), "ia.part", qui, declare, r)
                if not str(ia.get("indices") or "").strip():
                    r.avertir("{} : citez les indices (et indices contraires) qui fondent l'estimation de l'IA.".format(qui))
                a_trancher = " ".join(ch.normaliser(x) for x in e.get("aTrancher") or [])
                if proba and proba[0] >= 50 and "oral" not in a_trancher and "echange" not in a_trancher:
                    r.avertir("{} : probabilité d'usage de l'IA élevée, proposez un échange oral dans « aTrancher ».".format(qui))
        elif ia:
            r.avertir("{} ({}) : laissez le bloc « ia » vide.".format(qui, STATUTS[statut].lower()))

        eleves.append(dict(e, nom=e["nom"].strip(), prenom=str(e.get("prenom") or "").strip(), statut=statut, _qui=qui,
                           _notes=notes, _total=total if statut == "evalue" else None, ia=ia, _rendu=dict(e.get("rendu") or {}),
                           libelles=dict(e.get("libelles") or {}), aTrancher=list(e.get("aTrancher") or [])))
        for cle in eleves[-1]["libelles"]:
            if cle not in ids:
                r.erreur("{} : « libelles.{} » ne correspond à aucun critère du barème.".format(qui, cle))

    pris = {"synthèse", "barème"}
    for e in sorted(eleves, key=lambda x: x["_qui"].lower()):
        e["_onglet"] = nom_onglet(e, pris)
    n.update(controle=controle, _criteres=criteres, _sur=sur, _etudiants=eleves, promotion=dict(n.get("promotion") or {}))
    return n, r


# ---------------------------------------------------------------- classeur


def ref_onglet(nom: str) -> str:
    return "'{}'".format(nom.replace("'", "''"))


def rendre_xlsx(n: dict) -> bytes:
    """Classeur : Synthèse, Barème, puis un onglet par étudiant (ordre alphabétique)."""
    ch.exiger("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    t = ch.typographie
    gras = Font(bold=True)
    entete = PatternFill("solid", fgColor="E8E8EC")
    gris = PatternFill("solid", fgColor="F2F2F2")
    haut = Alignment(wrap_text=True, vertical="top")
    controle, criteres, sur = n["controle"], n["_criteres"], n["_sur"]

    wb = Workbook()
    synthese = wb.active
    synthese.title = "Synthèse"
    ws_bareme = wb.create_sheet("Barème")

    # Barème : la règle de chaque ligne, appliquée à tous.
    ws_bareme.append(["Critère", "Max", "Pénalité", "Règle appliquée à tous"])
    for c in criteres:
        ws_bareme.append([t(c["libelle"]), c["max"], c["penalite"], t(c.get("regle", ""))])
    ws_bareme.append(["Total", sur, None, None])
    for k, largeur in zip("ABCD", (44, 8, 10, 100)):
        ws_bareme.column_dimensions[k].width = largeur
    for ligne in ws_bareme.iter_rows():
        for cellule in ligne:
            cellule.alignment = haut
    for cellule in ws_bareme[1] + ws_bareme[ws_bareme.max_row]:
        cellule.font = gras
    for cellule in ws_bareme[1]:
        cellule.fill = entete

    # Un onglet par étudiant : A critère, B note, C max, D justification.
    lignes_total = {}
    for e in sorted(n["_etudiants"], key=lambda x: x["_onglet"].lower()):
        ws = wb.create_sheet(e["_onglet"])
        ws.append(["Critère", "Note", "Max", "Justification (fait observé, test, date)"])
        for c in criteres:
            note = e["_notes"].get(c["id"], {})
            ws.append([
                t(e["libelles"].get(c["id"]) or c["libelle"]),
                note.get("note") if e["statut"] == "evalue" else None,
                c["max"] if c["max"] is not None else c["penalite"],
                t(note.get("justification", "")),
            ])
        ligne_total = ws.max_row + 1
        lignes_total[e["_onglet"]] = ligne_total
        if e["statut"] == "evalue":
            ws.append(["Total", "=SUM(B2:B{})".format(ligne_total - 1), sur, None])
        else:
            ws.append(["Total", STATUTS[e["statut"]].lower(), sur, None])
        for cellule in ws[1] + ws[ligne_total]:
            cellule.font = gras
        for cellule in ws[1]:
            cellule.fill = entete

        def bloc(libelle, texte, ligne_vide=True):
            if ligne_vide:
                ws.append([])
            ws.append([libelle, t(texte)])
            k = ws.max_row
            ws.merge_cells("B{0}:D{0}".format(k))
            ws["A{}".format(k)].font = gras
            return k

        projet_libre = " (projet libre : lignes renommées)" if e["libelles"] else ""
        bloc("Commentaire" + projet_libre, e.get("commentaire", ""))
        rendu = e["_rendu"]
        morceaux = [x for x in (
            rendu.get("date"),
            "GitHub : " + rendu["github"] if rendu.get("github") else "",
            "Déploiement : " + rendu["deploiement"] if rendu.get("deploiement") else "",
            "hors délai" if rendu.get("horsDelai") else "",
        ) if x]
        bloc("Rendu", " · ".join(morceaux) or "aucun lien fourni", ligne_vide=False)
        if e["aTrancher"]:
            bloc("À trancher par l'enseignant", "\n".join("- " + x for x in e["aTrancher"]), ligne_vide=False)

        # Bloc IA : texte sur fond gris, hors du total.
        ws.append([])
        ia = e["ia"]
        parties = " ".join(x for x in (ia.get("parties", ""), "Usage déclaré : " + ia["declare"] if ia.get("declare") else "") if x)
        for libelle, valeur, detail in ((LIBELLE_PROBABILITE, ia.get("probabilite", ""), ia.get("indices", "")), (LIBELLE_PART, ia.get("part", ""), parties)):
            ws.append([libelle, t(valeur), None, t(detail)])
            k = ws.max_row
            ws["B{}".format(k)].number_format = "@"
            for cellule in ws[k]:
                cellule.fill = gris
            ws["A{}".format(k)].font = gras

        for k, largeur in zip("ABCD", (44, 14, 8, 90)):
            ws.column_dimensions[k].width = largeur
        for ligne in ws.iter_rows():
            for cellule in ligne:
                cellule.alignment = haut
        ws.freeze_panes = "A2"

    # Synthèse : de la meilleure note à la plus basse, puis les cas particuliers.
    synthese.append([t(controle["intitule"]) + (" · " + t(controle["matiere"]) if controle.get("matiere") else "")])
    synthese["A1"].font = Font(bold=True, size=14)
    infos = [x for x in ("Énoncé : " + controle["enonce"] if controle.get("enonce") else "", "Date : " + controle["date"] if controle.get("date") else "") if x]
    synthese.append([" · ".join(infos)])
    synthese.append([])
    synthese.append(["Étudiant", "Note", "Sur", "Statut", "IA (probabilité)", "IA (part estimée)", "Rendu", "Points à trancher"])
    for cellule in synthese[4]:
        cellule.font = gras
        cellule.fill = entete
    ordre = {"evalue": 0, "non-evaluable": 1, "sans-rendu": 2}
    for e in sorted(n["_etudiants"], key=lambda x: (ordre[x["statut"]], -(x["_total"] or 0), x["_qui"].lower())):
        note = "={}!B{}".format(ref_onglet(e["_onglet"]), lignes_total[e["_onglet"]]) if e["statut"] == "evalue" else None
        rendu = e["_rendu"]
        synthese.append([
            e["_qui"], note, sur, STATUTS[e["statut"]], t(e["ia"].get("probabilite", "")), t(e["ia"].get("part", "")),
            t(" · ".join(x for x in (rendu.get("date"), "hors délai" if rendu.get("horsDelai") else "") if x)),
            t("\n".join(e["aTrancher"])),
        ])
        for col in "EF":
            synthese["{}{}".format(col, synthese.max_row)].number_format = "@"
    synthese.append([])
    synthese.append([RAPPEL_IA])
    synthese["A{}".format(synthese.max_row)].font = Font(italic=True)
    promotion = n["promotion"]
    for titre, cle in (("Défauts récurrents de la promotion", "defauts"), ("Interprétations du barème retenues", "interpretations")):
        if promotion.get(cle):
            synthese.append([])
            synthese.append([titre])
            synthese["A{}".format(synthese.max_row)].font = gras
            for x in promotion[cle]:
                synthese.append(["- " + t(x)])
    for k, largeur in zip("ABCDEFGH", (28, 8, 6, 15, 16, 16, 20, 70)):
        synthese.column_dimensions[k].width = largeur
    for ligne in synthese.iter_rows(min_row=4):
        for cellule in ligne:
            cellule.alignment = haut
    synthese.freeze_panes = "A5"

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------- programme principal


def nom_fichier(n: dict) -> str:
    return "notation-" + ch.slug(n["controle"]["intitule"]) + ".xlsx"


def rendre(chemin_json: str, sortie: str | None = None, formats=("xlsx",)) -> list[Path]:
    n, rapport = normaliser_notation(ch.lire_json(chemin_json))
    rapport.afficher()
    if not rapport.ok:
        raise SystemExit(1)
    dossier = Path(sortie) if sortie else Path(chemin_json).resolve().parent
    dossier.mkdir(parents=True, exist_ok=True)
    fichiers = []
    if "xlsx" in formats:
        f = dossier / nom_fichier(n)
        f.write_bytes(rendre_xlsx(n))
        fichiers.append(f)
    return fichiers


def main() -> None:
    parser = argparse.ArgumentParser(description="Notation de rendus étudiants : JSON → classeur .xlsx (un onglet par étudiant).")
    parser.add_argument("json", help="fichier notation.json (voir references/format-notation.md)")
    parser.add_argument("--sortie", help="dossier de sortie (par défaut : celui du JSON)")
    parser.add_argument("--verifier", action="store_true", help="contrôle le JSON sans rien produire")
    args = parser.parse_args()
    if args.verifier:
        n, rapport = normaliser_notation(ch.lire_json(args.json))
        rapport.afficher()
        if rapport.ok:
            print("Notation valide : {} étudiants, barème sur {}.".format(len(n["_etudiants"]), affichage(n["_sur"])))
        raise SystemExit(0 if rapport.ok else 1)
    for f in rendre(args.json, args.sortie):
        print(f)


if __name__ == "__main__":
    main()
