"""Plan de formation : rendu DOCX, identique dans Word et dans Google Docs.

Mise en page du modèle « Plan de formation » de Skazy Formation, transposée en DOCX :

- les aplats aux coins arrondis (bandeau et ses pastilles, sticker, encadré, cartes, encart, boîte de
  contact) sont des images PNG ancrées derrière le texte, que l'import de Google Docs conserve ;
- le texte est dans des paragraphes et des tableaux sans bordure, sans marge haute ni basse de cellule :
  les espacements sont des espaces avant et après de paragraphe, et les interlignes des multiples,
  que Word et Google Docs lisent de la même façon (voir references/google-docs.md) ;
- un paragraphe d'espacement de hauteur connue porte les fonds d'un tableau : chaque fond est décalé
  de la hauteur du paragraphe, il commence donc là où commence le tableau.

Les mesures sont en points (1 pt = 1/72 de pouce ; le A4 fait 595,3 × 841,9 pt).
"""

from __future__ import annotations

import io

import charte as ch
from charte import typographie

POLICE = "Raleway"
POLICE_EXTRA = "Raleway ExtraBold"
HAUTEUR_LIGNE = 1.174  # hauteur de ligne naturelle de Raleway, rapportée à la taille de police
PAGE_L, PAGE_H = 595.3, 841.9
MARGE = 48.0  # marges gauche et droite (17 mm)
UTILE = PAGE_L - 2 * MARGE
DPI = 90  # résolution des fonds : nets à l'écran, légers (le base64 du DOCX reste recopiable)
EMU = 12700  # EMU par point

# Page 1
HAUT_1 = 36.55  # marge haute : ligne du titre
BANDEAU_1 = 39.7 * 72 / 25.4  # hauteur de l'aplat du bandeau (150 px du modèle)
PASTILLES_Y = 36.0 * 72 / 25.4  # haut de la rangée des pastilles (21 pt de haut, à cheval sur le bandeau)
PASTILLE_H = 21.0
PIED_1 = 55.6  # hauteur de la rangée du pied de la page 1
DISTANCE_PIED_1 = 6.25
# Page 2
HAUT_2 = 19.55
BANDEAU_2 = 19.6 * 72 / 25.4
PIED_2_CONTACT = 84.0
PIED_2_FILET = 31.2
DISTANCE_PIED_2 = 7.35
BAS_2 = 172.5  # marge basse de la page 2 : place du pied (boîte de contact et filet)

ICONES_FORMATS = ("users", "video", "blend")

# Chasse des caractères de Raleway Bold, en millièmes de cadratin : largeur des pastilles du DOCX,
# sans dépendre de la police installée.
CHASSE_GRAS = dict(zip(
    ' !"#$%&\'()*+,-./0123456789:;<=>?@ABCDEFGHIJKLMNOPQRSTUVWXYZ[\\]^_`abcdefghijklmnopqrstuvwxyz{|}~\xa0\xa1\xa2\xa3\xa4\xa5\xa6\xa7\xa8\xa9\xaa\xab\xac\xad\xae\xaf\xb0\xb1\xb2\xb3\xb4\xb5\xb6\xb7\xb8\xb9\xba\xbb\xbc\xbd\xbe\xbf\xc0\xc1\xc2\xc3\xc4\xc5\xc6\xc7\xc8\xc9\xca\xcb\xcc\xcd\xce\xcf\xd0\xd1\xd2\xd3\xd4\xd5\xd6\xd7\xd8\xd9\xda\xdb\xdc\xdd\xde\xdf\xe0\xe1\xe2\xe3\xe4\xe5\xe6\xe7\xe8\xe9\xea\xeb\xec\xed\xee\xef\xf0\xf1\xf2\xf3\xf4\xf5\xf6\xf7\xf8\xf9\xfa\xfb\xfc\xfd\xfe\xff’‘“”–—…•œŒ€  ',
    (
        240, 307, 393, 721, 634, 786, 712, 235, 310, 309, 345, 421, 232, 418, 225, 698, 614, 494, 564, 558, 571, 554, 607, 561,
        604, 589, 247, 259, 515, 436, 515, 486, 845, 672, 680, 686, 716, 601, 574, 725, 745, 286, 491, 669, 584, 865, 763, 756,
        623, 755, 665, 613, 619, 750, 676, 1053, 646, 660, 627, 316, 637, 316, 585, 475, 255, 574, 634, 564, 645, 599, 370, 634,
        604, 255, 261, 558, 317, 931, 604, 606, 634, 634, 387, 494, 384, 618, 550, 839, 531, 550, 493, 317, 272, 317, 567, 240,
        305, 592, 545, 552, 712, 275, 532, 457, 841, 444, 649, 612, 603, 841, 419, 285, 472, 416, 407, 255, 664, 657, 225, 326,
        298, 443, 649, 883, 913, 991, 479, 672, 672, 672, 672, 672, 672, 966, 686, 601, 601, 601, 601, 286, 286, 286, 286, 721,
        763, 756, 756, 756, 756, 756, 477, 756, 750, 750, 750, 750, 660, 604, 558, 574, 574, 574, 574, 574, 574, 937, 564, 599,
        599, 599, 599, 255, 255, 255, 255, 613, 604, 606, 606, 606, 606, 606, 511, 606, 618, 618, 618, 618, 550, 601, 550, 239,
        246, 434, 425, 603, 928, 602, 363, 1011, 1085, 798, 608, 602,
    ),
))


def chasse_gras(texte: str, taille: float) -> float:
    """Largeur (pt) d'un texte en Raleway Bold."""
    return sum(CHASSE_GRAS.get(c, 600) for c in texte) * taille / 1000


def nombre_de_lignes(texte: str, taille: float, largeur: float) -> int:
    """Lignes d'un texte en Raleway Bold coupé aux espaces dans une largeur donnée (pt)."""
    lignes, courante = 1, 0.0
    espace = chasse_gras(" ", taille)
    for mot in texte.split(" "):
        l = chasse_gras(mot, taille)
        if courante and courante + espace + l > largeur:
            lignes, courante = lignes + 1, l
        else:
            courante += (espace if courante else 0) + l
    return lignes


def ligne(taille: float, multiple: float = 1.0) -> float:
    """Hauteur (pt) d'une ligne de Raleway en interligne multiple."""
    return taille * HAUTEUR_LIGNE * multiple


# ---------------------------------------------------------------- fonds dessinés


def _toile(largeur_pt: float, hauteur_pt: float, fond="#FFFFFF"):
    from PIL import Image

    echelle = DPI / 72 * 4  # dessin en quadruple résolution, réduit à la fin (bords lissés)
    return Image.new("RGBA", (round(largeur_pt * echelle), round(hauteur_pt * echelle)), ch.rgb(fond) + (255,)), echelle


def _fin(im, transparente=False) -> bytes:
    from PIL import Image

    im = im.resize((round(im.width / 4), round(im.height / 4)), Image.LANCZOS)
    if transparente:
        return ch.png(im)
    # Aplats : 16 couleurs suffisent (bords lissés compris), l'image pèse 5 à 10 fois moins.
    tampon = io.BytesIO()
    im.convert("RGB").quantize(16, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(tampon, "PNG", optimize=True)
    return tampon.getvalue()


def fond_arrondi(largeur_pt, hauteur_pt, couleur, rayon_pt=9.0, bordure=None, barre=None, icone=None, encre=None, teinte=None) -> bytes:
    """Aplat arrondi posé sur du blanc : fond de l'encadré, des cartes, de l'encart et de la boîte de contact.

    bordure : couleur d'un liseré de 0,75 pt ; barre : couleur d'une barre de 3 pt en haut (cartes des
    modalités) ; icone : icône Lucide, dans une case teintée (cartes) ou seule à gauche (encart).
    """
    from PIL import Image, ImageDraw

    im, e = _toile(largeur_pt, hauteur_pt)
    d = ImageDraw.Draw(im)
    boite = (0, 0, im.width - 1, im.height - 1)
    rayon = round(rayon_pt * e)
    if barre:
        d.rounded_rectangle(boite, rayon, fill=ch.rgb(barre))
        masque = Image.new("L", im.size, 0)
        ImageDraw.Draw(masque).rounded_rectangle(boite, rayon, fill=255)
        corps = Image.new("RGBA", im.size, (255, 255, 255, 0))
        dc = ImageDraw.Draw(corps)
        dc.rectangle((0, round(3.75 * e), im.width, im.height), fill=ch.rgb(couleur) + (255,))
        im.paste(corps, (0, 0), Image.composite(corps.getchannel("A"), Image.new("L", im.size, 0), masque))
        if bordure:
            d.rounded_rectangle(boite, rayon, outline=ch.rgb(bordure), width=max(1, round(0.75 * e)))
    else:
        d.rounded_rectangle(boite, rayon, fill=ch.rgb(couleur), outline=ch.rgb(bordure) if bordure else None,
                            width=max(1, round(0.75 * e)) if bordure else 0)
    if icone and teinte:  # case d'icône des cartes : 34 px (25,5 pt) à 16 px du bord, icône de 19 px
        x0, y0, cote = round(12.75 * e), round(14.25 * e), round(25.5 * e)
        d.rounded_rectangle((x0, y0, x0 + cote, y0 + cote), round(7.5 * e), fill=ch.rgb(teinte))
        ic = round(14.25 * e)
        im.alpha_composite(ch.icone_image(icone, encre, ic), (x0 + (cote - ic) // 2, y0 + (cote - ic) // 2))
    elif icone:  # icône seule (encart), centrée verticalement : 20 px à 18 px du bord gauche
        ic = round(15 * e)
        im.alpha_composite(ch.icone_image(icone, encre, ic), (round(13.5 * e), (im.height - ic) // 2))
    return _fin(im)


def bandeau_1(couleur, pastilles: list[float]) -> bytes:
    """Bandeau de la page 1 (pleine largeur) et ses pastilles à cheval sur le bord, cerclées de blanc."""
    from PIL import ImageDraw

    hauteur = PASTILLES_Y + PASTILLE_H + 3
    im, e = _toile(PAGE_L, hauteur)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, im.width, round(BANDEAU_1 * e)), fill=ch.rgb(couleur))
    x = MARGE
    for largeur in pastilles:
        cercle = 2.25
        d.rounded_rectangle((round((x - cercle) * e), round((PASTILLES_Y - cercle) * e), round((x + largeur + cercle) * e),
                             round((PASTILLES_Y + PASTILLE_H + cercle) * e)), round((PASTILLE_H / 2 + cercle) * e), fill=(255, 255, 255))
        d.rounded_rectangle((round(x * e), round(PASTILLES_Y * e), round((x + largeur) * e), round((PASTILLES_Y + PASTILLE_H) * e)),
                            round(PASTILLE_H / 2 * e), fill=ch.rgb(ch.COULEURS["fondGris"]))
        x += largeur + 6
    return _fin(im)


def bandeau_2(couleur, encre, pastille_l: float, pastille_y: float, pastille_h: float, hauteur: float = BANDEAU_2) -> bytes:
    """Bandeau de la page 2 et la pastille « Modalités et infos pratiques », alignée sur la marge droite."""
    from PIL import ImageDraw

    im, e = _toile(PAGE_L, hauteur, couleur)
    d = ImageDraw.Draw(im)
    x1 = PAGE_L - MARGE
    d.rounded_rectangle((round((x1 - pastille_l) * e), round(pastille_y * e), round(x1 * e), round((pastille_y + pastille_h) * e)),
                        round(pastille_h / 2 * e), fill=ch.rgb(encre))
    return _fin(im)


def carre(couleur) -> bytes:
    """Puce carrée des intertitres de la page 2 (7,5 pt)."""
    from PIL import ImageDraw

    im, e = _toile(7.5, 7.5)
    ImageDraw.Draw(im).rounded_rectangle((0, 0, im.width - 1, im.height - 1), round(1.5 * e), fill=ch.rgb(couleur))
    return _fin(im)


# ---------------------------------------------------------------- rendu


def rendre_docx(p: dict) -> bytes:
    ch.exiger("docx", "PIL")
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Pt, RGBColor

    c = p["_cat"]
    g = ch.COULEURS
    NOIR, TEXTE, PIED = g["texte"], g["texte2"], g["texte4"]

    doc = Document()
    doc.core_properties.title = p["intitule"]
    doc.core_properties.author = "Skazy Formation"
    doc.core_properties.subject = "Plan de formation"

    # Style Normal : Raleway 10 pt, sans espacement. Police et langue par défaut : Google Docs les reprend
    # pour le texte ajouté après import.
    normal = doc.styles["Normal"]
    normal.font.name = POLICE
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(*ch.rgb(TEXTE))
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    defaut = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    fonts = defaut.find(qn("w:rFonts"))
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), POLICE)
    langue = defaut.find(qn("w:lang"))
    if langue is not None:
        langue.set(qn("w:val"), "fr-FR")
    pdefaut = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:pPrDefault"))
    if pdefaut is not None:
        espacement = pdefaut.find("{}/{}".format(qn("w:pPr"), qn("w:spacing")))
        if espacement is not None:
            espacement.getparent().remove(espacement)

    s1 = doc.sections[0]
    mise_en_page(s1, HAUT_1, 72, DISTANCE_PIED_1)

    # ------------------------------------------------ utilitaires

    def police(run, taille, gras=False, italique=False, couleur=TEXTE, famille=POLICE):
        run.font.size = Pt(taille)
        run.font.bold = gras or None
        run.font.italic = italique or None
        run.font.color.rgb = RGBColor(*ch.rgb(couleur))
        if famille != POLICE:
            rpr = run._element.get_or_add_rPr()
            f = OxmlElement("w:rFonts")
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                f.set(qn(attr), famille)
            rpr.insert(0, f)
        return run

    def texte(par, contenu, taille, couleur=TEXTE, gras=False, italique=False, couleur_gras=NOIR):
        for segment, en_gras in ch.segments_gras(typographie(contenu)):
            police(par.add_run(segment), taille, gras=gras or en_gras, italique=italique,
                   couleur=couleur_gras if en_gras and not gras else couleur)
        return par

    def marque(par, taille):
        """Taille de la marque de paragraphe : hauteur d'un paragraphe vide, dans Word comme dans Google Docs."""
        ppr = par._p.get_or_add_pPr()
        rpr = ppr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ppr.append(rpr)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(round(taille * 2)))
        rpr.append(sz)

    def para(conteneur, avant=0.0, apres=0.0, multiple=1.0, gauche=0.0, droite=0.0, alignement=None, taille=None, premier=False):
        par = conteneur.paragraphs[0] if premier and conteneur.paragraphs else conteneur.add_paragraph()
        pf = par.paragraph_format
        pf.space_before, pf.space_after = Pt(avant), Pt(apres)
        pf.line_spacing = multiple
        if gauche:
            pf.left_indent = Pt(gauche)
        if droite:
            pf.right_indent = Pt(droite)
        if alignement is not None:
            par.alignment = alignement
        if taille:
            marque(par, taille)
        return par

    def espace(conteneur, hauteur, premier=False):
        """Paragraphe vide de hauteur donnée (pt) : marque de 10 pt, interligne multiple ajusté."""
        return para(conteneur, multiple=hauteur / ligne(10), taille=10, premier=premier)

    compteur = [5000]

    def ancre(par, donnees, largeur, hauteur, x, y, page=False, devant=False):
        """Image ancrée derrière le texte (ou devant), position en points depuis la page ou le paragraphe."""
        rid, _ = par.part.get_or_add_image(io.BytesIO(donnees))
        compteur[0] += 1
        n = compteur[0]
        xml = (
            '<w:drawing {ns}><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{rh}" '
            'behindDoc="{derriere}" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="{rh_}"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="{rv}"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
            '<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            '<wp:docPr id="{n}" name="Fond {n}" descr="Fond décoratif"/><wp:cNvGraphicFramePr/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
            '<pic:pic><pic:nvPicPr><pic:cNvPr id="{n}" name="fond{n}.png"/><pic:cNvPicPr/></pic:nvPicPr>'
            '<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
            "</pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing>"
        ).format(
            ns=nsdecls("w", "wp", "a", "pic", "r"), rh=n, derriere="0" if devant else "1",
            rh_="page" if page else "column", rv="page" if page else "paragraph", x=round(x * EMU), y=round(y * EMU), cx=round(largeur * EMU), cy=round(hauteur * EMU), n=n, rid=rid,
        )
        run = par.add_run()
        if page:
            # Google Docs garde une image ancrée à la page comme un caractère de la taille du passage :
            # à 10 pt, il donnait 4 mm de haut au paragraphe minuscule qui porte le bandeau de la page 2.
            run.font.size = Pt(1)
        run._r.append(parse_xml(xml))

    def image(par, donnees, largeur):
        run = par.add_run()
        run.font.size = Pt(1)
        run.add_picture(io.BytesIO(donnees), width=Pt(largeur))
        return run

    def tableau(conteneur, largeurs, hauteur=None, filet=None):
        """Tableau à largeurs fixes, sans bordure ni marge de cellule ; hauteur minimale de rangée (pt)."""
        t = conteneur.add_table(rows=1, cols=len(largeurs)) if conteneur is doc else conteneur.add_table(1, len(largeurs), Pt(sum(largeurs)))
        t.autofit = False
        tblpr = t._element.tblPr
        for nom in ("w:tblLayout", "w:tblW", "w:tblInd", "w:tblCellMar", "w:tblBorders", "w:tblStyle", "w:tblLook"):
            for e in tblpr.findall(qn(nom)):
                tblpr.remove(e)
        tblpr.append(parse_xml('<w:tblW {} w:w="{}" w:type="dxa"/>'.format(nsdecls("w"), round(sum(largeurs) * 20))))
        tblpr.append(parse_xml('<w:tblInd {} w:w="0" w:type="dxa"/>'.format(nsdecls("w"))))
        bords = "".join('<w:{} w:val="nil"/>'.format(b) for b in ("left", "bottom", "right", "insideH", "insideV"))
        haut = '<w:top w:val="single" w:sz="6" w:space="0" w:color="{}"/>'.format(filet.lstrip("#")) if filet else '<w:top w:val="nil"/>'
        tblpr.append(parse_xml("<w:tblBorders {}>{}{}</w:tblBorders>".format(nsdecls("w"), haut, bords)))
        tblpr.append(parse_xml('<w:tblLayout {} w:type="fixed"/>'.format(nsdecls("w"))))
        tblpr.append(parse_xml("<w:tblCellMar {}>{}</w:tblCellMar>".format(
            nsdecls("w"), "".join('<w:{} w:w="0" w:type="dxa"/>'.format(b) for b in ("top", "left", "bottom", "right")))))
        for gc, largeur in zip(t._element.tblGrid.findall(qn("w:gridCol")), largeurs):
            gc.set(qn("w:w"), str(round(largeur * 20)))
        for cel, largeur in zip(t.rows[0].cells, largeurs):
            cel.width = Pt(largeur)
        if hauteur:
            t.rows[0].height = Pt(hauteur)
            t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        return t

    def cellule(t, i, centre=False):
        cel = t.rows[0].cells[i]
        cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if centre else WD_CELL_VERTICAL_ALIGNMENT.TOP
        return cel

    def lien(par, url, libelle, taille, couleur, gras=False):
        rid = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        run = police(par.add_run(libelle), taille, gras=gras, couleur=couleur)
        h.append(run._element)
        par._p.append(h)

    num_puce = puces(doc)

    def puce(conteneur, contenu, taille, multiple, retrait, apres=0.0, decalage=0.0):
        """Point de liste : puce à 2,25 pt du bord (plus décalage), texte au retrait (plus décalage)."""
        par = para(conteneur, apres=apres, multiple=multiple, taille=taille)
        ppr = par._p.get_or_add_pPr()
        ppr.insert(0, parse_xml('<w:numPr {}><w:ilvl w:val="0"/><w:numId w:val="{}"/></w:numPr>'.format(nsdecls("w"), num_puce)))
        par.paragraph_format.left_indent = Pt(retrait + decalage)
        par.paragraph_format.first_line_indent = Pt(2.25 - retrait)
        texte(par, contenu, taille)
        return par

    def intertitre(conteneur, libelle, avant=0.0, apres=7.5, premier=False):
        par = para(conteneur, avant=avant, apres=apres, multiple=1.011, premier=premier)
        image(par, carre(c["couleur"]), 7.5)
        police(par.add_run(" "), 12)
        police(par.add_run(typographie(libelle)), 12, gras=True, couleur=NOIR)
        return par

    # ------------------------------------------------ page 1
    pastilles = [chasse_gras(typographie(t), 10) + 21 for t in (p["etiquette"], p["niveau"])]

    # Titre centré dans le bandeau, comme dans le modèle (data-fit 42,36,26) : 31,5 pt sur une ligne,
    # 27 pt au plus sur deux, 19,5 pt au moins ; la marge haute de la page place son centre.
    intitule = typographie(p["intitule"])
    largeur_titre = UTILE - (82.5 if p["sticker"] else 0)
    taille_titre = 31.5
    if nombre_de_lignes(intitule, taille_titre, largeur_titre) > 1:
        taille_titre = 27.0
        while nombre_de_lignes(intitule, taille_titre, largeur_titre) > 2 and taille_titre > 19.5:
            taille_titre -= 0.75
    hauteur_titre = ligne(taille_titre, 0.913) * nombre_de_lignes(intitule, taille_titre, largeur_titre)
    haut_titre = HAUT_1 + ligne(31.5, 0.913) / 2 - hauteur_titre / 2
    s1.top_margin = Pt(haut_titre)
    titre = para(doc, apres=0, multiple=0.913, droite=82.5 if p["sticker"] else 0, premier=True)
    ancre(titre, bandeau_1(c["couleur"], pastilles), PAGE_L, PASTILLES_Y + PASTILLE_H + 3, 0, 0, page=True)
    if p["sticker"]:
        donnees = ch.sticker_png(c["sticker"], c["icone"], taille=300, rotation=8)
        cote = 26.5 * 72 / 25.4 / 0.652
        ancre(titre, donnees, cote, cote, 183.1 * 72 / 25.4 - cote / 2, 31.2 * 72 / 25.4 - cote / 2, page=True, devant=True)
    police(titre.add_run(intitule), taille_titre, gras=True, couleur=c["surCouleur"])
    titre.paragraph_format.space_after = Pt(max(0, PASTILLES_Y - haut_titre - hauteur_titre))

    t = tableau(doc, [pastilles[0], 6, pastilles[1], UTILE - pastilles[0] - pastilles[1] - 6], hauteur=PASTILLE_H)
    for i, valeur in ((0, p["etiquette"]), (2, p["niveau"])):
        pp = para(cellule(t, i, centre=True), multiple=0.992, alignement=WD_ALIGN_PARAGRAPH.CENTER, premier=True)
        police(pp.add_run(typographie(valeur)), 10, gras=True, couleur=NOIR)
    for i in (1, 3):
        espace(cellule(t, i), 1, premier=True)

    texte(para(doc, avant=19.5, multiple=1.322, droite=49.5), p["description"], 10.5)
    police(para(doc, avant=18, apres=6, multiple=0.983).add_run("Objectifs"), 13, gras=True, couleur=NOIR)
    for n, o in enumerate(p["objectifs"]):
        puce(doc, o, 10.5, 1.237, 15, apres=0 if n == len(p["objectifs"]) - 1 else 2.25)

    # Encadré prérequis / public / durée : fond arrondi porté par un paragraphe de 16,5 pt.
    porteur = espace(doc, 16.5)
    ancre(porteur, fond_arrondi(UTILE, 73.5, g["fondGris"]), UTILE, 73.5, 0, 16.5)
    t = tableau(doc, [215.25, 180, 104.25], hauteur=73.5)
    for i, (libelle, valeur) in enumerate((("Prérequis", p["prerequis"]), ("Public visé", p["public"]), ("Durée", p["duree"]))):
        cel = cellule(t, i)
        gauche, droite = (16.5, 6.75) if i == 0 else ((6.75, 16.5) if i == 2 else (6.75, 6.75))
        pt_ = para(cel, avant=13.5, apres=4.5, multiple=1.045, gauche=gauche, droite=droite, premier=True)
        police(pt_.add_run(libelle), 11, gras=True, couleur=NOIR)
        texte(para(cel, apres=13.5, multiple=1.205, gauche=gauche, droite=droite), valeur, 10)

    police(para(doc, avant=33, apres=9, multiple=0.983).add_run("Programme"), 13, gras=True, couleur=NOIR)
    charge = p["_charge"]
    taille = 10 if charge <= 1 else max(8, 10 / min(charge, 1.25))
    coupure = p["_coupure"]
    t = tableau(doc, [UTILE / 2, UTILE / 2])
    numero = 0
    for col, parties in enumerate((p["programme"][:coupure], p["programme"][coupure:])):
        cel = cellule(t, col)
        gauche, droite = (0, 15) if col == 0 else (15, 0)
        if not parties:
            espace(cel, 1, premier=True)
        for k, partie in enumerate(parties):
            numero += 1
            pt_ = para(cel, avant=0 if k == 0 else 9, apres=2.25, multiple=1.18, gauche=gauche, droite=droite, premier=k == 0)
            pt_.paragraph_format.keep_with_next = True
            police(pt_.add_run("{:02d}".format(numero)), taille, couleur=c["encre"], famille=POLICE_EXTRA)
            police(pt_.add_run(" "), taille)
            libelle = typographie(partie["titre"]) + (" · " + typographie(partie["duree"]) if partie.get("duree") else "")
            police(pt_.add_run(libelle), taille, gras=True, italique=True, couleur=NOIR)
            for point in partie["points"]:
                pp = puce(cel, point, taille, 1.18, 13.5, decalage=gauche)
                if droite:
                    pp.paragraph_format.right_indent = Pt(droite)

    # Pied de la page 1 : filet, adresses, logo.
    pied = s1.footer
    pied.is_linked_to_previous = False
    espace(pied, 1.2, premier=True)
    t = tableau(pied, [UTILE - 60, 60], hauteur=PIED_1, filet=g["filet"])
    pg = para(cellule(t, 0, centre=True), multiple=1.0, premier=True)
    lien(pg, ch.CONTACT["url"], ch.CONTACT["site"], 8.5, PIED)
    police(pg.add_run(" · "), 8.5, couleur=PIED)
    lien(pg, "mailto:" + ch.CONTACT["email"], ch.CONTACT["email"], 8.5, PIED)
    pd = para(cellule(t, 1, centre=True), alignement=WD_ALIGN_PARAGRAPH.RIGHT, taille=1, premier=True)
    image(pd, ch.chemin_logo("couleur").read_bytes(), 32.6)
    para(pied, taille=1)  # Word exige un paragraphe après un tableau

    # ------------------------------------------------ page 2
    m = p["_modalites"]
    if m:
        s2 = doc.add_section(WD_SECTION.NEW_PAGE)
        mise_en_page(s2, HAUT_2, BAS_2, DISTANCE_PIED_2)
        libelle_pastille = "MODALITÉS ET INFOS PRATIQUES"
        pastille_l = chasse_gras(libelle_pastille, 9) + 61
        pastille_h, pastille_y = 18.4, 6.5 * 72 / 25.4

        # Titre du bandeau : 16,5 pt, puis 15 et 14,25 pt (data-fit 22,20,19) ; sur deux lignes, le
        # bandeau grandit d'autant et la pastille reste au milieu de la rangée.
        largeur_t = UTILE - pastille_l - 18
        taille2 = next((t_ for t_ in (16.5, 15.0) if nombre_de_lignes(intitule, t_, largeur_t) == 1), 14.25)
        surplus = max(0.0, ligne(taille2, 0.981) * nombre_de_lignes(intitule, taille2, largeur_t) - 19.5)
        tete = para(doc, multiple=1.0, taille=1)
        ancre(tete, bandeau_2(c["couleur"], c["encre"], pastille_l, pastille_y + surplus / 2, pastille_h, BANDEAU_2 + surplus),
              PAGE_L, BANDEAU_2 + surplus, 0, 0, page=True)
        t = tableau(doc, [largeur_t, 18, pastille_l], hauteur=19.5)
        pt_ = para(cellule(t, 0, centre=True), multiple=0.981, premier=True)
        police(pt_.add_run(intitule), taille2, gras=True, couleur=c["surCouleur"])
        espace(cellule(t, 1), 1, premier=True)
        # Capitales sans jambage : centrée seule, la ligne place le texte 0,86 mm trop bas dans le badge.
        # 4,9 pt d'espace après le remontent de moitié, au centre optique du badge.
        pp = para(cellule(t, 2, centre=True), apres=4.9, multiple=1.0, alignement=WD_ALIGN_PARAGRAPH.CENTER, premier=True)
        police(pp.add_run(libelle_pastille), 9, gras=True, couleur="#FFFFFF")

        # Effectif, lieu, tarif : trois cartes grises.
        carte_l, ecart = (UTILE - 21) / 3, 10.5
        porteur = espace(doc, 40.7)
        fond = fond_arrondi(carte_l, 48, g["fondGris"])
        for k in range(3):
            ancre(porteur, fond, carte_l, 48, k * (carte_l + ecart), 40.7)
        t = tableau(doc, [carte_l, ecart, carte_l, ecart, carte_l], hauteur=48)
        for k, (libelle, valeur) in enumerate((("Effectif", m["effectif"]), ("Lieu", m["lieu"]), ("Tarif", m["tarif"]))):
            cel = cellule(t, 2 * k)
            pl = para(cel, avant=13, apres=3.75, multiple=0.977, gauche=13.5, droite=13.5, premier=True)
            police(pl.add_run(libelle.upper()), 8.5, couleur=c["encre"], famille=POLICE_EXTRA)
            police(para(cel, apres=6, multiple=1.045, gauche=13.5, droite=13.5).add_run(typographie(valeur)), 11, gras=True, couleur=NOIR)
        for k in (1, 3):
            espace(cellule(t, k), 1, premier=True)

        formats = m["formats"]
        if formats:
            intertitre(doc, "Modalités et lieu de la formation", avant=23.5, apres=0)
            n = len(formats)
            carte_l = (UTILE - (n - 1) * ecart) / n
            porteur = espace(doc, 7.5)
            largeurs = []
            for k, f in enumerate(formats):
                ancre(porteur, fond_arrondi(carte_l, 165, "#FFFFFF", rayon_pt=10.5, bordure="#DEE2E6", barre=c["couleur"], icone=f["icone"],
                                            encre=c["encre"], teinte=c["teinte"]), carte_l, 165, k * (carte_l + ecart), 7.5)
                largeurs += [carte_l, ecart]
            t = tableau(doc, largeurs[:-1], hauteur=165)
            for k, f in enumerate(formats):
                cel = cellule(t, 2 * k)
                pf_ = para(cel, avant=48, apres=3.75, multiple=1.034, gauche=12.75, droite=12.75, premier=True)
                police(pf_.add_run(typographie(f["titre"])), 10.5, gras=True, couleur=NOIR)
                texte(para(cel, apres=6, multiple=1.26, gauche=12.75, droite=12.75), f["texte"], 9.5)
            for k in range(1, 2 * n - 1, 2):
                espace(cellule(t, k), 1, premier=True)

        if m["accompagnement"]:
            porteur = espace(doc, 10.5)
            ancre(porteur, fond_arrondi(UTILE, 34.5, c["teinte"], icone="handshake", encre=c["encre"]), UTILE, 34.5, 0, 10.5)
            t = tableau(doc, [UTILE], hauteur=34.5)
            pa = para(cellule(t, 0, centre=True), multiple=1.248, gauche=37.5, droite=13.5, premier=True)
            police(pa.add_run("Accompagnement personnalisé : "), 10, gras=True, couleur=NOIR)
            texte(pa, m["accompagnement"], 10, couleur=NOIR)

        espace(doc, 21)
        t = tableau(doc, [UTILE / 2, UTILE / 2])
        gauche = cellule(t, 0)
        intertitre(gauche, "Méthodes et moyens pédagogiques", premier=True).paragraph_format.right_indent = Pt(15)
        texte(para(gauche, multiple=1.286, droite=15), m["methodes"], 10)
        if m["moyens"]:
            espace(gauche, 6)
            for k, moyen in enumerate(m["moyens"]):
                puce(gauche, moyen, 10, 1.248, 13.5, apres=0 if k == len(m["moyens"]) - 1 else 1.5).paragraph_format.right_indent = Pt(15)
        droite = cellule(t, 1)
        premier = True
        for libelle, valeur in (("Encadrement", m["encadrement"]), ("Évaluation des acquis", m["evaluation"]), ("Mode de validation", m["validation"])):
            if not valeur:
                continue
            intertitre(droite, libelle, avant=0 if premier else 15, premier=premier).paragraph_format.left_indent = Pt(15)
            premier = False
            texte(para(droite, multiple=1.286, gauche=15), valeur, 10)
        if premier:
            espace(droite, 1, premier=True)

        # Pied de la page 2 : boîte de contact, filet, RIDET et DFPC.
        pied2 = s2.footer
        pied2.is_linked_to_previous = False
        porteur = espace(pied2, 1.2, premier=True)
        ancre(porteur, fond_arrondi(UTILE, PIED_2_CONTACT, c["teinte"], rayon_pt=12), UTILE, PIED_2_CONTACT, 0, 1.2)
        t = tableau(pied2, [UTILE - 102.75, 102.75], hauteur=PIED_2_CONTACT)
        cel = cellule(t, 0, centre=True)
        pq = para(cel, apres=4.5, multiple=1.3125, gauche=21, droite=18, premier=True)
        police(pq.add_run("Une question sur cette formation ?"), 13.5, gras=True, couleur=NOIR)
        pc = para(cel, multiple=1.258, gauche=21, droite=18)
        lien(pc, "mailto:" + ch.CONTACT["email"], ch.CONTACT["email"], 10, TEXTE)
        police(pc.add_run(" · "), 10)
        lien(pc, "tel:" + ch.CONTACT["telephone"].replace(" ", ""), typographie(ch.CONTACT["telephone"]), 10, TEXTE)
        lien(para(cel, multiple=1.258, gauche=21, droite=18), ch.CONTACT["url"], ch.CONTACT["site"], 10, c["encre"], gras=True)
        pd = para(cellule(t, 1, centre=True), alignement=WD_ALIGN_PARAGRAPH.RIGHT, droite=21, taille=1, premier=True)
        image(pd, ch.chemin_logo("couleur").read_bytes(), 62.4)
        espace(pied2, 19.8)
        t = tableau(pied2, [UTILE], hauteur=PIED_2_FILET, filet=g["filet"])
        pr = para(cellule(t, 0, centre=True), multiple=1.0, premier=True)
        police(pr.add_run("RIDET {} · N° DFPC {}".format(ch.CONTACT["ridet"], ch.CONTACT["dfpc"])), 8.5, couleur=PIED)
        para(pied2, taille=1)

    # Paragraphe final minuscule (Word en exige un après un tableau).
    para(doc, taille=1)
    # Le paragraphe qui porte le saut de section, après le programme, est réduit lui aussi.
    for p_ in doc.element.body.iter(qn("w:p")):
        ppr = p_.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None and not p_.xpath(".//w:t"):
            espacement = parse_xml('<w:spacing {} w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'.format(nsdecls("w")))
            ppr.insert(0, espacement)
            ppr.insert(ppr.index(ppr.find(qn("w:sectPr"))), parse_xml('<w:rPr {}><w:sz w:val="2"/></w:rPr>'.format(nsdecls("w"))))

    mode_compatibilite_15(doc)
    tampon = io.BytesIO()
    doc.save(tampon)
    # Allégé (gabarit python-docx, images) pour Google Docs et les connecteurs Drive : voir docx_google_docs.py.
    from docx_google_docs import preparer_docx

    return preparer_docx(tampon.getvalue())[0]


def mise_en_page(section, haut, bas, distance_pied) -> None:
    from docx.shared import Pt

    section.page_width, section.page_height = Pt(PAGE_L), Pt(PAGE_H)
    section.left_margin = section.right_margin = Pt(MARGE)
    section.top_margin, section.bottom_margin = Pt(haut), Pt(bas)
    section.header_distance, section.footer_distance = Pt(0), Pt(distance_pied)


def puces(doc) -> int:
    """Liste à puces « • » en Arial, sans taille propre (celle du texte) : numId à utiliser."""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    numerotation = doc.part.numbering_part.element
    numerotation.append(parse_xml(
        '<w:abstractNum {} w:abstractNumId="90"><w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="300" w:hanging="255"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/><w:color w:val="3B3B42"/></w:rPr>'
        "</w:lvl></w:abstractNum>".format(nsdecls("w"))
    ))
    # Les w:num viennent après tous les w:abstractNum.
    numerotation.append(parse_xml('<w:num {} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'.format(nsdecls("w"))))
    abstraits = numerotation.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}abstractNum")
    for a in abstraits:
        numerotation.remove(a)
    for i, a in enumerate(abstraits):
        numerotation.insert(i, a)
    return 90


def mode_compatibilite_15(doc) -> None:
    """Mode Word 2013+ : le retrait d'un tableau se mesure à son bord (et non au texte de la cellule)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    reglages = doc.settings.element
    compat = reglages.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        reglages.append(compat)
    for e in compat.findall(qn("w:compatSetting")):
        if e.get(qn("w:name")) == "compatibilityMode":
            compat.remove(e)
    e = OxmlElement("w:compatSetting")
    e.set(qn("w:name"), "compatibilityMode")
    e.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    e.set(qn("w:val"), "15")
    compat.append(e)
