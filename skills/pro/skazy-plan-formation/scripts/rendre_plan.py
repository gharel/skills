#!/usr/bin/env python3
"""Plan de formation Skazy Formation : plan.json → plan.html (A4) et plan.docx (Word, Google Docs).

Usage :
    python scripts/rendre_plan.py plan.json [--sortie DOSSIER] [--formats html,docx] [--verifier]

Le JSON suit references/format-plan.md. Le HTML reproduit le modèle « 3. Plan de formation »
du Skazy Formation Design System (Claude Design) ; le DOCX en reprend la mise en page avec de vrais
en-têtes et pieds de page (bandeau couleur pleine largeur), lisibles dans Word et Google Docs.
"""

from __future__ import annotations

import argparse
import io
import math
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import charte as ch  # noqa: E402
from charte import esc, html_gras, typographie  # noqa: E402

FORMAT = "skazy-plan"

# ---------------------------------------------------------------- lecture et contrôle


def normaliser_plan(brut: dict) -> tuple[dict, ch.Rapport]:
    """Complète les valeurs par défaut et contrôle le contenu (règles de references/regles-redaction.md)."""
    r = ch.Rapport()
    p = dict(brut.get("plan", brut))
    if p.get("format") not in (None, FORMAT):
        r.erreur("format attendu « {} », reçu « {} ».".format(FORMAT, p.get("format")))
    for champ in ("intitule", "categorie", "description", "objectifs", "prerequis", "public", "duree", "programme"):
        if not p.get(champ):
            r.erreur("champ obligatoire manquant : « {} ».".format(champ))
    if not r.ok:
        return p, r
    try:
        p["_cat"] = ch.categorie(p["categorie"])
    except ValueError as e:
        r.erreur(str(e))
        return p, r
    try:
        p["niveau"] = ch.niveau(p.get("niveau"))
    except ValueError as e:
        r.erreur(str(e))
        return p, r
    p["categorie"] = p["_cat"]["nom"]
    p["etiquette"] = p.get("etiquette") or p["categorie"]
    p["sticker"] = p.get("sticker", True) is not False

    if len(p["intitule"]) > 60:
        r.avertir("intitulé long ({} caractères) : visez 60 au plus pour tenir sur deux lignes.".format(len(p["intitule"])))
    if len(p["description"]) > 420:
        r.avertir("description longue ({} caractères) : deux à quatre lignes suffisent.".format(len(p["description"])))
    if ch.normaliser(p["description"]).startswith(("cette formation", "ce parcours", "cette session")):
        r.avertir("la description commence par « Cette formation » : commencez par un verbe d'action.")
    n = len(p["objectifs"])
    if not 3 <= n <= 5:
        r.avertir("{} objectifs : le modèle en prévoit 3 à 5.".format(n))
    vagues = [o for o in p["objectifs"] if ch.normaliser(ch.sans_gras(o)).split(" ")[0] in {"comprendre", "connaitre", "savoir", "apprendre", "decouvrir"}]
    for o in vagues:
        r.avertir("objectif peu observable (Bloom) : « {} ». Préférez identifier, utiliser, créer, analyser…".format(o))

    programme = []
    for i, partie in enumerate(p["programme"], 1):
        if isinstance(partie, str):
            partie = {"titre": partie, "points": []}
        if not partie.get("titre"):
            r.erreur("partie {} du programme sans titre.".format(i))
        points = partie.get("points") or []
        if not 2 <= len(points) <= 6:
            r.avertir("partie {} « {} » : {} points (2 à 6 attendus).".format(i, partie.get("titre", ""), len(points)))
        programme.append({"titre": partie.get("titre", ""), "points": points, "duree": partie.get("duree", "")})
    p["programme"] = programme
    if not 2 <= len(programme) <= 8:
        r.avertir("{} parties au programme : 2 à 8 attendues.".format(len(programme)))

    modalites = p.get("modalites", {})
    if modalites is False or p.get("page2") is False:
        p["_modalites"] = None
    else:
        defaut = ch.CHARTE["modalitesParDefaut"]
        p["_modalites"] = {k: (modalites or {}).get(k) or defaut[k] for k in defaut}

    p["_charge"] = estimer_charge(p)
    if p["_charge"] > 1.0:
        r.avertir(
            "le contenu de la page 1 dépasse d'environ {:.0f} % la place disponible : la taille du programme "
            "sera réduite, raccourcissez les points ou supprimez une partie.".format((p["_charge"] - 1) * 100)
        )
    return p, r


def lignes(texte: str, largeur_car: int) -> int:
    return max(1, math.ceil(len(ch.sans_gras(texte)) / largeur_car))


def hauteur_parties(parties: list[dict]) -> float:
    """Hauteur (px) d'une colonne du programme à 13 px : titre + points + marge."""
    h = 0.0
    for partie in parties:
        h += 18.5 * lignes(partie["titre"], 40) + 3
        h += sum(18.5 * lignes(pt, 44) for pt in partie["points"])
        h += 12
    return h


def couper_programme(programme: list[dict]) -> int:
    """Indice de coupure qui équilibre les deux colonnes (ordre de lecture conservé)."""
    meilleur, ecart = len(programme), float("inf")
    for i in range(len(programme) + 1):
        a, b = hauteur_parties(programme[:i]), hauteur_parties(programme[i:])
        if max(a, b) < ecart:
            meilleur, ecart = i, max(a, b)
    return meilleur


def estimer_charge(p: dict) -> float:
    """Rapport entre la hauteur estimée du contenu de la page 1 et la place disponible (1.0 = plein)."""
    fixe = 150 + 40 + 84 + 58 + 78
    fixe += 21.7 * lignes(p["description"], 92)
    fixe += 52 + sum(23.3 * lignes(o, 96) for o in p["objectifs"])
    infos = max(lignes(p["prerequis"], 30), lignes(p["public"], 26), 1)
    fixe += 18.9 * infos
    i = couper_programme(p["programme"])
    programme = max(hauteur_parties(p["programme"][:i]), hauteur_parties(p["programme"][i:]))
    return (fixe + programme) / 1123


# ---------------------------------------------------------------- HTML


CSS_PLAN = """
*{box-sizing:border-box}
html,body{margin:0}
body{background:#E9E9EC;font-family:Raleway,system-ui,sans-serif;color:#1F1F24;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.page{position:relative;width:794px;height:1123px;margin:32px auto;background:#fff;display:flex;flex-direction:column;overflow:hidden;box-shadow:0 12px 40px rgba(0,0,0,.14)}
.bandeau{position:relative;background:var(--c);color:var(--co);padding:46px 64px 0;height:150px;flex-shrink:0}
.bandeau h1{margin:0;font-size:42px;font-weight:700;line-height:1.05;letter-spacing:-.01em;max-width:560px}
.pastilles{position:absolute;left:64px;bottom:-14px;display:flex;gap:8px}
.pastilles span{background:#F1F1F3;color:#1F1F24;font-size:13px;font-weight:700;padding:6px 14px;border-radius:999px;box-shadow:0 0 0 3px #fff}
.sticker{position:absolute;top:68px;right:52px;width:100px;height:100px;transform:rotate(8deg);background:#fff;border-radius:28px;padding:7px;box-shadow:0 6px 14px rgba(0,0,0,.18)}
.sticker div{width:100%;height:100%;background:var(--sk);border:3px solid #1F1F24;border-radius:22px;display:flex;align-items:center;justify-content:center}
.sticker i{width:48px;height:48px;background:#1F1F24;-webkit-mask:var(--si) center/contain no-repeat;mask:var(--si) center/contain no-repeat}
.corps{flex:1;display:flex;flex-direction:column;padding:40px 64px 0;min-height:0}
.description{margin:0;font-size:14px;line-height:1.55;color:#3B3B42;max-width:600px}
h2{margin:24px 0 8px;font-size:17px;font-weight:700}
.objectifs{margin:0;padding:0 0 0 20px;display:flex;flex-direction:column;gap:3px;font-size:14px;line-height:1.45;color:#3B3B42}
strong{color:#1F1F24}
.infos{margin-top:22px;background:#F4F4F5;border-radius:12px;padding:18px 22px;display:grid;grid-template-columns:minmax(0,1.5fr) minmax(0,1.3fr) minmax(108px,max-content);gap:18px}
.infos .t{font-size:15px;font-weight:700;margin-bottom:6px}
.infos .v{font-size:13px;line-height:1.45;color:#3B3B42}
.infos .duree .v{white-space:nowrap}
.h-prog{margin:26px 0 12px}
.programme{column-count:2;column-gap:40px;font-size:13px;line-height:1.42;color:#3B3B42}
.partie{break-inside:avoid;margin-bottom:12px}
.partie .titre{font-weight:700;font-style:italic;color:#1F1F24;margin-bottom:3px}
.partie .num{color:var(--ci);font-style:normal;font-weight:800;margin-right:6px}
.partie .duree-partie{font-style:normal;font-weight:600;color:#6E6E78;margin-left:6px}
.partie ul{margin:0;padding:0 0 0 18px}
.pied{margin-top:auto;padding:18px 0 26px;display:flex;align-items:center;justify-content:space-between;gap:16px;border-top:1px solid #EEEEF0;font-size:11px;color:#6E6E78}
.pied img{height:34px;width:auto}
.bandeau2{background:var(--c);color:var(--co);padding:26px 64px;display:flex;justify-content:space-between;align-items:baseline;gap:20px;flex-shrink:0}
.bandeau2 .i{font-size:22px;font-weight:700}
.bandeau2 .m{font-size:13px;font-weight:700;letter-spacing:.12em;text-transform:uppercase;white-space:nowrap}
.corps2{flex:1;display:flex;flex-direction:column;padding:36px 64px 0}
.deux{display:grid;grid-template-columns:1fr 1fr;gap:26px 40px}
.deux h2{margin:0 0 8px;font-size:16px;display:flex;gap:10px;align-items:center}
.deux h2 span{width:10px;height:10px;border-radius:2px;background:var(--c)}
.deux p{margin:0;font-size:13px;line-height:1.55;color:#3B3B42}
.cartes{margin-top:30px;display:grid;grid-template-columns:repeat(3,1fr);gap:14px}
.cartes>div{background:#F4F4F5;border-radius:12px;padding:16px 18px}
.cartes .t{font-size:11px;font-weight:800;letter-spacing:.12em;text-transform:uppercase;color:var(--ci);margin-bottom:6px}
.cartes .v{font-size:15px;font-weight:700}
.contact{margin-top:30px;background:var(--ct);border-radius:16px;padding:24px 28px;display:grid;grid-template-columns:1fr auto;gap:24px;align-items:center}
.contact .q{font-size:18px;font-weight:700;display:block;margin-bottom:6px}
.contact .c{font-size:13px;line-height:1.55;color:#3B3B42}
.contact a{color:var(--ci);font-weight:700}
.contact img{height:76px;width:auto}
.outils{position:fixed;top:16px;right:16px;z-index:10;display:flex;gap:8px}
.outils button{font:700 14px Raleway,system-ui,sans-serif;border:0;border-radius:999px;padding:10px 18px;background:#1F1F24;color:#fff;cursor:pointer;box-shadow:0 6px 18px rgba(0,0,0,.2)}
.outils button:focus-visible{outline:3px solid var(--c);outline-offset:2px}
@page{size:A4;margin:0}
@media print{body{background:none}.page{margin:0;box-shadow:none;break-after:page;width:210mm;height:297mm}.outils{display:none}}
"""

JS_AJUSTER = """
// Si le contenu déborde de la page 1, on réduit la taille du programme (13 px → 10 px au plus).
addEventListener('load', () => {
  const page = document.querySelector('.p1');
  const prog = page && page.querySelector('.programme');
  if (!prog) return;
  let taille = 13;
  while (page.scrollHeight > page.clientHeight + 1 && taille > 10) {
    taille -= 0.25;
    prog.style.fontSize = taille + 'px';
  }
});
"""


def rendre_html(p: dict) -> str:
    c = p["_cat"]
    variables = "--c:{};--ci:{};--ct:{};--co:{};--sk:{};--si:url({})".format(
        c["couleur"], c["encre"], c["teinte"], c["surCouleur"], c["sticker"], ch.icone_data_uri(c["icone"])
    )
    i = couper_programme(p["programme"])
    parties = []
    for n, partie in enumerate(p["programme"], 1):
        duree = '<span class="duree-partie">{}</span>'.format(esc(partie["duree"])) if partie.get("duree") else ""
        points = "".join("<li>{}</li>".format(html_gras(pt)) for pt in partie["points"])
        coupure = ' style="break-before:column"' if n == i + 1 and i not in (0, len(p["programme"])) else ""
        parties.append(
            '<div class="partie"{}><div class="titre"><span class="num">{:02d}</span>{}{}</div><ul>{}</ul></div>'.format(
                coupure, n, esc(partie["titre"]), duree, points
            )
        )
    sticker = '<div class="sticker" aria-hidden="true"><div><i></i></div></div>' if p["sticker"] else ""
    logo = ch.logo_svg_data_uri("couleur")
    page1 = """
<section class="page p1" aria-label="Page 1, programme">
<div class="bandeau"><h1>{intitule}</h1><div class="pastilles"><span>{etiquette}</span><span>{niveau}</span></div></div>
{sticker}
<div class="corps">
<p class="description">{description}</p>
<h2>Objectifs</h2>
<ul class="objectifs">{objectifs}</ul>
<div class="infos">
<div><div class="t">Prérequis</div><div class="v">{prerequis}</div></div>
<div><div class="t">Public visé</div><div class="v">{public}</div></div>
<div class="duree"><div class="t">Durée</div><div class="v">{duree}</div></div>
</div>
<h2 class="h-prog">Programme</h2>
<div class="programme">{programme}</div>
<div class="pied"><span>{site} · {email}</span><img src="{logo}" alt="Skazy Formation"></div>
</div>
</section>""".format(
        intitule=esc(p["intitule"]),
        etiquette=esc(p["etiquette"]),
        niveau=esc(p["niveau"]),
        sticker=sticker,
        description=html_gras(p["description"]),
        objectifs="".join("<li>{}</li>".format(html_gras(o)) for o in p["objectifs"]),
        prerequis=html_gras(p["prerequis"]),
        public=html_gras(p["public"]),
        duree=esc(p["duree"]),
        programme="\n".join(parties),
        site=ch.CONTACT["site"],
        email=ch.CONTACT["email"],
        logo=logo,
    )
    page2 = ""
    m = p["_modalites"]
    if m:
        page2 = """
<section class="page p2" aria-label="Page 2, modalités et informations pratiques">
<div class="bandeau2"><span class="i">{intitule}</span><span class="m">Modalités et infos pratiques</span></div>
<div class="corps2">
<div class="deux">
<div><h2><span></span>Méthodes pédagogiques</h2><p>{methodes}</p></div>
<div><h2><span></span>Évaluation</h2><p>{evaluation}</p></div>
</div>
<div class="cartes">
<div><div class="t">Effectif</div><div class="v">{effectif}</div></div>
<div><div class="t">Lieu</div><div class="v">{lieu}</div></div>
<div><div class="t">Tarif</div><div class="v">{tarif}</div></div>
</div>
<div class="contact"><div><span class="q">Une question sur cette formation&nbsp;?</span><span class="c">{email} · {tel}<br><a href="{url}">{site}</a></span></div><img src="{logo}" alt="Skazy Formation"></div>
<div class="pied"><span>RIDET {ridet} · N° DFPC {dfpc}</span></div>
</div>
</section>""".format(
            intitule=esc(p["intitule"]),
            methodes=html_gras(m["methodes"]),
            evaluation=html_gras(m["evaluation"]),
            effectif=esc(m["effectif"]),
            lieu=esc(m["lieu"]),
            tarif=esc(m["tarif"]),
            email=ch.CONTACT["email"],
            tel=esc(ch.CONTACT["telephone"]),
            url=ch.CONTACT["url"],
            site=ch.CONTACT["site"],
            ridet=esc(ch.CONTACT["ridet"]),
            dfpc=esc(ch.CONTACT["dfpc"]),
            logo=logo,
        )
    return """<!doctype html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{titre} · Plan de formation · Skazy Formation</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Raleway:ital,wght@0,400;0,500;0,600;0,700;0,800;1,600;1,700&amp;display=swap" rel="stylesheet">
<style>:root{{{variables}}}{css}</style>
</head>
<body>
<div class="outils"><button type="button" onclick="print()">Imprimer ou enregistrer en PDF</button></div>
<main>{page1}{page2}</main>
<script>{js}</script>
</body>
</html>
""".format(titre=esc(p["intitule"]), variables=variables, css=CSS_PLAN, page1=page1, page2=page2, js=JS_AJUSTER)


# ---------------------------------------------------------------- DOCX


def rendre_docx(p: dict) -> bytes:
    ch.exiger("docx", "PIL")
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    c = p["_cat"]
    LARGEUR, MARGE = 210, 17  # mm : marges de page nulles, retrait du texte de 17 mm (64 px du modèle)
    UTILE = LARGEUR - 2 * MARGE
    POLICE = ch.CHARTE["police"]
    gris = ch.COULEURS

    doc = Document()
    doc.core_properties.title = p["intitule"]
    doc.core_properties.author = "Skazy Formation"
    doc.core_properties.subject = "Plan de formation"

    # Styles de base : Raleway partout, retrait de 17 mm pour que le texte ajouté garde la mise en page.
    for nom in ("Normal", "List Bullet"):
        st = doc.styles[nom]
        st.font.name = POLICE
        st.font.size = Pt(10)
        st.font.color.rgb = RGBColor(*ch.rgb(gris["texte2"]))
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), POLICE)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
    doc.styles["Normal"].paragraph_format.left_indent = Mm(MARGE)
    doc.styles["Normal"].paragraph_format.right_indent = Mm(MARGE)

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(LARGEUR), Mm(297)
    section.left_margin = section.right_margin = Mm(0)
    section.top_margin = section.bottom_margin = Mm(0)
    section.header_distance = section.footer_distance = Mm(0)
    section.different_first_page_header_footer = True

    # ------------------------------------------------ utilitaires OOXML

    def police(run, taille, gras=False, italique=False, couleur=None, majuscules=False, espacement=None):
        run.font.name = POLICE
        run.font.size = Pt(taille)
        run.font.bold = gras
        run.font.italic = italique
        rpr = run._element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), POLICE)
        if couleur:
            run.font.color.rgb = RGBColor(*ch.rgb(couleur))
        if majuscules:
            run.font.all_caps = True
        if espacement is not None:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(int(espacement * 20)))
            rpr.append(sp)
        return run

    def texte(par, contenu, taille, couleur=gris["texte2"], gras_couleur=gris["texte"], **kw):
        for segment, est_gras in ch.segments_gras(contenu):
            police(par.add_run(segment), taille, gras=est_gras or kw.get("gras", False),
                   italique=kw.get("italique", False), couleur=gras_couleur if est_gras else couleur)
        return par

    def para(conteneur, retrait=True, avant=0, apres=0, interligne=None, alignement=None):
        par = conteneur.add_paragraph()
        pf = par.paragraph_format
        if not retrait:
            pf.left_indent = pf.right_indent = Mm(0)
        pf.space_before, pf.space_after = Pt(avant), Pt(apres)
        if interligne:
            pf.line_spacing = interligne
        if alignement:
            par.alignment = alignement
        return par

    def fond(cellule, couleur):
        tcpr = cellule._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), couleur.lstrip("#"))
        tcpr.append(shd)

    def marges_cellule(cellule, haut=0, bas=0, gauche=0, droite=0):
        tcpr = cellule._element.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for cote, val in (("top", haut), ("left", gauche), ("bottom", bas), ("right", droite)):
            e = OxmlElement("w:" + cote)
            e.set(qn("w:w"), str(int(val * 56.7)))  # mm → twips
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tcpr.append(mar)

    def bordures(table, haut=None):
        """Supprime les bordures du tableau ; haut = couleur d'un filet supérieur facultatif."""
        tblpr = table._element.tblPr
        b = OxmlElement("w:tblBorders")
        for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + cote)
            if cote == "top" and haut:
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "6")
                e.set(qn("w:color"), haut.lstrip("#"))
            else:
                e.set(qn("w:val"), "nil")
            b.append(e)
        tblpr.append(b)

    def tableau(conteneur, colonnes_mm, retrait_mm=0.0, largeur_mm=None):
        """Tableau à largeurs fixes (mm), sans bordure, décalé de retrait_mm."""
        largeur_mm = largeur_mm or sum(colonnes_mm)
        if hasattr(conteneur, "add_table") and conteneur is not doc:
            t = conteneur.add_table(rows=1, cols=len(colonnes_mm), width=Mm(largeur_mm))
        else:
            t = conteneur.add_table(rows=1, cols=len(colonnes_mm))
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        tblpr = t._element.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblpr.append(layout)
        tw = tblpr.find(qn("w:tblW"))
        if tw is None:
            tw = OxmlElement("w:tblW")
            tblpr.append(tw)
        tw.set(qn("w:w"), str(int(largeur_mm * 56.7)))
        tw.set(qn("w:type"), "dxa")
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(int(retrait_mm * 56.7)))
        ind.set(qn("w:type"), "dxa")
        tblpr.append(ind)
        mar = OxmlElement("w:tblCellMar")
        for cote in ("left", "right"):
            e = OxmlElement("w:" + cote)
            e.set(qn("w:w"), "0")
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblpr.append(mar)
        grille = t._element.tblGrid
        for gc, mm in zip(grille.findall(qn("w:gridCol")), colonnes_mm):
            gc.set(qn("w:w"), str(int(mm * 56.7)))
        for cellule, mm in zip(t.rows[0].cells, colonnes_mm):
            cellule.width = Mm(mm)
        bordures(t)
        if conteneur is not doc:
            # En-tête ou pied : le tableau passe avant le paragraphe vide créé par défaut.
            premier = conteneur.paragraphs[0]._p
            premier.addprevious(t._tbl)
            premier.getparent().append(premier)
        return t

    def premier_para(cellule):
        par = cellule.paragraphs[0]
        par.paragraph_format.left_indent = par.paragraph_format.right_indent = Mm(0)
        return par

    def puce(conteneur, contenu, taille, retrait_mm, apres=1.5):
        par = conteneur.add_paragraph(style="List Bullet")
        pf = par.paragraph_format
        pf.left_indent = Mm(retrait_mm)
        pf.first_line_indent = Mm(-4.5)
        pf.right_indent = Mm(0)
        pf.space_after = Pt(apres)
        pf.line_spacing = 1.2
        texte(par, contenu, taille)
        return par

    def image(par, donnees: bytes, hauteur_mm=None, largeur_mm=None):
        if largeur_mm:
            par.add_run().add_picture(io.BytesIO(donnees), width=Mm(largeur_mm))
        else:
            par.add_run().add_picture(io.BytesIO(donnees), height=Mm(hauteur_mm))

    def espace_apres(partie, points):
        """Dernier paragraphe (vide) d'un en-tête ou d'un pied : sert d'espacement, sans retrait."""
        par = partie.paragraphs[-1]
        pf = par.paragraph_format
        pf.left_indent = pf.right_indent = Mm(0)
        pf.space_before, pf.space_after = Pt(0), Pt(points)
        pf.line_spacing = 1.0
        police(par.add_run(""), 1)

    logo = ch.chemin_logo("couleur").read_bytes()

    # ------------------------------------------------ en-tête et pied de la page 1

    entete1 = section.first_page_header
    entete1.is_linked_to_previous = False
    t = tableau(entete1, [LARGEUR - 52, 52])
    rangee = t.rows[0]
    rangee.height = Mm(40)
    rangee.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    gauche, droite = rangee.cells
    for cel in (gauche, droite):
        fond(cel, c["couleur"])
        cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    marges_cellule(gauche, haut=11, gauche=MARGE, droite=4)
    marges_cellule(droite, haut=15, droite=13)
    titre = premier_para(gauche)
    titre.paragraph_format.line_spacing = 1.0
    police(titre.add_run(typographie(p["intitule"])), 24 if len(p["intitule"]) <= 40 else 21, gras=True, couleur=c["surCouleur"])
    pastilles = para(gauche, retrait=False, avant=7)
    for n, valeur in enumerate((p["etiquette"], p["niveau"])):
        if n:
            police(pastilles.add_run("  "), 9.5)
        r = police(pastilles.add_run("  {}  ".format(typographie(valeur))), 9.5, gras=True, couleur=gris["texte"])
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), gris["pastille"].lstrip("#"))
        r._element.get_or_add_rPr().append(shd)
    p_sticker = premier_para(droite)
    p_sticker.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if p["sticker"]:
        image(p_sticker, ch.sticker_png(c["sticker"], c["icone"], taille=170, rotation=8), 24)
    espace_apres(entete1, 6)

    pied1 = section.first_page_footer
    pied1.is_linked_to_previous = False
    t = tableau(pied1, [UTILE - 40, 40], retrait_mm=MARGE)
    bordures(t, haut=gris["filet"])
    g, d = t.rows[0].cells
    marges_cellule(g, haut=4.5, bas=6.5)
    marges_cellule(d, haut=3, bas=6)
    g.vertical_alignment = d.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    police(premier_para(g).add_run("{} · {}".format(ch.CONTACT["site"], ch.CONTACT["email"])), 8, couleur=gris["gris"])
    espace_apres(pied1, 0)
    pd = premier_para(d)
    pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    image(pd, logo, 9)

    # ------------------------------------------------ en-tête et pied des pages suivantes

    entete2 = section.header
    entete2.is_linked_to_previous = False
    t = tableau(entete2, [LARGEUR - 95, 95])
    g, d = t.rows[0].cells
    for cel in (g, d):
        fond(cel, c["couleur"])
        cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    marges_cellule(g, haut=7, bas=7, gauche=MARGE, droite=4)
    marges_cellule(d, haut=7, bas=7, droite=MARGE)
    police(premier_para(g).add_run(typographie(p["intitule"])), 13, gras=True, couleur=c["surCouleur"])
    pd = premier_para(d)
    pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    police(pd.add_run("Modalités et infos pratiques"), 8, gras=True, couleur=c["surCouleur"], majuscules=True, espacement=1.2)
    espace_apres(entete2, 6)

    pied2 = section.footer
    pied2.is_linked_to_previous = False
    t = tableau(pied2, [UTILE], retrait_mm=MARGE)
    bordures(t, haut=gris["filet"])
    cel = t.rows[0].cells[0]
    marges_cellule(cel, haut=4.5, bas=7)
    police(premier_para(cel).add_run("RIDET {} · N° DFPC {}".format(ch.CONTACT["ridet"], ch.CONTACT["dfpc"])), 8, couleur=gris["gris"])
    espace_apres(pied2, 0)

    # ------------------------------------------------ page 1 : contenu

    corps = doc
    description = corps.paragraphs[0] if corps.paragraphs else corps.add_paragraph()
    description.paragraph_format.space_before = Pt(16)
    description.paragraph_format.line_spacing = 1.45
    description.paragraph_format.right_indent = Mm(MARGE + 15)
    texte(description, p["description"], 10.5)

    def titre2(contenu, avant=14, apres=4):
        par = para(corps, avant=avant, apres=apres)
        police(par.add_run(contenu), 12.5, gras=True, couleur=gris["texte"])
        return par

    titre2("Objectifs")
    for o in p["objectifs"]:
        puce(corps, o, 10.5, MARGE + 5, apres=1)

    para(corps, avant=0, apres=0).paragraph_format.space_after = Pt(8)
    t = tableau(corps, [UTILE * 0.38, UTILE * 0.34, UTILE * 0.28], retrait_mm=MARGE)
    for n, (libelle, valeur) in enumerate((("Prérequis", p["prerequis"]), ("Public visé", p["public"]), ("Durée", p["duree"]))):
        cel = t.rows[0].cells[n]
        fond(cel, gris["fondGris"])
        marges_cellule(cel, haut=4.5, bas=4.5, gauche=6 if n == 0 else 2.4, droite=2.4 if n < 2 else 6)
        police(premier_para(cel).add_run(libelle), 11, gras=True, couleur=gris["texte"])
        pv = para(cel, retrait=False, avant=3, interligne=1.3)
        texte(pv, valeur, 9.75)

    titre2("Programme", avant=16, apres=6)
    charge = p["_charge"]
    taille_prog = 9.75 if charge <= 1 else max(8, 9.75 / min(charge, 1.25))
    i = couper_programme(p["programme"])
    colonnes = [p["programme"][:i], p["programme"][i:]]
    t = tableau(corps, [(UTILE - 10) / 2, 10, (UTILE - 10) / 2], retrait_mm=MARGE)
    numero = 0
    for col, cel in zip(colonnes, (t.rows[0].cells[0], t.rows[0].cells[2])):
        cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        premier = True
        for partie in col:
            numero += 1
            pt = premier_para(cel) if premier else para(cel, retrait=False, avant=7)
            premier = False
            police(pt.add_run("{:02d}  ".format(numero)), taille_prog, gras=True, couleur=c["encre"])
            police(pt.add_run(typographie(partie["titre"])), taille_prog, gras=True, italique=True, couleur=gris["texte"])
            if partie.get("duree"):
                police(pt.add_run("  " + typographie(partie["duree"])), taille_prog - 1, gras=True, couleur=gris["gris"])
            pt.paragraph_format.space_after = Pt(2)
            pt.paragraph_format.keep_with_next = True
            for point in partie["points"]:
                puce(cel, point, taille_prog, 4.5, apres=0.5)

    # ------------------------------------------------ page 2 : modalités

    m = p["_modalites"]
    if m:
        saut = para(corps)
        saut.add_run().add_break(WD_BREAK.PAGE)
        saut.paragraph_format.space_after = Pt(10)
        t = tableau(corps, [(UTILE - 10) / 2, 10, (UTILE - 10) / 2], retrait_mm=MARGE)
        for cel, (libelle, valeur) in zip((t.rows[0].cells[0], t.rows[0].cells[2]), (("Méthodes pédagogiques", m["methodes"]), ("Évaluation", m["evaluation"]))):
            ph = premier_para(cel)
            police(ph.add_run("■  "), 10, couleur=c["couleur"])
            police(ph.add_run(libelle), 12, gras=True, couleur=gris["texte"])
            texte(para(cel, retrait=False, avant=5, interligne=1.45), valeur, 9.75)

        para(corps).paragraph_format.space_after = Pt(14)
        ecart = 3.7
        largeur_carte = (UTILE - 2 * ecart) / 3
        t = tableau(corps, [largeur_carte, ecart, largeur_carte, ecart, largeur_carte], retrait_mm=MARGE)
        for cel, (libelle, valeur) in zip(t.rows[0].cells[0::2], (("Effectif", m["effectif"]), ("Lieu", m["lieu"]), ("Tarif", m["tarif"]))):
            fond(cel, gris["fondGris"])
            marges_cellule(cel, haut=4.2, bas=4.2, gauche=4.8, droite=4.8)
            police(premier_para(cel).add_run(libelle), 8, gras=True, couleur=c["encre"], majuscules=True, espacement=1.2)
            police(para(cel, retrait=False, avant=3).add_run(typographie(valeur)), 11, gras=True, couleur=gris["texte"])

        para(corps).paragraph_format.space_after = Pt(14)
        t = tableau(corps, [UTILE - 45, 45], retrait_mm=MARGE)
        g, d = t.rows[0].cells
        for cel in (g, d):
            fond(cel, c["teinte"])
            cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        marges_cellule(g, haut=6.3, bas=6.3, gauche=7.4)
        marges_cellule(d, haut=4, bas=4, droite=7.4)
        police(premier_para(g).add_run("Une question sur cette formation ?"), 13.5, gras=True, couleur=gris["texte"])
        pc = para(g, retrait=False, avant=4, interligne=1.4)
        police(pc.add_run("{} · {}".format(ch.CONTACT["email"], typographie(ch.CONTACT["telephone"]))), 9.75)
        pc.add_run().add_break()
        lien(pc, ch.CONTACT["url"], ch.CONTACT["site"], police, c["encre"])
        pd = premier_para(d)
        pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        image(pd, logo, largeur_mm=34)

    # Paragraphe final minuscule (Word en exige un après un tableau).
    fin = para(corps)
    police(fin.add_run(""), 2)
    _ = WD_SECTION  # import gardé pour les extensions futures (sections multiples)

    mode_compatibilite_15(doc)
    for partie in [doc.element.body] + [
        getattr(section, nom)._element
        for nom in ("first_page_header", "first_page_footer", "header", "footer")
    ]:
        ordonner_ooxml(partie)

    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()


# Ordre des enfants imposé par le schéma OOXML (Word l'exige pour appliquer certaines propriétés).
ORDRE = {
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
}


def ordonner_ooxml(racine) -> None:
    from docx.oxml.ns import qn

    for balise, sequence in ORDRE.items():
        rangs = {qn("w:" + nom): i for i, nom in enumerate(sequence.split())}
        for element in racine.iter(qn("w:" + balise)):
            enfants = list(element)
            enfants.sort(key=lambda e: rangs.get(e.tag, len(rangs)))
            for e in enfants:
                element.append(e)


def mode_compatibilite_15(doc) -> None:
    """Mode Word 2013+ : le retrait d'un tableau se mesure à son bord (et non au texte de la cellule)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    reglages = doc.settings.element
    compat = reglages.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        reglages.append(compat)
    for e in compat.findall(qn("w:compatSetting")):
        if e.get(qn("w:name")) == "compatibilityMode":
            compat.remove(e)
    e = OxmlElement("w:compatSetting")
    e.set(qn("w:name"), "compatibilityMode")
    e.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    e.set(qn("w:val"), "15")
    compat.append(e)


def lien(par, url, libelle, police, couleur):
    """Ajoute un lien hypertexte cliquable (python-docx n'a pas d'API dédiée)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    rid = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = par.add_run(libelle)
    police(run, 9.75, gras=True, couleur=couleur)
    run.font.underline = True
    h.append(run._element)
    par._p.append(h)


# ---------------------------------------------------------------- programme principal


def rendre(chemin_json: str, sortie: str | None = None, formats=("html", "docx")) -> list[Path]:
    p, rapport = normaliser_plan(ch.lire_json(chemin_json))
    rapport.afficher()
    if not rapport.ok:
        raise SystemExit(1)
    dossier = Path(sortie) if sortie else Path(chemin_json).resolve().parent
    dossier.mkdir(parents=True, exist_ok=True)
    base = "plan-" + ch.slug(p["intitule"])
    fichiers = []
    if "html" in formats:
        f = dossier / (base + ".html")
        f.write_text(rendre_html(p), encoding="utf-8")
        fichiers.append(f)
    if "docx" in formats:
        f = dossier / (base + ".docx")
        f.write_bytes(rendre_docx(p))
        fichiers.append(f)
    return fichiers


def main() -> None:
    parser = argparse.ArgumentParser(description="Plan de formation Skazy Formation : JSON → HTML et DOCX.")
    parser.add_argument("json", help="fichier plan.json (voir references/format-plan.md)")
    parser.add_argument("--sortie", help="dossier de sortie (par défaut : celui du JSON)")
    parser.add_argument("--formats", default="html,docx", help="formats à produire : html, docx (séparés par des virgules)")
    parser.add_argument("--verifier", action="store_true", help="contrôle le JSON sans rien produire")
    args = parser.parse_args()
    if args.verifier:
        _, rapport = normaliser_plan(ch.lire_json(args.json))
        rapport.afficher()
        print("Plan valide." if rapport.ok else "Plan à corriger.")
        raise SystemExit(0 if rapport.ok else 1)
    for f in rendre(args.json, args.sortie, [x.strip() for x in args.formats.split(",") if x.strip()]):
        print(f)


if __name__ == "__main__":
    main()
