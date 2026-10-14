#!/usr/bin/env python3
"""Quiz de formation : quiz.json (format Skazy Quiz v1) → JSON Skazy Quiz, lien de partage,
fichier Kahoot (.xlsx), script Google Forms (.gs) et corrigé lisible (.md).

Usage :
    python scripts/rendre_quiz.py quiz.json [--sortie DOSSIER] [--formats skazy,kahoot,forms,corrige] [--verifier]

Le format est celui du projet Quiz de Skazy Formation (https://gharel.github.io/quiz/), décrit dans
references/format-quiz.md. Les règles de validation reprennent celles de son éditeur.
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import random
import secrets
import sys
import time
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import charte as ch  # noqa: E402

URL_QUIZ = "https://gharel.github.io/quiz/"
TYPES = ("single", "multiple", "truefalse", "text", "slider", "order")
LIBELLES = {"single": "QCM", "multiple": "Choix multiples", "truefalse": "Vrai ou faux", "text": "Réponse libre", "slider": "Curseur", "order": "Remise en ordre"}
TEMPS = [5, 10, 15, 20, 30, 45, 60, 90, 120]
TEMPS_DEFAUT = {"single": 20, "multiple": 30, "truefalse": 15, "text": 30, "slider": 30, "order": 40}
TEMPS_KAHOOT = [5, 10, 20, 30, 60, 90, 120, 240]
CATEGORIES_QUIZ = {"bureautique", "dev", "web", "gestion", "graphisme", "ia", "nocode", "webmarketing", "data", "general"}
MAX = {"titre": 120, "description": 300, "question": 240, "reponse": 90, "explication": 300, "unite": 12}
KAHOOT_QUESTION, KAHOOT_REPONSE = 120, 75


def uid(n: int = 8) -> str:
    return secrets.token_hex(n // 2 + 1)[:n]


def plus_proche(valeur: float, options: list[int]) -> int:
    """Valeur autorisée la plus proche ; à égalité, la plus longue (laisse le temps de lire)."""
    return min(options, key=lambda o: (abs(o - valeur), -o))


def categorie_quiz(nom: str, r: ch.Rapport) -> str:
    if (nom or "") in CATEGORIES_QUIZ:
        return nom
    if ch.normaliser(nom or "") in ("general", "culture generale"):
        return "general"
    try:
        return ch.categorie(nom)["quiz"]
    except ValueError:
        r.avertir("catégorie « {} » inconnue du Quiz : « general » (Culture générale) utilisée.".format(nom))
        return "general"


# ---------------------------------------------------------------- contrôle


def normaliser_quiz(brut: dict) -> tuple[dict, ch.Rapport]:
    """Contrôle un quiz au format Skazy Quiz v1 (enveloppe ou objet nu) et complète les valeurs par défaut."""
    r = ch.Rapport()
    q = dict(brut.get("quiz", brut))
    if not isinstance(q.get("title"), str) or not q["title"].strip():
        r.erreur("« title » (titre du quiz) est obligatoire.")
    if not isinstance(q.get("questions"), list) or not q["questions"]:
        r.erreur("« questions » est vide ou absent.")
    if not r.ok:
        return q, r
    if len(q["title"]) > MAX["titre"]:
        r.erreur("titre trop long ({} caractères, {} au plus).".format(len(q["title"]), MAX["titre"]))
    q["description"] = (q.get("description") or "").strip()
    if len(q["description"]) > MAX["description"]:
        r.erreur("description trop longue ({} caractères, {} au plus).".format(len(q["description"]), MAX["description"]))
    q["category"] = categorie_quiz(q.get("category", "general"), r)

    questions = []
    for i, brute in enumerate(q["questions"], 1):
        x = dict(brute)
        t = x.get("type")
        ou = "question {}".format(i)
        if t not in TYPES:
            r.erreur("{} : type « {} » inconnu ({}).".format(ou, t, ", ".join(TYPES)))
            continue
        texte = (x.get("text") or "").strip()
        if not texte:
            r.erreur("{} : intitulé manquant.".format(ou))
        elif len(texte) > MAX["question"]:
            r.erreur("{} : intitulé trop long ({} caractères, {} au plus).".format(ou, len(texte), MAX["question"]))
        x["text"] = texte
        if t in ("single", "multiple"):
            reponses = [str(a).strip() for a in x.get("answers") or []]
            correct = [int(k) for k in x.get("correct") or [] if str(k).lstrip("-").isdigit()]
            if len([a for a in reponses if a]) < 2:
                r.erreur("{} : proposez au moins deux réponses.".format(ou))
            if len(reponses) > 4:
                r.erreur("{} : 4 réponses au plus ({} reçues).".format(ou, len(reponses)))
            for a in reponses:
                if len(a) > MAX["reponse"]:
                    r.erreur("{} : réponse trop longue ({} caractères, {} au plus) : « {} ».".format(ou, len(a), MAX["reponse"], a))
            correct = sorted({k for k in correct if 0 <= k < len(reponses) and reponses[k]})
            if not correct:
                r.erreur("{} : indiquez la bonne réponse (« correct » : indices à partir de 0).".format(ou))
            if t == "single" and len(correct) > 1:
                r.erreur("{} : un QCM n'a qu'une bonne réponse ; utilisez le type « multiple ».".format(ou))
            if t == "multiple" and len(correct) < 2:
                r.erreur("{} : cochez au moins deux bonnes réponses (sinon, type « single »).".format(ou))
            x["answers"], x["correct"] = reponses, correct
        elif t == "truefalse":
            if not isinstance(x.get("truth"), bool):
                r.erreur("{} : « truth » doit valoir true ou false.".format(ou))
        elif t == "text":
            acceptees = [str(a).strip() for a in x.get("accepted") or [] if str(a).strip()]
            if not acceptees:
                r.erreur("{} : indiquez au moins une réponse acceptée.".format(ou))
            if len(acceptees) > 5:
                r.erreur("{} : 5 réponses acceptées au plus.".format(ou))
            x["accepted"] = acceptees
        elif t == "slider":
            s = dict(x.get("slider") or {})
            try:
                mini, maxi, pas, rep = (float(s[k]) for k in ("min", "max", "step", "answer"))
                tol = float(s.get("tolerance", 0))
            except (KeyError, TypeError, ValueError):
                r.erreur("{} : « slider » doit contenir min, max, step, answer (et tolerance).".format(ou))
                continue
            if not maxi > mini:
                r.erreur("{} : la valeur maximale doit dépasser la minimale.".format(ou))
            if not pas > 0:
                r.erreur("{} : le pas doit être supérieur à zéro.".format(ou))
            if not mini <= rep <= maxi:
                r.erreur("{} : la bonne réponse doit être entre le minimum et le maximum.".format(ou))
            if tol < 0:
                r.erreur("{} : la marge d'erreur ne peut pas être négative.".format(ou))
            if len(s.get("unit") or "") > MAX["unite"]:
                r.erreur("{} : unité trop longue ({} caractères au plus).".format(ou, MAX["unite"]))
            x["slider"] = {k: (int(v) if float(v).is_integer() else v) for k, v in (("min", mini), ("max", maxi), ("step", pas), ("answer", rep), ("tolerance", tol))}
            if s.get("unit"):
                x["slider"]["unit"] = s["unit"]
        elif t == "order":
            items = [str(a).strip() for a in x.get("items") or [] if str(a).strip()]
            if not 3 <= len(items) <= 4:
                r.erreur("{} : 3 ou 4 éléments à remettre en ordre ({} reçus).".format(ou, len(items)))
            for a in items:
                if len(a) > MAX["reponse"]:
                    r.erreur("{} : élément trop long ({} caractères au plus) : « {} ».".format(ou, MAX["reponse"], a))
            x["items"] = items
        temps = x.get("time") or TEMPS_DEFAUT[t]
        if temps not in TEMPS:
            nouveau = plus_proche(float(temps), TEMPS)
            r.avertir("{} : temps {} s ramené à {} s (valeurs possibles : {}).".format(ou, temps, nouveau, TEMPS))
            temps = nouveau
        x["time"] = temps
        x["points"] = x.get("points", 1) if x.get("points", 1) in (0, 1, 2) else 1
        if x.get("explanation") and len(x["explanation"]) > MAX["explication"]:
            r.erreur("{} : explication trop longue ({} caractères au plus).".format(ou, MAX["explication"]))
        if not x.get("explanation"):
            x.pop("explanation", None)
        x["id"] = x.get("id") or uid()
        questions.append(x)
    q["questions"] = questions
    controles_pedagogiques(questions, r)
    return q, r


def controles_pedagogiques(questions: list[dict], r: ch.Rapport) -> None:
    """Contrôles de qualité issus du skill qcm-creator (avertissements seulement)."""
    qcm = [q for q in questions if q["type"] == "single" and len(q.get("answers") or []) == 4]
    if len(qcm) >= 4:
        positions = [q["correct"][0] for q in qcm if q.get("correct")]
        plus_frequente = max(set(positions), key=positions.count)
        if positions.count(plus_frequente) > len(positions) / 2:
            r.avertir("la bonne réponse est en position {} dans {} QCM sur {} : variez sa place.".format("ABCD"[plus_frequente], positions.count(plus_frequente), len(positions)))
        plus_longue = sum(1 for q in qcm if q.get("correct") and len(q["answers"][q["correct"][0]]) == max(len(a) for a in q["answers"]) and len(set(map(len, q["answers"]))) > 1)
        if plus_longue > len(qcm) * 0.6:
            r.avertir("la bonne réponse est la plus longue dans {} QCM sur {} : équilibrez la longueur des propositions.".format(plus_longue, len(qcm)))
    vf = [q for q in questions if q["type"] == "truefalse" and isinstance(q.get("truth"), bool)]
    if len(vf) >= 4:
        vrais = sum(1 for q in vf if q["truth"])
        if abs(vrais - (len(vf) - vrais)) > max(1, len(vf) // 3):
            r.avertir("vrai ou faux déséquilibrés : {} vrais pour {} faux.".format(vrais, len(vf) - vrais))
    sans = sum(1 for q in questions if not q.get("explanation"))
    if sans:
        r.avertir("{} question(s) sans explication : ajoutez une justification courte (« explanation »).".format(sans))


# ---------------------------------------------------------------- Skazy Quiz


def export_skazy(q: dict) -> dict:
    maintenant = int(time.time() * 1000)
    quiz = {
        "id": q.get("id") or uid(10),
        "title": q["title"].strip(),
        "description": q["description"],
        "category": q["category"],
        "questions": q["questions"],
        "createdAt": q.get("createdAt") or maintenant,
        "updatedAt": maintenant,
    }
    return {"format": "skazy-quiz", "version": 1, "quiz": quiz}


def lien_partage(q: dict) -> str:
    """Lien qui ouvre le quiz dans l'application Quiz (mêmes données que encodeShare())."""
    questions = [{k: v for k, v in x.items() if not (k == "image" and str(v).startswith("data:"))} for x in q["questions"]]
    donnees = json.dumps({"title": q["title"], "description": q["description"], "category": q["category"], "questions": questions},
                         ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    compresseur = zlib.compressobj(9, zlib.DEFLATED, -15)
    brut = compresseur.compress(donnees) + compresseur.flush()
    code = base64.urlsafe_b64encode(brut).decode().rstrip("=")
    return "{}#/partage?q={}".format(URL_QUIZ, code)


def lire_lien(lien: str) -> dict:
    """Décode un lien de partage (contrôle et tests)."""
    code = lien.split("q=", 1)[1]
    brut = base64.urlsafe_b64decode(code + "=" * (-len(code) % 4))
    return json.loads(zlib.decompress(brut, -15).decode("utf-8"))


# ---------------------------------------------------------------- conversion en choix (Kahoot, Forms)


def en_choix(q: dict, graine: int) -> dict | None:
    """Ramène une question à des propositions + bonnes réponses ; None si impossible (réponse libre)."""
    t = q["type"]
    if t in ("single", "multiple"):
        reponses = [a for a in q["answers"] if a]
        bonnes = [reponses.index(q["answers"][k]) for k in q["correct"]]
        return {"texte": q["text"], "choix": reponses, "bonnes": bonnes, "multiple": t == "multiple", "converti": False}
    if t == "truefalse":
        return {"texte": q["text"], "choix": ["Vrai", "Faux"], "bonnes": [0 if q["truth"] else 1], "multiple": False, "converti": False}
    alea = random.Random(graine)
    if t == "slider":
        s = q["slider"]
        unite = (" " + s["unit"]) if s.get("unit") else ""
        pas = max(float(s["step"]), float(s.get("tolerance", 0)) * 2 + float(s["step"]))
        candidats = []
        for k in (-3, -2, -1, 1, 2, 3):
            v = float(s["answer"]) + k * pas
            if float(s["min"]) <= v <= float(s["max"]):
                candidats.append(v)
        alea.shuffle(candidats)
        valeurs = sorted([float(s["answer"])] + candidats[:3])
        fmt = lambda v: ("{:g}".format(v)).replace(".", ",") + unite  # noqa: E731
        return {"texte": q["text"], "choix": [fmt(v) for v in valeurs], "bonnes": [valeurs.index(float(s["answer"]))], "multiple": False, "converti": True}
    if t == "order":
        items = q["items"]
        bonne = list(items)
        permutations, essais = [bonne], 0
        while len(permutations) < 4 and essais < 50:
            p = list(items)
            alea.shuffle(p)
            if p not in permutations:
                permutations.append(p)
            essais += 1
        alea.shuffle(permutations)
        return {"texte": q["text"] + " (bon ordre ?)", "choix": [" → ".join(p) for p in permutations],
                "bonnes": [permutations.index(bonne)], "multiple": False, "converti": True}
    return None


def raccourcir(texte: str, maxi: int) -> str:
    return texte if len(texte) <= maxi else texte[: maxi - 1].rstrip() + "…"


# ---------------------------------------------------------------- Kahoot


def export_kahoot(q: dict, r: ch.Rapport) -> bytes:
    """Classeur au gabarit d'import Kahoot (en-têtes ligne 8, questions à partir de la ligne 9)."""
    ch.exiger("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "Quiz template"
    ws["B2"].font = Font(bold=True, size=16)
    ws["B3"] = "Add questions, at least two answer alternatives, time limit and choose correct answers (at least one). Have fun creating your awesome quiz!"
    ws["B4"] = "Remember: questions have a limit of 120 characters and answers can have 75 characters max. Text will turn red in Excel or Google Docs if you exceed this limit. If several answers are correct, separate them with a comma."
    ws["B5"] = "Généré par le skill skazy-quiz à partir de « {} ».".format(q["title"])
    ws["B6"] = "And remember,  if you're not using Excel you need to export to .xlsx format before you upload to Kahoot!"
    for cellule in ("B3", "B4", "B6"):
        ws.merge_cells("{}:H{}".format(cellule, cellule[1:]))
    entetes = [
        "Question - max 120 characters",
        "Answer 1 - max 75 characters",
        "Answer 2 - max 75 characters",
        "Answer 3 - max 75 characters",
        "Answer 4 - max 75 characters",
        "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs",
        "Correct answer(s) - choose at least one",
    ]
    for k, texte in enumerate(entetes):
        c = ws.cell(row=8, column=2 + k, value=texte)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="46178F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, largeur in zip("ABCDEFGH", (12.5, 49.3, 31.8, 31.8, 31.8, 31.8, 26.3, 33.2)):
        ws.column_dimensions[col].width = largeur
    validation = DataValidation(type="list", formula1='"5,10,20,30,60,90,120,240"', allow_blank=True)
    ws.add_data_validation(validation)

    ligne = 9
    for i, x in enumerate(q["questions"], 1):
        c = en_choix(x, i)
        if c is None:
            r.avertir("Kahoot : question {} (réponse libre) écartée, Kahoot n'importe que des questions à choix.".format(i))
            continue
        if any(len(a) > KAHOOT_REPONSE for a in c["choix"]) and x["type"] == "order":
            r.avertir("Kahoot : question {} (remise en ordre) écartée, ses propositions dépassent {} caractères une fois converties en QCM.".format(i, KAHOOT_REPONSE))
            continue
        if c["converti"]:
            r.avertir("Kahoot : question {} ({}) convertie en QCM à {} propositions.".format(i, LIBELLES[x["type"]], len(c["choix"])))
        if len(c["texte"]) > KAHOOT_QUESTION:
            r.avertir("Kahoot : question {} raccourcie à {} caractères.".format(i, KAHOOT_QUESTION))
        ws.cell(row=ligne, column=1, value=ligne - 8)
        ws.cell(row=ligne, column=2, value=raccourcir(c["texte"], KAHOOT_QUESTION))
        for k, a in enumerate(c["choix"][:4]):
            if len(a) > KAHOOT_REPONSE:
                r.avertir("Kahoot : question {}, réponse {} raccourcie à {} caractères.".format(i, k + 1, KAHOOT_REPONSE))
            ws.cell(row=ligne, column=3 + k, value=raccourcir(a, KAHOOT_REPONSE))
        ws.cell(row=ligne, column=7, value=plus_proche(x["time"], TEMPS_KAHOOT))
        ws.cell(row=ligne, column=8, value=",".join(str(k + 1) for k in c["bonnes"]))
        validation.add("G{}".format(ligne))
        ligne += 1
    if ligne == 9:
        r.avertir("Kahoot : aucune question exportable.")
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------- Google Forms


def export_forms(q: dict, r: ch.Rapport) -> str:
    """Script Google Apps Script qui crée le quiz dans Google Forms (à coller dans script.google.com)."""
    questions = []
    for i, x in enumerate(q["questions"], 1):
        item = {"texte": x["text"], "points": x["points"], "explication": x.get("explanation", "")}
        if x.get("image", "").startswith("https://"):
            item["image"] = x["image"]
        if x["type"] == "text":
            item.update(type="texte", acceptees=x["accepted"])
            r.avertir("Google Forms : question {} (réponse libre) ajoutée sans correction automatique (réponses acceptées dans le commentaire).".format(i))
        else:
            c = en_choix(x, i)
            if c["converti"]:
                r.avertir("Google Forms : question {} ({}) convertie en QCM.".format(i, LIBELLES[x["type"]]))
            item.update(type="cases" if c["multiple"] else "choix", texte=c["texte"], choix=c["choix"], bonnes=c["bonnes"])
        questions.append(item)
    donnees = json.dumps({"titre": q["title"], "description": q["description"], "questions": questions}, ensure_ascii=False, indent=2)
    return GABARIT_FORMS.replace("__TITRE__", q["title"].replace("*/", "* /")).replace("__DONNEES__", donnees)


GABARIT_FORMS = """/**
 * Quiz « __TITRE__ » pour Google Forms, généré par le skill skazy-quiz.
 *
 * 1. Ouvrez https://script.google.com et créez un projet (Nouveau projet).
 * 2. Remplacez le code par tout ce fichier, puis enregistrez.
 * 3. Choisissez la fonction creerQuiz et cliquez sur Exécuter ; autorisez l'accès à Google Forms.
 * 4. Les liens du formulaire (à partager et à modifier) s'affichent dans le journal d'exécution.
 *
 * Le formulaire est en mode quiz : points, bonnes réponses et commentaires sont déjà réglés.
 * Les réponses libres n'ont pas de correction automatique : les réponses acceptées sont dans le commentaire.
 */
const QUIZ = __DONNEES__;

function creerQuiz() {
  const form = FormApp.create(QUIZ.titre);
  form.setIsQuiz(true);
  form.setDescription(QUIZ.description || '');
  form.setProgressBar(true);
  QUIZ.questions.forEach((q, i) => {
    if (q.image) {
      try {
        form.addImageItem().setImage(UrlFetchApp.fetch(q.image)).setTitle('Question ' + (i + 1));
      } catch (e) {
        Logger.log('Image de la question ' + (i + 1) + ' ignorée : ' + e);
      }
    }
    let item;
    if (q.type === 'texte') {
      item = form.addTextItem();
      const reponses = 'Réponses acceptées : ' + q.acceptees.join(' / ');
      item.setGeneralFeedback(FormApp.createFeedback().setText(q.explication ? q.explication + ' ' + reponses : reponses).build());
    } else {
      item = q.type === 'cases' ? form.addCheckboxItem() : form.addMultipleChoiceItem();
      item.setChoices(q.choix.map((c, k) => item.createChoice(c, q.bonnes.indexOf(k) !== -1)));
      if (q.explication) {
        const retour = FormApp.createFeedback().setText(q.explication).build();
        item.setFeedbackForCorrect(retour);
        item.setFeedbackForIncorrect(retour);
      }
    }
    item.setTitle(q.texte).setRequired(true).setPoints(q.points);
  });
  Logger.log('Formulaire à partager : ' + form.getPublishedUrl());
  Logger.log('Formulaire à modifier : ' + form.getEditUrl());
}
"""


# ---------------------------------------------------------------- corrigé


def export_corrige(q: dict) -> str:
    lignes = ["# {}".format(q["title"]), ""]
    if q["description"]:
        lignes += [q["description"], ""]
    lignes += ["{} questions · catégorie {}".format(len(q["questions"]), q["category"]), ""]
    for i, x in enumerate(q["questions"], 1):
        lignes.append("## Question {} · {}".format(i, LIBELLES[x["type"]]))
        lignes.append("")
        lignes.append("**{}**".format(x["text"]))
        lignes.append("")
        t = x["type"]
        if t in ("single", "multiple"):
            for k, a in enumerate(x["answers"]):
                lignes.append("- {}. {}{}".format("ABCD"[k], a, " ✔" if k in x["correct"] else ""))
        elif t == "truefalse":
            lignes.append("- Réponse : **{}**".format("Vrai" if x["truth"] else "Faux"))
        elif t == "text":
            lignes.append("- Réponses acceptées : {}".format(" / ".join(x["accepted"])))
        elif t == "slider":
            s = x["slider"]
            lignes.append("- Réponse : **{}{}** (marge ± {}, curseur de {} à {})".format(s["answer"], " " + s["unit"] if s.get("unit") else "", s["tolerance"], s["min"], s["max"]))
        elif t == "order":
            lignes.append("- Bon ordre : {}".format(" → ".join(x["items"])))
        if x.get("explanation"):
            lignes += ["", "> {}".format(x["explanation"])]
        lignes.append("")
    return "\n".join(lignes)


# ---------------------------------------------------------------- programme principal


def rendre(chemin_json: str, sortie: str | None = None, formats=("skazy", "kahoot", "forms", "corrige")) -> tuple[list[Path], str]:
    q, rapport = normaliser_quiz(ch.lire_json(chemin_json))
    if not rapport.ok:
        rapport.afficher()
        raise SystemExit(1)
    dossier = Path(sortie) if sortie else Path(chemin_json).resolve().parent
    dossier.mkdir(parents=True, exist_ok=True)
    base = "quiz-" + ch.slug(q["title"])
    fichiers = []
    lien = lien_partage(q)
    if "skazy" in formats:
        f = dossier / (base + ".skazy-quiz.json")
        f.write_text(json.dumps(export_skazy(q), ensure_ascii=False, indent=2), encoding="utf-8")
        fichiers.append(f)
        f = dossier / (base + ".lien.txt")
        f.write_text(lien + "\n", encoding="utf-8")
        fichiers.append(f)
    if "kahoot" in formats:
        f = dossier / (base + ".kahoot.xlsx")
        f.write_bytes(export_kahoot(q, rapport))
        fichiers.append(f)
    if "forms" in formats:
        f = dossier / (base + ".google-forms.gs")
        f.write_text(export_forms(q, rapport), encoding="utf-8")
        fichiers.append(f)
    if "corrige" in formats:
        f = dossier / (base + ".corrige.md")
        f.write_text(export_corrige(q), encoding="utf-8")
        fichiers.append(f)
    rapport.afficher()
    return fichiers, lien


def main() -> None:
    parser = argparse.ArgumentParser(description="Quiz Skazy Formation : JSON → Skazy Quiz, Kahoot, Google Forms, corrigé.")
    parser.add_argument("json", help="fichier quiz.json au format Skazy Quiz v1 (voir references/format-quiz.md)")
    parser.add_argument("--sortie", help="dossier de sortie (par défaut : celui du JSON)")
    parser.add_argument("--formats", default="skazy,kahoot,forms,corrige", help="skazy, kahoot, forms, corrige (séparés par des virgules)")
    parser.add_argument("--verifier", action="store_true", help="contrôle le JSON sans rien produire")
    args = parser.parse_args()
    if args.verifier:
        q, rapport = normaliser_quiz(ch.lire_json(args.json))
        rapport.afficher()
        if rapport.ok:
            print("Quiz valide : {} questions.".format(len(q["questions"])))
        raise SystemExit(0 if rapport.ok else 1)
    fichiers, lien = rendre(args.json, args.sortie, [x.strip() for x in args.formats.split(",") if x.strip()])
    for f in fichiers:
        print(f)
    print("Lien du quiz : " + lien)


if __name__ == "__main__":
    main()
