import io
import json

import pytest
from conftest import SKILLS

import charte as ch
import support_modele as m

EXEMPLE = SKILLS / "skazy-support-formation" / "exemples" / "support-exemple.json"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
PX = 6350


def support(**modifs):
    donnees = json.loads(EXEMPLE.read_text(encoding="utf-8"))
    donnees.update(modifs)
    return m.normaliser_support(donnees)


def presentation(s):
    from pptx import Presentation
    from support_pptx import rendre_pptx

    return Presentation(io.BytesIO(rendre_pptx(s)))


@pytest.fixture(scope="module")
def anime():
    """Le support d'exemple rendu une seule fois : couples (slide du JSON, slide du PPTX)."""
    ch.IMAGES_LOCALES.clear()
    s, _ = support()
    return list(zip(s["_slides"], presentation(s).slides))


def clics(slide) -> list[list[int]]:
    """Identifiants des formes de chaque clic, lus dans le XML p:timing."""
    timing = slide._element.find(P + "timing")
    if timing is None:
        return []
    seq = timing.find(".//{}cTn[@nodeType='mainSeq']".format(P))
    return [list(dict.fromkeys(int(t.get("spid")) for t in clic.iter(P + "spTgt"))) for clic in seq.find(P + "childTnLst")]


def a_du_contenu(sh) -> bool:
    return sh.shape_type == 13 or (sh.has_text_frame and sh.text_frame.text.strip() != "")  # 13 : image


def dedans(a, b, marge=2 * PX) -> bool:
    """La forme a tient-elle dans le cadre de b ?"""
    return (
        a.left >= b.left - marge
        and a.top >= b.top - marge
        and a.left + a.width <= b.left + b.width + marge
        and a.top + a.height <= b.top + b.height + marge
    )


def test_exemple_couvre_les_30_mises_en_page():
    s, rapport = support()
    assert rapport.erreurs == []
    assert {x["type"] for x in s["_slides"]} == set(m.TYPES)


def test_quiz_en_deux_slides_et_numerotation():
    s, _ = support()
    quiz = [x for x in s["_slides"] if x["type"] == "quiz"]
    assert len(quiz) == 4
    assert [q["_reponse"] for q in quiz] == [False, True, False, True]
    assert (quiz[0]["numero"], quiz[0]["total"]) == (1, 2)
    assert [x["numero"] for x in s["_slides"] if x["type"] == "chapitre"] == [1, 2, 3]


def test_erreurs():
    slides = [{"type": "inconnu"}, {"type": "quiz", "question": "?", "propositions": ["a", "b"], "bonne": 5}]
    _, rapport = support(slides=slides)
    textes = " ".join(rapport.erreurs)
    assert "type inconnu" in textes
    assert "bonne" in textes


def test_alias_et_lettre_de_reponse():
    slides = [
        {"type": "couverture"},
        {"type": "Chiffre clé", "chiffre": "3", "phrase": "x"},
        {"type": "quiz", "question": "?", "propositions": ["a", "b"], "bonne": "B"},
    ]
    s, rapport = support(slides=slides)
    assert rapport.erreurs == []
    assert s["_slides"][1]["type"] == "chiffre-cle"
    assert s["_slides"][2]["bonne"] == 1


def test_html():
    from support_html import rendre_html

    s, _ = support()
    html = rendre_html(s)
    assert html.count('<section class="slide"') == len(s["_slides"])
    assert 'template class="notes-slide"' in html
    assert "--c:#A977DA" in html


def test_pptx_slides_notes_et_pied():
    from pptx import Presentation
    from support_pptx import rendre_pptx

    ch.IMAGES_LOCALES.clear()
    s, _ = support()
    prs = Presentation(io.BytesIO(rendre_pptx(s)))
    assert len(prs.slides) == len(s["_slides"])
    assert prs.slide_width == 12192000
    notes = [sl.notes_slide.notes_text_frame.text for sl in prs.slides if sl.has_notes_slide]
    assert any("Accueil des participants" in n for n in notes)
    textes = " ".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)
    assert "Excel : les bases" in textes.replace(" ", " ")
    assert "02" in textes


def test_animations_seulement_sur_les_slides_de_contenu(anime):
    from support_pptx import animee

    for donnees, slide in anime:
        if animee(donnees):
            assert clics(slide), donnees["type"]
        else:
            assert slide._element.find(P + "timing") is None, donnees["type"]
    types = {d["type"] for d, _ in anime if not animee(d)}
    assert types == {"couverture", "chapitre", "image", "questions", "quiz"}


def test_fondu_au_clic(anime):
    for _, slide in anime:
        for effet in slide._element.iter(P + "cTn"):
            if effet.get("presetClass") is None:
                continue
            assert (effet.get("presetClass"), effet.get("presetID")) == ("entr", "10")
            fondu = effet.find(".//{}animEffect".format(P))
            assert fondu.get("filter") == "fade" and fondu.get("transition") == "in"
            assert fondu.find(".//{}cTn".format(P)).get("dur") == "400"
        timing = slide._element.find(P + "timing")
        if timing is None:
            continue
        seq = timing.find(".//{}cTn[@nodeType='mainSeq']".format(P))
        for clic in seq.find(P + "childTnLst"):
            noeuds = [c.get("nodeType") for c in clic.iter(P + "cTn") if c.get("presetClass")]
            assert noeuds[0] == "clickEffect" and set(noeuds[1:]) <= {"withEffect"}


def test_titre_et_pied_restent_fixes(anime):
    for donnees, slide in anime:
        animes = {i for clic in clics(slide) for i in clic}
        for sh in slide.shapes:
            if sh.shape_id in animes:
                assert sh.top < 1000 * PX, donnees["type"]
        titre = ch.sans_gras(donnees.get("titre") or "").replace(" ", " ").replace(" ", " ")
        if titre:
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.replace(" ", " ").replace(" ", " ") == titre:
                    assert sh.shape_id not in animes, donnees["type"]


def test_chaque_clic_montre_du_contenu(anime):
    for donnees, slide in anime:
        formes = {sh.shape_id: sh for sh in slide.shapes}
        for clic in clics(slide):
            assert any(a_du_contenu(formes[i]) for i in clic), donnees["type"]


def test_aucun_conteneur_avant_son_contenu(anime):
    for donnees, slide in anime:
        numero = {i: k for k, clic in enumerate(clics(slide), 1) for i in clic}
        formes = list(slide.shapes)
        for r in formes:
            vide = r.shape_type == 1 and not r.text_frame.text.strip()  # 1 : forme automatique
            if not vide or min(r.width, r.height) < 40 * PX:
                continue
            contenu = [f for f in formes if f is not r and dedans(f, r)]
            if contenu:
                attendu = min(numero.get(f.shape_id, 0) for f in contenu)
                assert numero.get(r.shape_id, 0) == attendu, "{} : conteneur au clic {}, contenu au clic {}".format(
                    donnees["type"], numero.get(r.shape_id, 0), attendu
                )


def test_sans_animations():
    donnees = json.loads(EXEMPLE.read_text(encoding="utf-8"))
    comparatif = next(x for x in donnees["slides"] if x["type"] == "comparatif")

    def cadres(animations):
        s, _ = m.normaliser_support(dict(donnees, formation=dict(donnees["formation"], animations=animations), slides=[comparatif]))
        slide = presentation(s).slides[0]
        # Cadre de 1 680 px, trait de 2 px posé à l'intérieur (forme rentrée d'un demi-trait de chaque côté).
        return slide, [sh for sh in slide.shapes if abs(sh.width - 1680 * PX) <= 2 * PX and sh.height > 100 * PX]

    slide, cadre = cadres(False)
    assert slide._element.find(P + "timing") is None
    assert len(cadre) == 1
    # Animé, le tableau garde son cadre (comme le modèle), fixe : seules les lignes apparaissent au clic.
    slide, cadre = cadres(True)
    animes = {i for clic in clics(slide) for i in clic}
    assert clics(slide) and len(cadre) == 1 and cadre[0].shape_id not in animes


def test_pptx_allege():
    """Paquet allégé : une seule disposition, ni vignette ni propriétés, images transparentes en palette."""
    import zipfile

    from PIL import Image
    from support_pptx import rendre_pptx

    ch.IMAGES_LOCALES.clear()
    s, _ = support()
    donnees = rendre_pptx(s)
    z = zipfile.ZipFile(io.BytesIO(donnees))
    noms = z.namelist()
    assert [n for n in noms if n.startswith("ppt/slideLayouts/slideLayout")] == ["ppt/slideLayouts/slideLayout7.xml"]
    assert not any(n.startswith(("docProps/", "ppt/printerSettings/")) for n in noms)
    maitre = z.read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
    assert maitre.count("<p:sldLayoutId ") == 1
    for n in noms:
        if n.startswith("ppt/media/") and n.endswith(".png"):
            im = Image.open(io.BytesIO(z.read(n)))
            if "transparency" in im.info or im.mode in ("RGBA", "LA"):
                assert im.mode == "P", n
    assert len(presentation(s).slides) == len(s["_slides"])
