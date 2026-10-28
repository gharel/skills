"""Support de formation : rendu PPTX (PowerPoint, Google Slides, Keynote).

Reprend les mesures du modèle HTML 1920 × 1080 du design system (1 px = 6 350 EMU = 0,5 pt) : chaque
mise en page est calculée comme le navigateur la pose (boîtes de ligne, écarts, centrages). Les éléments
que Google Slides importe mal (rotations groupées, ombres, dégradés) sont composés en images avec Pillow.
Chaque slide reçoit les notes formateur du JSON. Les slides de contenu se dévoilent clic par clic :
chaque bloc (carte, étape, ligne de tableau…) apparaît en fondu, titre et pied restant fixes.
"""

from __future__ import annotations

import io
import math
import re
import unicodedata
from contextlib import contextmanager
from functools import lru_cache

import charte as ch
from support_modele import touches

ch.exiger("pptx", "PIL")

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml import parse_xml  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

PX = 6350
W, H = 1920, 1080
BAS = 1000  # haut du pied de slide
POLICE = ch.CHARTE["police"]
MONO = "Consolas"
C = ch.COULEURS
GAUCHE, CENTRE, DROITE = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
HAUT, MILIEU = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE

# Graisse CSS → (famille, gras) PowerPoint : 500, 600 et 800 sont des familles à part, installées avec Raleway.
GRAISSES = {
    400: (POLICE, False),
    500: (POLICE + " Medium", False),
    600: (POLICE + " SemiBold", False),
    700: (POLICE, True),
    800: (POLICE + " ExtraBold", False),
}


# PowerPoint ne crène pas Raleway (l'attribut kern reste sans effet) : ses lignes sont de 0,5 à 1,3 % plus
# larges que dans le navigateur. Les estimations de lignes acceptent 1,05 % de dépassement, et les zones de
# texte sont élargies d'autant (plus 0,2 % d'arrondi) : PowerPoint coupe aux mots prévus par la mise en page,
# en général les mêmes que le navigateur.
ELARGI_ESTIME = 1.0105
ELARGI_BOITE = 1.0125


def e(px: float) -> Emu:
    return Emu(int(round(px * PX)))


def couleur(h: str) -> RGBColor:
    return RGBColor(*ch.rgb(h))


def blanc(cat: dict) -> bool:
    return cat["surCouleur"].upper() == "#FFFFFF"


# ---------------------------------------------------------------- mesure du texte

# Chasse de Raleway (millièmes d'em, sans crénage, comme PowerPoint) pour les caractères ASCII 32 à 126,
# puis « » ’ … – — · € ° œ ; mesurée sur les fichiers de la police, une table par graisse.
CARACTERES = "".join(chr(c) for c in range(32, 127)) + "«»’…–—·€°œ"
_CHASSES = {
    400: (
        "255 252 292 680 619 714 667 197 269 268 321 415 225 422 189 604 612 428 521 534 546 545 603 535 "
        "593 586 188 205 508 438 508 476 837 676 666 686 715 609 589 716 738 248 481 643 583 877 772 751 "
        "622 749 653 607 610 755 679 1046 632 651 629 258 568 258 539 511 222 545 617 549 622 587 330 618 "
        "583 218 225 530 272 919 583 595 617 617 351 483 335 594 533 813 504 542 492 259 231 259 511 573 "
        "573 190 457 570 901 189 801 266 1034"
    ),
    500: (
        "253 262 324 692 622 733 681 210 283 283 329 416 211 423 202 633 614 450 535 540 558 548 606 535 "
        "598 589 201 221 510 439 510 479 840 677 671 682 716 605 583 714 741 260 485 652 584 874 770 748 "
        "623 746 657 609 614 755 680 1049 637 655 631 270 590 270 553 503 232 554 623 554 629 591 343 624 "
        "590 230 237 539 286 924 590 599 623 623 363 488 352 601 539 821 513 545 493 271 240 271 532 596 "
        "596 200 503 582 911 202 795 272 1028"
    ),
    600: (
        "247 284 357 706 628 758 696 222 296 296 337 419 221 421 213 664 614 471 549 549 564 551 607 548 "
        "601 589 223 239 512 438 512 482 843 675 675 684 716 603 579 719 743 272 488 660 584 870 767 752 "
        "623 750 661 611 616 752 678 1051 641 657 629 292 612 292 569 490 243 564 628 559 637 595 356 629 "
        "597 242 248 548 301 927 597 602 628 628 374 491 367 609 544 830 522 547 493 293 256 293 549 621 "
        "621 219 550 592 919 213 797 278 1020"
    ),
    700: (
        "240 307 393 721 634 786 712 235 310 309 345 421 232 418 225 698 614 494 564 558 571 554 607 561 "
        "604 589 247 259 515 436 515 486 845 672 680 686 716 601 574 725 745 286 491 669 584 865 763 756 "
        "623 755 665 613 619 750 676 1053 646 660 627 316 637 316 585 475 255 574 634 564 645 599 370 634 "
        "604 255 261 558 317 931 604 606 634 634 387 494 384 618 550 839 531 550 493 317 272 317 567 649 "
        "649 239 602 603 928 225 798 285 1011"
    ),
    800: (
        "233 333 431 738 641 816 729 250 325 324 354 424 243 416 239 735 614 518 580 569 578 558 608 576 "
        "607 589 274 280 518 435 518 490 848 670 684 688 716 599 569 732 747 300 494 679 585 860 759 760 "
        "623 759 670 615 621 747 674 1055 651 663 625 341 663 341 604 459 269 586 640 570 654 604 385 640 "
        "612 270 274 569 335 935 612 610 640 640 400 497 402 627 556 850 542 552 494 344 291 344 587 679 "
        "679 261 658 615 937 239 800 292 1002"
    ),
}
CHASSES = {g: dict(zip(CARACTERES, (int(v) / 1000 for v in t.split()))) for g, t in _CHASSES.items()}


def chasse(car: str, graisse=400, mono=False) -> float:
    """Avance d'un caractère, en em."""
    if mono:
        return 0.55
    table = CHASSES[min(CHASSES, key=lambda g: abs(g - graisse))]
    if car in table:
        return table[car]
    if car == "\u00a0":
        return table[" "]
    if car == " ":
        return 0.2
    return table.get(unicodedata.normalize("NFD", car)[:1], 0.56)


def caracteres(texte: str, graisse=400, majuscules=False) -> list[tuple[str, int]]:
    """Caractères du texte (typographie appliquée, **gras** en 700) avec leur graisse."""
    sortie = []
    for segment, gras in ch.segments_gras(texte or ""):
        g = max(graisse, 700) if gras else graisse
        sortie += [(x, g) for x in (segment.upper() if majuscules else segment)]
    return sortie


def largeur(texte: str, taille: float, graisse=400, mono=False, espacement=0.0, majuscules=False) -> float:
    """Largeur d'un texte sur une ligne, en px (espacement des lettres compris)."""
    return sum(chasse(x, g, mono) + espacement for x, g in caracteres(texte, graisse, majuscules)) * taille


def taille_ajustee(texte: str, largeur_dispo: float, taille: float, graisse=400, espacement=0.0) -> float:
    """Taille du modèle, réduite juste assez pour que le mot le plus long tienne dans la largeur : un mot
    n'est jamais coupé (« communicatio / n »). Le contenu d'exemple du modèle garde la taille du modèle."""
    mots = [m for m in re.split(r"[ \n]+", ch.sans_gras(ch.typographie(texte or ""))) if m]
    plus_long = max((largeur(m, taille, graisse, espacement=espacement) for m in mots), default=0)
    return taille if plus_long <= largeur_dispo * 0.98 else math.floor(taille * largeur_dispo * 0.98 / plus_long)


def lignes_coupees(texte: str, largeur_dispo: float, taille: float, graisse=400, mono=False, espacement=0.0,
                   majuscules=False) -> list[str]:
    """Coupe le texte comme un navigateur, aux espaces ordinaires (pas aux insécables) ; renvoie les lignes
    sans les marques de gras (dont la largeur est comptée en 700)."""
    sortie = []
    for paragraphe in str(texte or "").split("\n"):
        mots, mot = [], []
        for x, g in caracteres(paragraphe, graisse, majuscules):
            if x == " ":
                mots.append(mot)
                mot = []
            else:
                mot.append((x, g))
        mots.append(mot)
        lignes, courant, l_courant = [], [], 0.0
        for mot in mots:
            l_mot = sum(chasse(x, g, mono) + espacement for x, g in mot) * taille
            l_espace = (chasse(" ", graisse, mono) + espacement) * taille
            if courant and l_courant + l_espace + l_mot > largeur_dispo * ELARGI_ESTIME:
                lignes.append(courant)
                courant, l_courant = [mot], l_mot
            else:
                l_courant += (l_espace if courant else 0) + l_mot
                courant.append(mot)
        lignes.append(courant)
        sortie += [" ".join("".join(x for x, _ in m) for m in ligne) for ligne in lignes]
    return sortie


def nb_lignes(texte: str, largeur_dispo: float, taille: float, graisse=400, mono=False, espacement=0.0) -> int:
    if not texte:
        return 0
    if mono:  # code : pas de retour automatique (white-space: pre)
        return len(str(texte).split("\n"))
    return len(lignes_coupees(texte, largeur_dispo, taille, graisse, mono, espacement))


def hauteur(texte: str, largeur_dispo: float, taille: float, interligne=1.5, graisse=400, mono=False, espacement=0.0) -> float:
    return nb_lignes(texte, largeur_dispo, taille, graisse, mono, espacement) * taille * interligne


def equilibrer(texte: str, largeur_dispo: float, taille: float, graisse=400, espacement=0.0) -> str:
    """text-wrap: balance : mêmes nombres de lignes, mais longueurs égalisées (retours à la ligne explicites)."""
    if not texte or "**" in texte or "\n" in texte:
        return texte
    n = nb_lignes(texte, largeur_dispo, taille, graisse, espacement=espacement)
    if n < 2:
        return texte
    bas, haut = 0.0, largeur_dispo
    for _ in range(24):
        milieu = (bas + haut) / 2
        if nb_lignes(texte, milieu, taille, graisse, espacement=espacement) > n:
            bas = milieu
        else:
            haut = milieu
    return "\n".join(lignes_coupees(texte, haut + 1, taille, graisse, espacement=espacement))


# Ligne de base de la première ligne sous le haut d'une zone de texte PowerPoint, en em, selon l'interligne
# CSS (interligne PowerPoint = interligne CSS / 1,2). Mesuré dans PowerPoint 365, identique pour toutes
# les polices : PowerPoint retire sous 1,2 tout l'écart au-dessus de la ligne et n'en ajoute que 60 % au-delà,
# là où le navigateur partage l'écart moitié au-dessus, moitié au-dessous.
BASE_PPT = [(0.85, 0.63), (0.9, 0.67), (1.0, 0.75), (1.05, 0.81), (1.1, 0.85), (1.2, 0.95)]


def base_ppt(interligne: float) -> float:
    if interligne >= 1.2:
        return 0.95 + 0.6 * (interligne - 1.2)
    for (l0, b0), (l1, b1) in zip(BASE_PPT, BASE_PPT[1:]):
        if interligne <= l1:
            return b0 + (b1 - b0) * (interligne - l0) / (l1 - l0)
    return 0.95


def decalage(interligne: float, ancre=HAUT, mono=False) -> float:
    """Écart (en em) à ajouter au haut d'une zone de texte pour que ses lignes tombent où le navigateur les pose.

    Ligne de base CSS : demi-interligne + 0,84 em pour Raleway, 0,8 em pour Consolas (mesurés dans Chrome).
    Texte centré verticalement : PowerPoint le descend de 0,3 em par unité d'interligne au-delà de 1,2
    (et le remonte de 0,5 em en deçà) ; à 1,2, il est centré exactement comme dans le navigateur.
    """
    if ancre == MILIEU:
        return -(0.3 if interligne >= 1.2 else 0.5) * (interligne - 1.2)
    return (0.8 if mono else 0.84) + (interligne - 1) / 2 - base_ppt(interligne)


def ligne_de_base(taille: float, interligne: float) -> float:
    """Ligne de base CSS d'une ligne de Raleway sous le haut de sa boîte de ligne (alignement baseline)."""
    return taille * (0.84 + (interligne - 1) / 2)


# ---------------------------------------------------------------- formes de base


class Slide:
    """Enveloppe d'une slide python-pptx avec des primitives en pixels du modèle."""

    def __init__(self, deck: "Deck", fond: str = "#FFFFFF") -> None:
        self.deck = deck
        self.s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
        self.blocs: list[list[int]] = []
        if fond.upper() != "#FFFFFF":
            self.rect(0, 0, W, H, fond)

    # -- animations
    @contextmanager
    def bloc(self, *avec):
        """Les formes dessinées dans le `with` apparaissent ensemble, au même clic.

        `avec` : formes déjà dessinées (panneau, axe) qui apparaissent avec ce premier contenu,
        sans changer leur ordre d'empilement.
        """
        avant = {sh.shape_id for sh in self.s.shapes}
        yield
        ids = [sh.shape_id for sh in avec if sh is not None]
        ids += [sh.shape_id for sh in self.s.shapes if sh.shape_id not in avant and sh.shape_id not in ids]
        if ids:
            self.blocs.append(ids)

    def animer(self):
        """Écrit les animations des blocs (python-pptx n'a pas d'API pour cela)."""
        if not self.blocs:
            return
        formes = {sh.shape_id for sh in self.s.shapes if sh._element.tag == qn("p:sp")}
        sld = self.s._element
        timing = parse_xml(xml_animations(self.blocs, formes))
        fin = sld.find(qn("p:extLst"))
        if fin is not None:
            fin.addprevious(timing)
        else:
            sld.append(timing)

    # -- formes
    def rect(self, x, y, w, h, fond=None, rayon=0, trait=None, epaisseur=2, tirets=False, alpha=None, coins="tous"):
        """Rectangle, arrondi si `rayon` ; `coins` : « tous », « haut » (les deux du haut) ou un seul coin
        (« hg », « hd », « bg », « bd »). Le contour est posé à l'intérieur de la boîte, comme en CSS."""
        if trait:  # PowerPoint centre le trait sur le contour : on le rentre d'une demi-épaisseur
            x, y, w, h = x + epaisseur / 2, y + epaisseur / 2, w - epaisseur, h - epaisseur
            rayon = max(0, rayon - epaisseur / 2) if rayon else 0
        if not rayon:
            forme = MSO_SHAPE.RECTANGLE
        elif coins == "tous":
            forme = MSO_SHAPE.ROUNDED_RECTANGLE
        elif coins == "haut":
            forme = MSO_SHAPE.ROUND_2_SAME_RECTANGLE
        else:
            forme = MSO_SHAPE.ROUND_1_RECTANGLE  # coin haut droit, retourné pour les autres
        sh = self.s.shapes.add_shape(forme, e(x), e(y), e(w), e(h))
        if rayon:
            sh.adjustments[0] = min(0.5, rayon / max(1, min(w, h)))
            if coins in ("hg", "bg", "bd"):
                xfrm = sh._element.spPr.find(qn("a:xfrm"))
                if coins in ("hg", "bg"):
                    xfrm.set("flipH", "1")
                if coins in ("bg", "bd"):
                    xfrm.set("flipV", "1")
        if fond:
            sh.fill.solid()
            sh.fill.fore_color.rgb = couleur(fond)
            if alpha is not None:
                srgb = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))
        else:
            sh.fill.background()
        if trait:
            sh.line.color.rgb = couleur(trait)
            sh.line.width = e(epaisseur)
            if tirets:
                tirets_css(sh.line)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        sh.text_frame.text = ""
        return sh

    def rond(self, x, y, d, fond=None, trait=None, epaisseur=2):
        if trait:  # contour à l'intérieur du diamètre, comme en CSS
            x, y, d = x + epaisseur / 2, y + epaisseur / 2, d - epaisseur
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, e(x), e(y), e(d), e(d))
        if fond:
            sh.fill.solid()
            sh.fill.fore_color.rgb = couleur(fond)
        else:
            sh.fill.background()
        if trait:
            sh.line.color.rgb = couleur(trait)
            sh.line.width = e(epaisseur)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def polygone(self, points, fond, alpha=None):
        constructeur = self.s.shapes.build_freeform(e(points[0][0]), e(points[0][1]), scale=1.0)
        constructeur.add_line_segments([(e(x), e(y)) for x, y in points[1:]], close=True)
        sh = constructeur.convert_to_shape()
        sh.fill.solid()
        sh.fill.fore_color.rgb = couleur(fond)
        if alpha is not None:
            srgb = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def ligne(self, x, y, w, h, fond):
        return self.rect(x, y, w, max(h, 1), fond)

    def pointilles(self, x, y, w, coul, epaisseur=3):
        """Filet horizontal en tirets (bordure « dashed » CSS), de y à y + epaisseur."""
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, e(x), e(y + epaisseur / 2), e(x + w), e(y + epaisseur / 2))
        style = ln._element.find(qn("p:style"))  # style du thème : il ajoute une ombre au trait
        if style is not None:
            ln._element.remove(style)
        ln.line.color.rgb = couleur(coul)
        ln.line.width = e(epaisseur)
        tirets_css(ln.line)
        return ln

    def ombre(self, x, y, w, h, rayon, dy, flou, alpha, rond=False):
        """Ombre portée CSS (box-shadow: 0 dy flou rgba(0, 0, 0, alpha)) composée en image, sous la forme."""
        k = 0.5  # demi-résolution : l'ombre est floue
        marge = flou * 1.5 + 2
        toile = Image.new("RGBA", (int((w + 2 * marge) * k), int((h + 2 * marge) * k)), (0, 0, 0, 0))
        boite = (marge * k, marge * k, (marge + w) * k, (marge + h) * k)
        d = ImageDraw.Draw(toile)
        if rond:
            d.ellipse(boite, fill=(0, 0, 0, round(255 * alpha)))
        else:
            d.rounded_rectangle(boite, radius=rayon * k, fill=(0, 0, 0, round(255 * alpha)))
        toile = toile.filter(ImageFilter.GaussianBlur(flou / 2 * k))
        return self.image(ch.png(toile), x - marge, y + dy - marge, w + 2 * marge, h + 2 * marge)

    # -- texte
    def texte(self, x, y, w, h, contenu, taille, coul=None, gras=False, italique=False, align=GAUCHE, ancre=HAUT,
              interligne=1.5, gras_coul=None, majuscules=False, espacement=0.0, police=None, lien=None, forme=None,
              graisse=None, souligne=False, gras_souligne=False, retour=True, ligne_raleway=False):
        """Zone de texte posée comme une boîte de ligne CSS : (x, y) est le haut de la première ligne,
        `interligne` le line-height CSS, `graisse` le font-weight (700 si `gras`).

        `contenu` : texte (avec **gras** et sauts de ligne), ou liste de morceaux (texte, options) pour une ligne
        à plusieurs styles. `ligne_raleway` : la ligne porte un morceau en police mono, mais sa ligne de base
        est celle de Raleway (code dans une phrase)."""
        coul = coul or C["texte"]
        graisse = graisse or (700 if gras else 400)
        mono = police == MONO and not ligne_raleway
        if interligne < 0.85:  # une seule ligne : même ligne de base qu'avec un interligne de 1
            y += (interligne - 1) / 2 * taille
            interligne = 1.0
        if forme is None:
            y += decalage(interligne, ancre, mono) * taille
            if retour:  # compense l'absence de crénage (voir ELARGI_BOITE), sans déplacer le texte
                extra = w * (ELARGI_BOITE - 1)
                x -= extra / 2 if align == CENTRE else (extra if align == DROITE else 0)
                w += extra
            cible = self.s.shapes.add_textbox(e(x), e(y), e(w), e(h))
        else:
            cible = forme
        tf = cible.text_frame
        tf.word_wrap = retour
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ancre
        if isinstance(contenu, list):
            paragraphes = [[(ch.typographie(t), dict({"coul": coul, "graisse": graisse, "police": police}, **o)) for t, o in contenu]]
        else:
            paragraphes = []
            for ligne in str(contenu if contenu is not None else "").split("\n"):
                morceaux = []
                for segment, est_gras in ch.segments_gras(ligne) or [("", False)]:
                    opts = {"coul": coul, "graisse": graisse, "police": police, "souligne": souligne}
                    if est_gras:
                        opts.update(graisse=max(graisse, 700), coul=gras_coul or coul, souligne=souligne or gras_souligne)
                    morceaux.append((segment, opts))
                paragraphes.append(morceaux)
        for k, morceaux in enumerate(paragraphes):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = round(interligne / 1.2, 4)  # interligne simple de PowerPoint : 1,2 em
            for segment, opts in morceaux:
                r = p.add_run()
                r.text = segment
                f = r.font
                g = opts.get("graisse") or 400
                if opts.get("police"):
                    f.name, f.bold = opts["police"], g >= 600
                else:
                    f.name, f.bold = GRAISSES[min(GRAISSES, key=lambda x: abs(x - g))]
                f.size = Pt(taille / 2)
                f.italic = italique
                f.color.rgb = couleur(opts.get("coul") or coul)
                if opts.get("souligne"):
                    f.underline = True
                rpr = r._r.get_or_add_rPr()
                if majuscules:
                    rpr.set("cap", "all")
                if espacement:
                    rpr.set("spc", str(int(round(espacement * taille / 2 * 100))))
        if lien:
            # Lien porté par la forme : la couleur du texte reste celle de la charte.
            cible.click_action.hyperlink.address = lien
        return cible

    def pile(self, x, y, w, elements, ecart=0):
        """Empile des textes : [(contenu, taille, options)] ; renvoie le y de fin."""
        for contenu, taille, opts in elements:
            if not contenu:
                continue
            h = hauteur_opts(contenu, w, taille, opts)
            self.texte(x, y, w, h + 4, contenu, taille, **opts)
            y += h + ecart
        return y - ecart

    def pastille(self, x, y, texte, taille=24, fond=None, coul=None, pad_x=28, pad_y=10, majuscules=True, droite=False,
                 graisse=800, espacement=None):
        """Pastille du modèle : hauteur = taille × 1,5 + 2 × pad_y, largeur au texte."""
        fond = fond or self.deck.c["couleur"]
        coul = coul or self.deck.c["surCouleur"]
        esp = (0.12 if majuscules else 0.0) if espacement is None else espacement
        w = largeur(texte, taille, graisse, espacement=esp, majuscules=majuscules) + 2 * pad_x
        h = taille * 1.5 + 2 * pad_y
        if droite:
            x -= w
        sh = self.rect(x, y, w, h, fond, rayon=h / 2)
        self.texte(x, y, w, h, texte, taille, coul, graisse=graisse, align=CENTRE, ancre=MILIEU, majuscules=majuscules,
                   espacement=esp, forme=sh, interligne=1.2, retour=False)
        return w, h

    def numero(self, x, y, d, n, taille, fond=None, coul=None):
        sh = self.rond(x, y, d, fond or self.deck.c["couleur"])
        self.texte(x, y, d, d, str(n), taille, coul or self.deck.c["surCouleur"], graisse=800, align=CENTRE, ancre=MILIEU,
                   forme=sh, interligne=1.2, retour=False)

    # -- images
    def image(self, donnees: bytes, x, y, w, h):
        return self.s.shapes.add_picture(io.BytesIO(donnees), e(x), e(y), e(w), e(h))

    def photo(self, url, x, y, w, h, libelle="Photo", position="center", saturation=1.0, contraste=1.0):
        """Photo en « cover » ; seules la couverture et « Questions » sont saturées (filtres CSS du modèle)."""
        donnees = self.deck.telecharger(url)
        if not donnees:
            return self.vide(x, y, w, h, libelle, rayon=0, fond=C["photoVide"])
        try:
            img = recadrer(donnees, w, h, position)
        except Exception:
            return self.vide(x, y, w, h, libelle, rayon=0, fond=C["photoVide"])
        if saturation != 1.0:
            img = ImageEnhance.Color(img).enhance(saturation)
        if contraste != 1.0:  # contrast() CSS : écart au gris moyen multiplié
            img = img.point(lambda v: max(0, min(255, round((v - 127.5) * contraste + 127.5))))
        return self.image(jpeg(img, max(1, int(w * 1.25))), x, y, w, h)

    def visuel(self, url, x, y, w, h, libelle, rayon=16):
        """Capture ou schéma ajusté (contain) ; emplacement vide à remplacer sinon."""
        donnees = self.deck.telecharger(url) if url else None
        if not donnees:
            return self.vide(x, y, w, h, libelle, rayon)
        try:
            img = Image.open(io.BytesIO(donnees)).convert("RGB")
        except Exception:
            return self.vide(x, y, w, h, libelle, rayon)
        ratio = min(w / img.width, h / img.height)
        iw, ih = img.width * ratio, img.height * ratio
        return self.image(jpeg(img, int(iw * 1.25)), x + (w - iw) / 2, y + (h - ih) / 2, iw, ih)

    def vide(self, x, y, w, h, libelle, rayon=16, fond=None):
        self.rect(x, y, w, h, fond or C["fondCarte"], rayon=rayon, trait=C["barreGrise"], epaisseur=3, tirets=True)
        ic = min(56, h / 4)
        taille = 24 if w > 200 else 18
        haut = y + h / 2 - (ic + 14 + taille * 1.5) / 2
        self.image(ch.png(ch.icone_image("image", "#B5B5BD", 128)), x + (w - ic) / 2, haut, ic, ic)
        self.texte(x + 16, haut + ic + 14, w - 32, taille * 1.5 * 2, libelle, taille, C["gris"], graisse=600, align=CENTRE)

    def icone(self, nom, x, y, taille, coul="#1F1F24"):
        self.image(ch.png(ch.icone_image(nom, coul, 256)), x, y, taille, taille)

    def sticker(self, x, y, taille, rotation, icone=None):
        if not self.deck.f["stickers"]:
            return
        c = self.deck.c
        img = ch.sticker_image(c["sticker"], icone or c["icone"], taille=taille * 2, rotation=rotation)
        w, h = img.width / 2, img.height / 2
        self.image(ch.png(img), x + taille / 2 - w / 2, y + taille / 2 - h / 2, w, h)

    # -- pied et notes
    def pied(self, couleur_cat=False):
        c, f = self.deck.c, self.deck.f
        n = "{:02d}".format(len(self.deck.prs.slides))
        if couleur_cat:
            self.rect(0, BAS, W, 80, c["couleur"])
            coul = coul_n = c["surCouleur"]
            graisse, variante = 700, ("blanc" if blanc(c) else "noir")
        else:
            self.rect(0, BAS, W, 80, C["fondGris"])
            coul, coul_n, graisse, variante = C["texte4"], C["texte"], 600, "couleur"
        self.texte(64, BAS, 820, 80, f["intitule"], 24, coul, graisse=graisse, ancre=MILIEU)
        # Logo du modèle : PNG haut de 64 px marges comprises, centré ; on place sa partie visible.
        self.image(self.deck.logo(variante), *boite_logo_pied(variante))
        self.texte(W - 64 - 200, BAS, 200, 80, n, 24, coul_n, graisse=700, align=DROITE, ancre=MILIEU, retour=False)

    def notes(self, texte):
        if isinstance(texte, list):
            texte = "\n".join(texte)
        if texte:
            self.s.notes_slide.notes_text_frame.text = ch.typographie(ch.sans_gras(texte))


def tirets_css(ligne) -> None:
    """Tirets d'une bordure « dashed » du navigateur : deux épaisseurs de trait, une d'espace."""
    ligne.dash_style = MSO_LINE_DASH_STYLE.DASH  # crée a:prstDash à sa place dans a:ln, remplacé ci-dessous
    prst = ligne._get_or_add_ln().find(qn("a:prstDash"))
    prst.addprevious(parse_xml('<a:custDash xmlns:a="{}"><a:ds d="200000" sp="100000"/></a:custDash>'.format(A_NS)))
    prst.getparent().remove(prst)


def hauteur_opts(contenu, w, taille, opts) -> float:
    return hauteur(contenu, w, taille, opts.get("interligne", 1.5), opts.get("graisse") or (700 if opts.get("gras") else 400),
                   opts.get("police") == MONO, opts.get("espacement", 0.0))


@lru_cache(maxsize=None)
def boite_logo_pied(variante: str) -> tuple[float, float, float, float]:
    """Partie visible du logo du pied : <img height:64px> centrée, entre y 1008 et 1072."""
    img = Image.open(ch.chemin_logo("couleur" if variante == "couleur" else "blanc"))
    k = 64 / img.height
    gauche, haut, droite, bas = img.getchannel("A").getbbox()
    x0 = (W - img.width * k) / 2
    return x0 + gauche * k, 1008 + haut * k, (droite - gauche) * k, (bas - haut) * k


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
DUREE_FONDU = 400  # ms


def xml_animations(blocs: list[list[int]], formes: set[int]) -> str:
    """Séquence principale PowerPoint : un clic par bloc, fondu en entrée pour chacune de ses formes.

    La première forme du bloc part au clic, les autres avec elle. `formes` : identifiants des `p:sp`,
    qui reçoivent un `p:bldP` (forme et texte animés d'un seul tenant) ; les images n'en ont pas.
    """
    compteur = iter(range(3, 1_000_000))
    clics = []
    for bloc in blocs:
        id_clic, id_groupe = next(compteur), next(compteur)
        effets = []
        for k, spid in enumerate(bloc):
            id_effet, id_set, id_fondu = next(compteur), next(compteur), next(compteur)
            cible = '<p:tgtEl><p:spTgt spid="{}"/></p:tgtEl>'.format(spid)
            effets.append(
                '<p:par><p:cTn id="{}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{}">'
                '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                '<p:set><p:cBhvr><p:cTn id="{}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{}'
                "<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
                '<p:to><p:strVal val="visible"/></p:to></p:set>'
                '<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{}" dur="{}"/>{}</p:cBhvr></p:animEffect>'
                "</p:childTnLst></p:cTn></p:par>".format(
                    id_effet, "clickEffect" if k == 0 else "withEffect", id_set, cible, id_fondu, DUREE_FONDU, cible
                )
            )
        clics.append(
            '<p:par><p:cTn id="{}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="{}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{}'
            "</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>".format(id_clic, id_groupe, "".join(effets))
        )
    vus: list[int] = []
    for spid in (i for bloc in blocs for i in bloc):
        if spid in formes and spid not in vus:
            vus.append(spid)
    construction = "".join('<p:bldP spid="{}" grpId="0" animBg="1"/>'.format(i) for i in vus)
    return (
        '<p:timing xmlns:p="{}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{}</p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        "</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>{}</p:timing>"
    ).format(P_NS, "".join(clics), "<p:bldLst>{}</p:bldLst>".format(construction) if construction else "")


def url_image(url: str) -> str:
    """URL réellement téléchargée (photos Unsplash en 1600 px, suffisant pour une demi-slide)."""
    return re.sub(r"([?&])w=\d+", r"\1w=1600", url)


def recadrer(donnees: bytes, largeur_: float, hauteur_: float, position="center") -> Image.Image:
    """background-size: cover, centré verticalement ; `position` « right » ou « left » cale l'image sur ce bord."""
    image = Image.open(io.BytesIO(donnees)).convert("RGB")
    ratio = largeur_ / hauteur_
    if image.width / image.height > ratio:
        nouvelle = int(image.height * ratio)
        reste = image.width - nouvelle
        gauche = reste if position.startswith("right") else (0 if position.startswith("left") else reste // 2)
        return image.crop((gauche, 0, gauche + nouvelle, image.height))
    nouvelle = int(image.width / ratio)
    haut = (image.height - nouvelle) // 2
    return image.crop((0, haut, image.width, haut + nouvelle))


def jpeg(img: Image.Image, largeur_max: int) -> bytes:
    if img.width > largeur_max:
        img = img.resize((largeur_max, int(img.height * largeur_max / img.width)), Image.LANCZOS)
    tampon = io.BytesIO()
    img.convert("RGB").save(tampon, format="JPEG", quality=82, optimize=True)
    return tampon.getvalue()


class Deck:
    def __init__(self, s: dict) -> None:
        self.s = s
        self.f = s["formation"]
        self.c = self.f["_cat"]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = e(W), e(H)
        self.prs.core_properties.title = self.f["intitule"]
        self.prs.core_properties.author = self.f["formateur"]
        self._cache: dict[str, bytes | None] = {}
        self._logos: dict[str, bytes] = {}

    def telecharger(self, url):
        if not url:
            return None
        if url not in self._cache:
            self._cache[url] = ch.telecharger(url_image(url))
        return self._cache[url]

    def logo(self, variante: str) -> bytes:
        """Logo rogné (sans marges) : « couleur », « blanc », ou « noir » (filtre brightness(0) du modèle)."""
        if variante not in self._logos:
            if variante == "noir":
                img = Image.open(io.BytesIO(ch.logo_rogne("blanc"))).convert("RGBA")
                noir = Image.new("RGBA", img.size, (0, 0, 0, 255))
                noir.putalpha(img.getchannel("A"))
                self._logos[variante] = ch.png(noir)
            else:
                self._logos[variante] = ch.logo_rogne(variante)
        return self._logos[variante]


# ---------------------------------------------------------------- éléments composés


def titre(sl: Slide, x, y, w, texte, taille=72, coul=None, interligne=1.05, espacement=-0.01, graisse=600):
    """Titre h2 du modèle : 72 px, graisse 600, interligne 1,05, lettres resserrées de 0,01 em."""
    h = hauteur(texte, w, taille, interligne, graisse)
    sl.texte(x, y, w, h + 6, texte, taille, coul or C["texte"], graisse=graisse, interligne=interligne, espacement=espacement)
    return y + h


def logo_entete(sl: Slide):
    """Logo de la couverture et de « Questions » : 322 × 106 px, ajusté à gauche ; noir sur les couleurs claires."""
    donnees = sl.deck.logo("blanc" if blanc(sl.deck.c) else "noir")
    img = Image.open(io.BytesIO(donnees))
    k = min(322 / img.width, 106 / img.height)
    sl.image(donnees, 1080, 110 + (106 - img.height * k) / 2, img.width * k, img.height * k)


def contact(sl: Slide):
    """Formateur et e-mail (lien mailto) en bas de la couverture et de « Questions »."""
    f, co = sl.deck.f, sl.deck.c["surCouleur"]
    sl.texte(1080, 898.8, 720, 41, f["formateur"], 28, co, graisse=700, interligne=1.45)
    sl.texte(1080, 939.4, 720, 41, f["email"], 28, co, interligne=1.45, lien="mailto:" + str(f["email"]).strip())


def formater_point(p) -> str:
    if isinstance(p, dict):
        return "**{}** : {}".format(p.get("label"), p.get("texte", "")) if p.get("label") else p.get("texte", "")
    return p or ""


def puces(sl: Slide, x, y, w, items, taille, ecart, interligne, marge, ecart_x=20, coul=C["texte3"]):
    """Liste à puces carrées de 12 px : carré décalé de `marge`, texte à 12 + ecart_x du bord."""
    lt = w - 12 - ecart_x
    for item in items:
        t = formater_point(item)
        h = hauteur(t, lt, taille, interligne)
        with sl.bloc():
            sl.rect(x, y + marge, 12, 12, sl.deck.c["couleur"], rayon=2)
            sl.texte(x + 12 + ecart_x, y, lt, h + 4, t, taille, coul, interligne=interligne, gras_coul=C["texte"],
                     lien=item.get("lien") if isinstance(item, dict) else None)
        y += h + ecart
    return y - ecart


def hauteur_puces(items, w, taille, ecart, interligne, ecart_x=20):
    return sum(hauteur(formater_point(i), w - 12 - ecart_x, taille, interligne) for i in items) + ecart * max(0, len(items) - 1)


def polaroid(url_donnees: bytes | None, largeur_=440, hauteur_=380, rotation=-4) -> Image.Image:
    """Photo façon polaroïd (cadre blanc, ombre 0 16px 40px rgba(0,0,0,.14), rotation), composée en une image."""
    k = 2
    w, h = largeur_ * k, hauteur_ * k
    marge = 60 * k
    toile = Image.new("RGBA", (w + 2 * marge, h + 2 * marge), (0, 0, 0, 0))
    ombre = Image.new("RGBA", toile.size, (0, 0, 0, 0))
    ImageDraw.Draw(ombre).rectangle((marge, marge + 16 * k, marge + w, marge + h + 16 * k), fill=(0, 0, 0, 36))
    toile.alpha_composite(ombre.filter(ImageFilter.GaussianBlur(20 * k)))
    carte = Image.new("RGBA", (w, h), (255, 255, 255, 255))
    pw, ph = w - 36 * k, h - (18 + 56) * k
    if url_donnees:
        try:
            carte.paste(recadrer(url_donnees, pw, ph).resize((pw, ph), Image.LANCZOS), (18 * k, 18 * k))
        except Exception:
            ImageDraw.Draw(carte).rectangle((18 * k, 18 * k, 18 * k + pw, 18 * k + ph), fill=ch.rgb(C["photoVide"]))
    else:
        ImageDraw.Draw(carte).rectangle((18 * k, 18 * k, 18 * k + pw, 18 * k + ph), fill=ch.rgb(C["photoVide"]))
    toile.alpha_composite(carte, (marge, marge))
    return toile.rotate(-rotation, resample=Image.BICUBIC, expand=True)


# ---------------------------------------------------------------- mises en page


def couverture(sl: Slide, s):
    f, c = sl.deck.f, sl.deck.c
    sl.photo(s.get("photo") or f["_couverture"], 0, 0, 960, H, saturation=1.25, contraste=1.06)
    sl.rect(960, 0, 960, H, c["couleur"])
    logo_entete(sl)
    co = c["surCouleur"]
    # Bloc centré entre le logo (bas à 216) et le contact (haut à 898,8) : titre, 32, pastilles, 32, durée.
    t = taille_ajustee(f["intitule"], 720, 104, 700, -0.02)
    h_titre = hauteur(f["intitule"], 720, t, 1.0, 700, espacement=-0.02)
    y = (216 + 898.8) / 2 - (h_titre + 32 + 60 + 32 + 42) / 2
    sl.texte(1080, y, 720, h_titre + 8, equilibrer(f["intitule"], 720, t, 700, -0.02), t, co, graisse=700,
             interligne=1.0, espacement=-0.02)
    y += h_titre + 32
    w, _ = sl.pastille(1080, y, f["categorie"], 24, "#FFFFFF", c["encre"], pad_x=30, pad_y=12, majuscules=False, graisse=700)
    sl.pastille(1080 + w + 14, y, f["niveau"], 24, "#FFFFFF", c["encre"], pad_x=30, pad_y=12, majuscules=False, graisse=700)
    sl.texte(1080, y + 60 + 32, 720, 42, f["duree"], 28, co, graisse=500)
    contact(sl)


def faisons_connaissance(sl: Slide, s):
    d, f, c = sl.deck, sl.deck.f, sl.deck.c
    # Colonne gauche (640 px), centrée verticalement entre 96 et 928.
    titre_s = s.get("titre") or "Faisons connaissance"
    infos = s.get("infos") or []
    h_infos = hauteur_puces(infos, 640, 28, 18, 1.4, ecart_x=18) if infos else 0
    bloc = hauteur(titre_s, 640, 72, 1.05, 600) + 40 + 200 + (40 + h_infos if infos else 0)
    y = titre(sl, 120, 512 - bloc / 2, 640, titre_s) + 40
    with sl.bloc():
        if s.get("photo") and d.telecharger(s["photo"]):
            img = recadrer(d.telecharger(s["photo"]), 400, 400).resize((400, 400), Image.LANCZOS)
            masque = Image.new("L", (400, 400), 0)
            ImageDraw.Draw(masque).ellipse((0, 0, 399, 399), fill=255)
            portrait = img.convert("RGBA")
            portrait.putalpha(masque)
            sl.image(ch.png(portrait), 120, y, 200, 200)
        else:
            sl.rond(120, y, 200, c["teinte"])
            sl.icone("user-round", 172, y + 52, 96, c["encre"])
        # Nom (38 px, ligne de 57) et rôle (26 px, ligne de 39, 6 px plus bas), centrés sur le portrait.
        sl.texte(352, y + 49, 408, 57, f["formateur"], 38, graisse=700)
        sl.texte(352, y + 112, 408, 39, f["role"], 26, C["texte4"])
    y += 240
    if infos:
        puces(sl, 120, y, 640, infos, 28, 18, 1.4, marge=14, ecart_x=18)
    # Panneau « Tour de table » : il apparaît avec son libellé et la première question.
    panneau = sl.rect(856, 96, 944, 832, c["teinte"], rayon=32)
    tour = s["tourDeTable"]
    textes = [hauteur(t, 708, 36, 1.5, 500) for t in tour]
    hauteurs = [max(64, h) for h in textes]
    y = 512 - (36 + 36 + sum(hauteurs) + 36 * (len(tour) - 1)) / 2
    libelle = sl.texte(928, y, 800, 36, "Tour de table", 24, c["encre"], graisse=800, majuscules=True, espacement=0.14)
    y += 36 + 36
    for k, (t, ht, h) in enumerate(zip(tour, textes, hauteurs), 1):
        with sl.bloc(*((panneau, libelle) if k == 1 else ())):
            sl.numero(928, y + (h - 64) / 2, 64, k, 28)
            sl.texte(1020, y + (h - ht) / 2, 708, ht + 4, t, 36, graisse=500)
        y += h + 36
    sl.sticker(1680, 70, 150, 8, "hand")


def objectifs(sl: Slide, s):
    c = sl.deck.c
    sl.photo(s.get("photo"), 1160, 0, 760, BAS)
    y = titre(sl, 120, 96, 944, s.get("titre") or "Objectifs pédagogiques") + 20
    sl.texte(120, y, 944, 45, s.get("intro") or "À la fin de la formation, vous serez capable de :", 30, C["texte4"])
    taille = 34 if len(s["objectifs"]) <= 4 else 30
    hauteurs = [hauteur(o, 868, taille, 1.4) for o in s["objectifs"]]
    y = 920 - sum(hauteurs) - 34 * (len(hauteurs) - 1)
    dn = ligne_de_base(taille, 1.4) - ligne_de_base(28, 1.5)  # numéro aligné sur la ligne de base du texte
    for k, (o, h) in enumerate(zip(s["objectifs"], hauteurs), 1):
        with sl.bloc():
            sl.texte(120, y + dn, 70, 42, "{:02d}".format(k), 28, c["encre"], graisse=800)
            sl.texte(196, y, 868, h + 4, o, taille, C["texte3"], interligne=1.4, gras_coul=C["texte"])
        y += h + 34
    sl.sticker(1075, 150, 170, -10, "target")


def programme(sl: Slide, s):
    d, f, c = sl.deck, sl.deck.f, sl.deck.c
    y = titre(sl, 120, 96, 620, s.get("titre") or "Programme") + 16
    sl.texte(120, y, 620, 45, s.get("duree") or f["duree"], 30, C["texte4"])
    # Polaroïd centré dans la zone libre (x 60 à 740, sous la durée jusqu'à 920).
    img = polaroid(d.telecharger(s.get("photo")))
    w, h = img.width / 2, img.height / 2
    cy = (y + 45 + 920) / 2
    sl.image(ch.png(img), 400 - w / 2, cy - h / 2, w, h)
    sl.sticker(530, 690, 140, 10)
    panneau = sl.rect(820, 0, 1100, BAS, c["couleur"])
    parties = s["parties"]
    taille = 32 if len(parties) <= 8 else 28
    pad = 20 if len(parties) <= 8 else 14
    co = c["surCouleur"]
    # Chaque ligne est une grille 64 px | 1fr | auto, écart 20 : le titre prend la place laissée par la durée.
    lignes = []
    for k, p in enumerate(parties, 1):
        titre_p = p["titre"] if isinstance(p, dict) else p
        duree = (p.get("duree") or "") if isinstance(p, dict) else ""
        wd = largeur(duree, 26, 500) if duree else 0
        lt = 880 - 64 - 20 - (20 + wd if duree else 0)
        ht = hauteur(titre_p, lt, taille, 1.5, 600)
        lignes.append((titre_p, duree, wd, lt, ht, ht + 2 * pad + (2 if k < len(parties) else 0)))
    y = BAS / 2 - sum(x[-1] for x in lignes) / 2
    filet = ch.melange(co, c["couleur"], 0.14)
    dn = ligne_de_base(taille, 1.5) - ligne_de_base(26, 1.5)
    for k, (titre_p, duree, wd, lt, ht, h) in enumerate(lignes, 1):
        with sl.bloc(*((panneau,) if k == 1 else ())):
            sl.texte(930, y + pad, 64, ht, "{}.".format(k), taille, co, graisse=800)
            sl.texte(1014, y + pad, lt, ht + 4, titre_p, taille, co, graisse=600)
            if duree:
                sl.texte(1810 - wd - 20, y + pad + dn, wd + 20, 39, duree, 26, co, graisse=500, align=DROITE, retour=False)
            if k < len(parties):
                sl.ligne(930, y + h - 2, 880, 2, filet)
        y += h


def chapitre(sl: Slide, s):
    c = sl.deck.c
    sl.photo(s.get("photo"), 1160, 0, 760, BAS)
    sl.texte(120, 88, 760, 170, "{:02d}".format(int(s["numero"])), 200, c["encre"], graisse=800, interligne=0.85,
             espacement=-0.04, retour=False)
    if s.get("duree"):
        sl.pastille(1064, 88, s["duree"], 28, c["teinte"], c["encre"], pad_x=30, pad_y=12, majuscules=False, droite=True)
    taille = 88 if len(s["titre"]) <= 40 else 72
    titre(sl, 120, 298, 944, s["titre"], taille, interligne=1.02, espacement=-0.015)
    # Points en bas de colonne (bas à 912), carré centré sur la ligne.
    points = s.get("points") or []
    hauteurs = [hauteur(p, 912, 32, 1.5) for p in points]
    y = 912 - sum(hauteurs) - 18 * max(0, len(points) - 1)
    for p, h in zip(points, hauteurs):
        sl.rect(120, y + h / 2 - 6, 12, 12, c["couleur"], rayon=2)
        sl.texte(152, y, 912, h + 4, p, 32, C["texte3"], gras_coul=C["texte"])
        y += h + 18
    sl.sticker(1080, 770, 160, -8)


def contenu(sl: Slide, s):
    x, w = 120, 824
    y = titre(sl, x, 96, w, s["titre"]) + 36
    if s.get("texte"):
        with sl.bloc():
            y = sl.pile(x, y, w, [(s["texte"], 32, {"coul": C["texte3"], "interligne": 1.55, "gras_coul": C["texte"]})]) + 36
    puces(sl, x, y, w, s.get("points") or [], 30, 22, 1.45, marge=15)
    with sl.bloc():
        sl.rect(1040, 188, 760, 640, C["fondCarte"], rayon=28)
        sl.visuel(s.get("image"), 1068, 216, 704, 584, s.get("legendeImage") or "Capture d'écran ou schéma")


def chiffre_cle(sl: Slide, s):
    c = sl.deck.c
    chiffre = str(s["chiffre"])
    taille = 280 if len(chiffre) <= 3 else (220 if len(chiffre) <= 5 else 160)
    h_phrase = hauteur(s["phrase"], 780, 44, 1.2, 700)
    y = 508 - (taille * 0.85 + 32 + h_phrase) / 2
    with sl.bloc():
        tb = sl.texte(120, y, 780, taille * 0.85, chiffre, taille, c["encre"], graisse=800, interligne=0.85, espacement=-0.04,
                      retour=False)
        if s.get("unite"):
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = s["unite"]
            r.font.name, r.font.bold = GRAISSES[800]
            r.font.size = Pt(taille * 150 / 280 / 2)
            r.font.color.rgb = couleur(c["encre"])
            # letter-spacing hérité du chiffre (−0,04 em de 280 px).
            r._r.get_or_add_rPr().set("spc", str(int(round(-0.04 * taille / 2 * 100))))
        sl.texte(120, y + taille * 0.85 + 32, 780, h_phrase + 4, equilibrer(s["phrase"], 780, 44, 700), 44, graisse=700,
                 interligne=1.2)
    # Colonne droite : explication, repères, source ; centrée sur 508.
    reperes = (s.get("reperes") or [])[:2]
    h_expl = hauteur(s.get("explication", ""), 780, 34, 1.5)
    h_rep = 32 + 64 + 10 + max([hauteur(r.get("texte"), 306, 26, 1.4) for r in reperes] or [0]) + 32
    bloc = h_expl + (40 + h_rep if reperes else 0) + (40 + 36 if s.get("source") else 0)
    y = 508 - bloc / 2
    if s.get("explication"):
        with sl.bloc():
            sl.texte(1020, y, 780, h_expl + 4, s["explication"], 34, C["texte3"], gras_coul=C["texte"])
        y += h_expl + 40
    if reperes:
        for k, r in enumerate(reperes):
            x = 1020 + k * (378 + 24)
            with sl.bloc():
                sl.rect(x, y, 378, h_rep, c["teinte"], rayon=24)
                sl.texte(x + 36, y + 32, 306, 64, r.get("valeur"), 64, c["encre"], graisse=800, interligne=1.0, retour=False)
                sl.texte(x + 36, y + 106, 306, h_rep - 138 + 4, r.get("texte"), 26, C["texte2"], interligne=1.4)
        y += h_rep + 40
    if s.get("source"):
        with sl.bloc():
            sl.texte(1020, y, 780, 36, "Source : " + s["source"], 24, C["gris"])


def frise(sl: Slide, s):
    c = sl.deck.c
    y = titre(sl, 120, 96, 1680, s["titre"]) + 20
    if s.get("intro"):
        h_i = hauteur(s["intro"], 1680, 30, 1.5)
        sl.texte(120, y, 1680, h_i + 4, s["intro"], 30, C["texte4"], gras_coul=C["texte"])
        y += h_i + 20
    milieu = (y + 920) / 2
    axe = sl.rect(120, milieu - 3, 1680, 6, c["teinte"], rayon=3)
    etapes = s["etapes"]
    n = len(etapes)
    lc = (1680 - 32 * (n - 1)) / n
    for k, et in enumerate(etapes):
        x = 120 + k * (lc + 32)
        derniere = k == n - 1
        h_t = hauteur(et.get("texte"), lc, 28, 1.4)
        # Date (ligne de 52), 10 px, texte : au-dessus de l'axe (bas à 32 px du rond) ou au-dessous.
        yb = milieu - 18 - 32 - (52 + 10 + h_t) if k % 2 == 0 else milieu + 18 + 32
        with sl.bloc(*((axe,) if k == 0 else ())):
            sl.rond(x, milieu - 18, 36, c["couleur"] if derniere else "#FFFFFF", trait=c["couleur"], epaisseur=6)
            t = taille_ajustee(et.get("date"), lc, 52, 800)
            sl.texte(x, yb + (52 - t), lc, 52, et.get("date"), t, c["encre"] if derniere else C["texte"], graisse=800, interligne=1.0)
            sl.texte(x, yb + 62, lc, h_t + 4, et.get("texte"), 28, C["texte3"], interligne=1.4, gras_coul=C["texte"])


def graphique(sl: Slide, s):
    c = sl.deck.c
    y = titre(sl, 120, 96, 960, s["titre"]) + 40
    barres = s["barres"]
    valeurs = [float(b.get("valeur", 0)) for b in barres]
    maxi = max(valeurs) or 1
    cible = s.get("miseEnAvant", len(barres) - 1)
    n = len(barres)
    base = 861  # bas des barres ; l'axe (3 px) va de 861 à 864, les années à 884
    lc = (944 - 28 * (n - 1)) / n
    zone = base - y - 56
    with sl.bloc():
        for k, b in enumerate(barres):
            x = 128 + k * (lc + 28)
            hb = zone * round(valeurs[k] / maxi, 2)  # même arrondi que le HTML
            if hb > 0:
                sl.rect(x, base - hb, lc, hb, c["couleur"] if k == cible else C["barreGrise"], rayon=min(12, hb), coins="haut")
            valeur = b.get("affichage") or _nombre(b.get("valeur"))
            sl.texte(x - 20, base - hb - 14 - 42, lc + 40, 42, valeur, 28, c["encre"] if k == cible else C["texte4"],
                     graisse=800, align=CENTRE, retour=False)
            sl.texte(x - 14, 884, lc + 28, 36, b.get("label"), 24, C["gris"], graisse=600, align=CENTRE)
        sl.rect(120, base, 960, 3, C["texte"])
    elements = [
        (s.get("surtitre") or "À retenir", 24, {"coul": c["encre"], "graisse": 800, "majuscules": True, "espacement": 0.14}),
        (s.get("chiffre", ""), 128, {"coul": c["encre"], "graisse": 800, "interligne": 0.9, "espacement": -0.02}),
        (s.get("message", ""), 32, {"coul": C["texte3"], "gras_coul": C["texte"]}),
        ("Source : " + s["source"] if s.get("source") else "", 24, {"coul": C["gris"]}),
    ]
    presents = [x for x in elements if x[0]]
    total = sum(hauteur_opts(t, 600, tl, o) for t, tl, o in presents) + 28 * (len(presents) - 1)
    with sl.bloc():
        sl.pile(1200, 508 - total / 2, 600, elements, 28)


def _nombre(v) -> str:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v or "")
    return ("{:,.0f}" if x == int(x) else "{:,.1f}").format(x).replace(",", " ").replace(".", ",")


def image(sl: Slide, s):
    sl.photo(s.get("image"), 0, 0, W, H)
    surtitre = s.get("surtitre") or "Légende"
    texte_ = s.get("texte", "")
    # Carte de légende : largeur au contenu (928 px de texte au plus), bas à 960, ombre 0 20px 50px .25.
    lt = min(928, max(largeur(surtitre, 24, 800, espacement=0.14, majuscules=True), largeur(s["titre"], 56, 700),
                      largeur(texte_, 28) if texte_ else 0) + 2)
    h_t = hauteur(s["titre"], lt, 56, 1.1, 700)
    h_x = hauteur(texte_, lt, 28, 1.45) if texte_ else 0
    w, h = lt + 112, 48 + 36 + 16 + h_t + (16 + h_x if h_x else 0) + 48
    y = 960 - h
    sl.ombre(120, y, w, h, 28, 20, 50, 0.25)
    sl.rect(120, y, w, h, "#FFFFFF", rayon=28)
    sl.pile(176, y + 48, lt, [
        (surtitre, 24, {"coul": sl.deck.c["encre"], "graisse": 800, "majuscules": True, "espacement": 0.14}),
        (equilibrer(s["titre"], lt, 56, 700), 56, {"graisse": 700, "interligne": 1.1}),
        (texte_, 28, {"coul": C["texte3"], "interligne": 1.45, "gras_coul": C["texte"]}),
    ], 16)
    n = "{:02d}".format(len(sl.deck.prs.slides))
    sl.pastille(W - 64, 980, n, 24, "#FFFFFF", C["texte"], pad_x=22, pad_y=8, majuscules=False, droite=True, graisse=700)


def texte(sl: Slide, s):
    c = sl.deck.c
    paragraphes = s["paragraphes"]
    sl.rect(120, 96, 96, 10, c["couleur"], rayon=5)
    y = titre(sl, 120, 138, 520, s["titre"]) + 32
    if s.get("intro"):
        sl.texte(120, y, 520, hauteur(s["intro"], 520, 30, 1.5) + 4, s["intro"], 30, C["texte4"], gras_coul=C["texte"])
    deux = len(paragraphes) > 1
    lc = 484 if deux else 1040
    hauteurs = [hauteur(p, lc, 30, 1.6) for p in paragraphes]
    # Répartition en deux colonnes équilibrées.
    coupe = len(paragraphes)
    if deux:
        meilleur = None
        for i in range(1, len(paragraphes)):
            m = max(sum(hauteurs[:i]) + 28 * (i - 1), sum(hauteurs[i:]) + 28 * (len(paragraphes) - i - 1))
            if meilleur is None or m < meilleur:
                meilleur, coupe = m, i
    colonnes = [list(zip(paragraphes[:coupe], hauteurs[:coupe])), list(zip(paragraphes[coupe:], hauteurs[coupe:]))]
    total = max(sum(h for _, h in col) + 28 * max(0, len(col) - 1) for col in colonnes)
    for k, col in enumerate(colonnes):
        y = 508 - total / 2
        for p, h in col:
            with sl.bloc():
                sl.texte(760 + k * (484 + 72), y, lc, h + 4, p, 30, C["texte2"], interligne=1.6, gras_coul=C["texte"])
            y += h + 28


def citation(sl: Slide, s):
    c = sl.deck.c
    co = c["surCouleur"]
    sl.rect(0, 0, W, BAS, c["couleur"])
    taille = 72 if len(s["citation"]) <= 100 else 60
    h_q = hauteur(s["citation"], 1500, taille, 1.2, 600)
    y = 500 - (110 + 56 + h_q + 56 + 120) / 2
    sl.texte(200, y, 300, 132, "«", 220, co, graisse=800, interligne=0.6, retour=False)
    y += 110 + 56
    with sl.bloc():
        sl.texte(200, y, 1500, h_q + 8, s["citation"], taille, co, graisse=600, interligne=1.2, espacement=-0.01)
    y += h_q + 56
    with sl.bloc():
        sl.rond(200, y, 120, "#FFFFFF")
        sl.icone("user-round", 228, y + 28, 64)
        sl.texte(348, y + 34.5, 1200, 51, s.get("auteur") or "", 34, co, graisse=700)


def avant_apres(sl: Slide, s):
    c = sl.deck.c
    y = titre(sl, 120, 96, 1680, s["titre"]) + 48
    lc = (1680 - 48) / 2
    h_box = 920 - y - 28 - 56
    yb = y + h_box + 28
    for k, (cle, libelle) in enumerate((("avant", "Avant"), ("apres", "Après"))):
        d = s.get(cle) or {}
        x = 120 + k * (lc + 48)
        fond, coul = (C["bordure"], C["texte"]) if k == 0 else (c["couleur"], c["surCouleur"])
        with sl.bloc():
            sl.rect(x, y, lc, h_box, C["fondCarte"], rayon=28)
            sl.visuel(d.get("image"), x + 20, y + 20, lc - 40, h_box - 40, "Capture " + libelle.lower())
            w, _ = sl.pastille(x, yb, libelle, 24, fond, coul, pad_x=26)
            sl.texte(x + w + 24, yb, lc - w - 24, 56, d.get("texte", ""), 30, C["texte3"], ancre=MILIEU, interligne=1.4,
                     gras_coul=C["texte"])


def notion_cle(sl: Slide, s):
    c = sl.deck.c
    sl.rect(0, 0, W, BAS, c["teinte"])
    taille = 112 if len(s["titre"]) <= 26 else 84
    h_t = hauteur(s["titre"], 1600, taille, 1.0, 700, espacement=-0.02)
    h_d = hauteur(s["definition"], 1400, 44, 1.4)
    # Termes : cartes de 78 px en flex-wrap (écart 20), une nouvelle ligne quand la largeur de 1 600 px est atteinte.
    cartes, x, dy = [], 160, 0
    for t in s.get("termes") or []:
        lt = min(1600, 28 + largeur(t.get("terme"), 28, 700, mono=True) + 14 + largeur(t.get("definition"), 28) + 28)
        if x > 160 and x + lt > 1760:
            x, dy = 160, dy + 78 + 20
        cartes.append((x, dy, lt, t))
        x += lt + 20
    h_termes = (16 + dy + 78) if cartes else 0
    y = 500 - (56 + 40 + h_t + 40 + h_d + (40 + h_termes if cartes else 0)) / 2
    sl.pastille(160, y, s.get("surtitre") or "Notion clé", 24)
    y += 56 + 40
    sl.texte(160, y, 1600, h_t + 8, s["titre"], taille, graisse=700, interligne=1.0, espacement=-0.02)
    y += h_t + 40
    with sl.bloc():
        # Le gras de la définition garde la couleur du texte (strong sans couleur dans le modèle).
        sl.texte(160, y, 1400, h_d + 6, s["definition"], 44, C["texte2"], interligne=1.4, gras_coul=C["texte2"])
    y += h_d + 40 + 16
    for x, dy, lt, t in cartes:
        with sl.bloc():
            sl.rect(x, y + dy, lt, 78, "#FFFFFF", rayon=18)
            sl.texte(x + 28, y + dy + 18, lt - 40, 42, [
                (t.get("terme", "") + " ", {"police": MONO, "graisse": 700, "coul": c["encre"]}),
                (t.get("definition", ""), {"coul": C["texte3"]}),
            ], 28, retour=False, ligne_raleway=True)


def code(sl: Slide, s):
    c = sl.deck.c
    elements = s.get("elements") or []
    intro = s.get("intro") or ""
    # Colonne gauche (724 px) centrée sur 508 : titre, 36, intro, 36, éléments (code puis texte, écart 24).
    hs_el = [40.6 + hauteur(x.get("texte"), 724, 28, 1.45) for x in elements]
    h_intro = hauteur(intro, 724, 32, 1.55) if intro else 0
    h_g = hauteur(s["titre"], 724, 72, 1.05, 600) + (36 + h_intro if intro else 0)
    h_g += (36 + sum(hs_el) + 24 * (len(elements) - 1)) if elements else 0
    y = titre(sl, 120, 508 - h_g / 2, 724, s["titre"]) + 36
    if intro:
        sl.texte(120, y, 724, h_intro + 4, intro, 32, C["texte3"], interligne=1.55, gras_coul=C["texte"])
        y += h_intro + 36
    for x, h in zip(elements, hs_el):
        with sl.bloc():
            sl.texte(120, y, 724, 40.6, x.get("code"), 28, c["encre"], graisse=700, police=MONO, interligne=1.45,
                     ligne_raleway=True, retour=False)
            sl.texte(120, y + 40.6, 724, h - 40.6 + 4, x.get("texte"), 28, C["texte3"], interligne=1.45, gras_coul=C["texte"])
        y += h + 24
    # Encadré sombre (padding 52 56, écarts 22) : libellé, syntaxe ; filet (bas du code + 34), « Exemple »
    # (filet + 36), exemple ; résultat (bas du code + 22). Il apparaît avec la syntaxe, puis l'exemple et le résultat.
    h_syn = 51 * len(str(s["syntaxe"]).split("\n"))
    h_ex = 51 * len(str(s["exemple"]).split("\n")) if s.get("exemple") else 0
    h_res = hauteur(s.get("resultat", ""), 748, 28, 1.5) if s.get("resultat") else 0
    h = 52 + 36 + 22 + h_syn + (128 + h_ex if h_ex else 0) + (22 + h_res if h_res else 0) + 52
    y = 508 - h / 2
    libelle = {"coul": C["codeLibelle"], "graisse": 700, "majuscules": True, "espacement": 0.14}
    with sl.bloc():
        sl.rect(940, y, 860, h, C["codeFond"], rayon=28)
        sl.texte(996, y + 52, 748, 36, s.get("libelle") or "Syntaxe", 24, **libelle)
        code_colore(sl, 996, y + 110, 748, h_syn, s["syntaxe"], c["couleur"])
    y += 110 + h_syn
    if h_ex:
        with sl.bloc():
            sl.ligne(996, y + 34, 748, 2, ch.melange("#FFFFFF", C["codeFond"], 0.12))
            sl.texte(996, y + 70, 748, 36, "Exemple", 24, **libelle)
            code_colore(sl, 996, y + 128, 748, h_ex, s["exemple"], c["couleur"])
        y += 128 + h_ex
    if h_res:
        with sl.bloc():
            sl.texte(996, y + 22, 748, h_res + 4, s["resultat"], 28, C["codeTexte"], gras_coul="#FFFFFF")


def code_colore(sl: Slide, x, y, w, h, source, coul_fonction):
    """Code en Consolas 34 px, interligne 1,5, sans retour automatique : fonctions en couleur, chaînes en jaune."""
    lignes = []
    for ligne in (source or "").split("\n"):
        morceaux = []
        for morceau in re.split(r'("[^"]*"|\'[^\']*\'|\b[A-Z][A-Z0-9_.]{1,}(?=\())', ligne):
            if not morceau:
                continue
            if morceau.startswith(("\"", "'")):
                coul = C["codeChaine"]
            elif re.fullmatch(r"[A-Z][A-Z0-9_.]{1,}", morceau):
                coul = coul_fonction
            else:
                coul = "#FFFFFF"
            morceaux.append((morceau, coul))
        lignes.append(morceaux or [("", "#FFFFFF")])
    tb = sl.texte(x, y, w, h + 6, "", 34, police=MONO, retour=False)
    tf = tb.text_frame
    for k, morceaux in enumerate(lignes):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.line_spacing = round(1.5 / 1.2, 4)
        for morceau, coul in morceaux:
            r = p.add_run()
            r.text = morceau
            r.font.name, r.font.size = MONO, Pt(17)
            r.font.color.rgb = couleur(coul)


def capture(sl: Slide, s):
    c = sl.deck.c
    y0 = titre(sl, 120, 88, 1680, s["titre"]) + 44
    y1 = 928
    zx, zy, zw, zh = 140, y0 + 20, 1100, y1 - y0 - 40
    with sl.bloc():
        sl.rect(120, y0, 1140, y1 - y0, C["fondCarte"], rayon=24)
        sl.visuel(s.get("image"), zx, zy, zw, zh, "Capture d'écran à annoter", 12)
    # Chaque repère (rond blanc de 64 px, rond couleur de 56 px, ombre) apparaît avec sa légende.
    reperes = s["reperes"]
    titres = [hauteur(r.get("titre"), 418, 32, 1.4, 700) for r in reperes]
    hauteurs = [ht + hauteur(r.get("texte"), 418, 28, 1.4) for r, ht in zip(reperes, titres)]
    y = (y0 + y1) / 2 - (sum(hauteurs) + 36 * (len(reperes) - 1)) / 2
    for k, (r, ht, h) in enumerate(zip(reperes, titres, hauteurs), 1):
        px = zx + zw * float(r.get("x", 15 + 70 * ((k - 1) % 2))) / 100 - 32
        py = zy + zh * float(r.get("y", 15 + 25 * ((k - 1) // 2))) / 100 - 32
        with sl.bloc():
            sl.ombre(px, py, 64, 64, 32, 6, 16, 0.2, rond=True)
            sl.rond(px, py, 64, "#FFFFFF")
            sl.numero(px + 4, py + 4, 56, k, 28)
            sl.texte(1332, y, 40, 48, str(k), 32, c["encre"], graisse=800, retour=False)
            sl.texte(1382, y, 418, ht + 4, r.get("titre"), 32, graisse=700, interligne=1.4)
            sl.texte(1382, y + ht, 418, h - ht + 4, r.get("texte"), 28, C["texte3"], interligne=1.4, gras_coul=C["texte"])
        y += h + 36


def pas_a_pas(sl: Slide, s):
    y = titre(sl, 120, 96, 1680, s["titre"]) + 64
    etapes = s["etapes"]
    n = len(etapes)
    lc = (1680 - 32 * (n - 1)) / n
    for k, et in enumerate(etapes):
        x = 120 + k * (lc + 32)
        d = et if isinstance(et, dict) else {"texte": et}
        with sl.bloc():
            sl.rect(x, y, lc, 260, C["fondCarte"], rayon=20)
            sl.visuel(d.get("image"), x + 12, y + 12, lc - 24, 236, "Capture étape {}".format(k + 1), 12)
            sl.numero(x, y + 288, 64, k + 1, 28)
            sl.texte(x, y + 380, lc, 300, d.get("texte"), 30, C["texte3"], interligne=1.4, gras_coul=C["texte"])


def comparatif(sl: Slide, s):
    c = sl.deck.c
    y = titre(sl, 120, 96, 1680, s["titre"]) + 56
    colonnes = (s["colonnes"] + ["", ""])[:2]
    lignes = s["lignes"]
    # Cadre de 2 px (rayon 24) ; cellules dedans : colonnes 440 | 618 | 618 à x 122, 562, 1180 ; padding 26 32.
    h_tete = max(hauteur(t, 554, 30, 1.4, 800) for t in colonnes) + 52
    hauteurs = []
    for ligne in lignes:
        v = ((ligne.get("valeurs") or []) + ["", ""])[:2]
        hauteurs.append(2 + 52 + max(hauteur(ligne.get("critere"), 376, 28, 1.4, 700), hauteur(v[0], 554, 28, 1.4),
                                     hauteur(v[1], 554, 28, 1.4), 28 * 1.4))
    total = 2 + h_tete + sum(hauteurs) + 2
    # Le cadre et l'en-tête restent fixes ; chaque ligne apparaît ensuite (animée ou non).
    sl.rect(120, y, 1680, total, "#FFFFFF", rayon=24, trait=C["filet"], epaisseur=2)
    yt = y + 2
    sl.rect(122, yt, 1058, h_tete, C["fondGris"], rayon=22, coins="hg")
    sl.rect(1180, yt, 618, h_tete, c["couleur"], rayon=22, coins="hd")
    sl.texte(594, yt + 26, 554, h_tete - 52, colonnes[0], 30, graisse=800, interligne=1.4)
    sl.texte(1212, yt + 26, 554, h_tete - 52, colonnes[1], 30, c["surCouleur"], graisse=800, interligne=1.4)
    yl = yt + h_tete
    for k, (ligne, h) in enumerate(zip(lignes, hauteurs)):
        v = ((ligne.get("valeurs") or []) + ["", ""])[:2]
        derniere = k == len(lignes) - 1
        with sl.bloc():
            # Dernière cellule teintée : coin bas droit arrondi, comme le cadre.
            sl.rect(1180, yl, 618, h, c["teinte"], rayon=22 if derniere else 0, coins="bd")
            sl.ligne(122, yl, 1676, 2, C["filet"])
            sl.texte(154, yl + 28, 376, h - 54, ligne.get("critere"), 28, graisse=700, interligne=1.4)
            sl.texte(594, yl + 28, 554, h - 54, v[0], 28, C["texte3"], interligne=1.4, gras_coul=C["texte"])
            sl.texte(1212, yl + 28, 554, h - 54, v[1], 28, interligne=1.4)
        yl += h
    if s.get("recommandation"):
        with sl.bloc():
            sl.texte(120, yl + 2 + 56, 1680, 90, "**Recommandation** : " + s["recommandation"], 28, C["texte4"],
                     gras_coul=c["encre"])


def raccourcis(sl: Slide, s):
    titre_s = s.get("titre") or "Raccourcis à connaître"
    remarque = s.get("remarque") or ""
    w_rem = min(800, largeur(remarque, 26) + 4) if remarque else 0
    y = titre(sl, 120, 96, 1680 - (w_rem + 40 if remarque else 0), titre_s)
    if remarque:  # alignée à droite, sur le bas du titre
        h_r = hauteur(remarque, w_rem, 26, 1.5)
        sl.texte(1800 - w_rem, y - h_r, w_rem, h_r + 4, remarque, 26, C["gris"], align=DROITE, gras_coul=C["texte"])
    y += 56
    plus = largeur("+", 24)
    for k, r in enumerate(s["raccourcis"]):
        x = 120 + (k % 2) * 860
        yy = y + (k // 2) * 105
        with sl.bloc():
            sl.rect(x, yy, 820, 89, C["fondCarte"], rayon=20)
            xk = x + 32
            for i, t in enumerate(touches(r)):
                if i:  # « + » à 10 px de chaque touche
                    sl.texte(xk + 10 - 4, yy + 20, plus + 8, 49, "+", 24, C["gris"], align=CENTRE, ancre=MILIEU, interligne=1.2,
                             retour=False)
                    xk += 10 + plus + 10
                # Touche : bordure de 2 px, 5 px en bas (deux rectangles arrondis superposés).
                lt = largeur(t, 26, 700) + 36 + 4
                sl.rect(xk, yy + 20, lt, 49, C["barreGrise"], rayon=12)
                sl.rect(xk + 2, yy + 22, lt - 4, 42, "#FFFFFF", rayon=10)
                sl.texte(xk + 2, yy + 22, lt - 4, 42, t, 26, graisse=700, align=CENTRE, ancre=MILIEU, interligne=1.2, retour=False)
                xk += lt
            sl.texte(x + 376, yy, 412, 89, r.get("action"), 30, ancre=MILIEU, interligne=1.5)


def bon_a_savoir(sl: Slide, s):
    from support_html import genre_encadre

    y = titre(sl, 120, 96, 1680, s.get("titre") or "Bon à savoir") + 56
    encadres = s["encadres"]
    n = len(encadres)
    lc = (1680 - 32 * (n - 1)) / n
    for k, en in enumerate(encadres):
        libelle, fond, encre = genre_encadre(en)
        fond = sl.deck.c["teinte"] if fond.startswith("var") else fond
        encre = sl.deck.c["encre"] if encre.startswith("var") else encre
        x = 120 + k * (lc + 32)
        with sl.bloc():
            sl.rect(x, y, lc, 920 - y, fond, rayon=28)
            sl.pile(x + 48, y + 48, lc - 96, [
                (en.get("libelle") or libelle, 24, {"coul": encre, "graisse": 800, "majuscules": True, "espacement": 0.14}),
                (en.get("titre", ""), 40, {"graisse": 700, "interligne": 1.15}),
                (en.get("texte", ""), 30, {"coul": C["texte2"], "gras_coul": C["texte"]}),
            ], 24)


def exercice(sl: Slide, s):
    c = sl.deck.c
    panneau = sl.rect(1160, 0, 760, BAS, C["fondCarte"])
    x, w = 120, 920
    sl.pastille(x, 88, "Exercice {}".format(s["numero"]), 24)
    y = titre(sl, x, 88 + 56 + 32, w, s["titre"], 64, interligne=1.08) + 32
    if s.get("contexte"):
        y = sl.pile(x, y, w, [(s["contexte"], 30, {"coul": C["texte3"], "gras_coul": C["texte"]})]) + 32
    y += 8
    for k, consigne in enumerate(s["consignes"], 1):
        ht = hauteur(consigne, w - 70, 30, 1.4)
        with sl.bloc():
            sl.numero(x, y, 48, k, 24)
            sl.texte(x + 70, y, w - 70, ht + 4, consigne, 30, C["texte3"], interligne=1.4, gras_coul=C["texte"])
        y += max(48, ht) + 20
    # Cartes : padding 28 32, libellé (36), 6, valeur ; colonne centrée sur 504.
    cartes = []
    if s.get("duree"):
        cartes.append(("Durée", s["duree"], 34, 700, None, 1.5, False))
    if s.get("fichier"):
        cartes.append(("Fichier", s["fichier"], 30, 600, MONO, 1.5, False))
    if s.get("livrable"):
        cartes.append(("Livrable", s["livrable"], 28, 400, None, 1.4, False))
    if s.get("bonus"):
        cartes.append(("Pour aller plus loin", s["bonus"], 28, 600, None, 1.4, True))
    valeurs = [hauteur(v, 496, t, lh, g, p == MONO) for _, v, t, g, p, lh, _ in cartes]
    hauteurs = [28 + 36 + 6 + hv + 28 for hv in valeurs]
    y = 504 - (sum(hauteurs) + 20 * max(0, len(cartes) - 1)) / 2
    # Le panneau gris apparaît avec la première carte, puis chaque carte à son tour.
    for k, ((libelle, valeur, taille, graisse, police, lh, plein), hv, h) in enumerate(zip(cartes, valeurs, hauteurs)):
        fond, coul_l, coul_v = (c["couleur"], c["surCouleur"], c["surCouleur"]) if plein else ("#FFFFFF", c["encre"], C["texte"])
        with sl.bloc(*((panneau,) if k == 0 else ())):
            sl.rect(1240, y, 560, h, fond, rayon=22)
            sl.texte(1272, y + 28, 496, 36, libelle, 24, coul_l, graisse=800, majuscules=True, espacement=0.1)
            sl.texte(1272, y + 70, 496, hv + 4, valeur, taille, coul_v, graisse=graisse, police=police, interligne=lh,
                     gras_coul=coul_v)
        y += h + 20


def exercice_etapes(sl: Slide, s):
    c = sl.deck.c
    # En-tête : bloc gauche (pastille, 24, titre, 24, contexte) et cercle de 240, centrés l'un sur l'autre.
    h_t = hauteur(s["titre"], 1200, 64, 1.08, 600)
    h_c = hauteur(s.get("contexte", ""), 1200, 30, 1.5) if s.get("contexte") else 0
    gauche = 56 + 24 + h_t + (24 + h_c if h_c else 0)
    h_entete = max(240, gauche)
    yg = 88 + (h_entete - gauche) / 2
    cy = 88 + (h_entete - 240) / 2
    sl.pastille(120, yg, "Exercice {}".format(s["numero"]), 24)
    titre(sl, 120, yg + 80, 1200, s["titre"], 64, interligne=1.08)
    if h_c:
        sl.texte(120, yg + 80 + h_t + 24, 1200, h_c + 4, s["contexte"], 30, C["texte3"], gras_coul=C["texte"])
    # Informations en ligne (écart 56, retour à la ligne avec 20 d'écart), sous un filet, en bas à 936.
    infos = [(lib, cle, s[cle]) for lib, cle in (("Fichier", "fichier"), ("Modalité", "modalite"), ("Livrable", "livrable")) if s.get(cle)]
    places, x, rang = [], 120, 0
    for lib, cle, val in infos:
        mono = cle == "fichier"
        wi = largeur("**{}** : ".format(lib), 28) + largeur(val, 28, mono=mono)
        if x > 120 and x + wi > 1800:
            x, rang = 120, rang + 1
        places.append((x, rang, wi, lib, val, mono))
        x += wi + 56
    h_infos = 2 + 28 + (rang + 1) * 42 + rang * 20 if infos else 0
    bas_etapes = 936 - (h_infos + 44 if infos else 0)
    y = 88 + h_entete + 44
    etapes = s["etapes"]
    n = len(etapes)
    lc = (1680 - 28 * (n - 1)) / n
    for k, et in enumerate(etapes):
        d = et if isinstance(et, dict) else {"texte": et}
        x = 120 + k * (lc + 28)
        with sl.bloc():
            sl.rect(x, y, lc, bas_etapes - y, C["fondCarte"], rayon=28)
            sl.pile(x + 44, y + 40, lc - 88, [
                (str(k + 1), 96, {"coul": c["encre"], "graisse": 800, "interligne": 0.9}),
                (d.get("titre", ""), 34, {"graisse": 700, "interligne": 1.2}),
                (d.get("texte", ""), 28, {"coul": C["texte3"], "interligne": 1.45, "gras_coul": C["texte"]}),
            ], 18)
    # Cercle de durée (240 px, trait de 10 à l'intérieur), puis infos : ils viennent après les étapes.
    duree = str(s.get("duree", "")).strip()
    nombre, _, unite = re.sub(r"[\s\u00a0\u202f]+", " ", duree).partition(" ")  # espace insécable comprise
    if not unite or " " in unite:  # « 20 minutes » : nombre et unité ; « 1 h 30 » : d'un bloc
        nombre, unite = duree, ""
    unite = unite or ("minutes" if nombre.isdigit() else "")
    tn = 84 if len(nombre) <= 3 else 56
    bloc = tn + (4 + 39 if unite else 0)
    with sl.bloc():
        sl.rond(1560, cy, 240, None, trait=c["couleur"], epaisseur=10)
        sl.texte(1560, cy + 120 - bloc / 2, 240, tn, nombre or "–", tn, c["encre"], graisse=800, align=CENTRE, interligne=1.0,
                 retour=False)
        if unite:
            sl.texte(1560, cy + 120 - bloc / 2 + tn + 4, 240, 39, unite, 26, C["texte4"], graisse=700, align=CENTRE, retour=False)
    if infos:
        yi = bas_etapes + 44
        with sl.bloc():
            sl.ligne(120, yi, 1680, 2, C["filet"])
            for x, rang, wi, lib, val, mono in places:
                valeur = (val, {"police": MONO, "coul": C["texte"]}) if mono else (ch.sans_gras(val), {})
                sl.texte(x, yi + 30 + rang * 62, wi + 20, 42, [
                    (lib, {"graisse": 700, "coul": c["encre"]}),
                    (" : ", {}),
                    valeur,
                ], 28, C["texte3"], retour=False, ligne_raleway=True)


def quiz(sl: Slide, s):
    c = sl.deck.c
    w, _ = sl.pastille(120, 96, "Quiz", 24)
    sl.texte(120 + w + 24, 103, 700, 42, "Question {} sur {}".format(s.get("numero", 1), s.get("total", 1)), 28, C["gris"],
             graisse=600, retour=False)
    y = titre(sl, 120, 200, 1500, s["question"], 64, interligne=1.12) + 48
    props = s["propositions"]
    rangs = (len(props) + 1) // 2
    ht = (920 - y - 28 * (rangs - 1)) / rangs
    l_bonne = largeur("Bonne réponse", 26, 800)
    for k, p in enumerate(props):
        x = 120 + (k % 2) * (826 + 28)
        yy = y + (k // 2) * (ht + 28)
        bonne = s["_reponse"] and k == s["bonne"]
        # Tuile : bordure 4 (padding 0 44) ; lettre 72 à x + 48, texte à x + 152, « Bonne réponse » calé à x + 778.
        wt = 778 - 152 - (32 + l_bonne if bonne else 0)
        h_p = hauteur(p, wt, 40, 1.5, 600)
        with sl.bloc():
            if bonne:
                sl.rect(x, yy, 826, ht, c["teinte"], rayon=28, trait=c["couleur"], epaisseur=4)
            else:
                sl.rect(x, yy, 826, ht, C["fondCarte"], rayon=28)
            lettre = sl.rect(x + 48, yy + ht / 2 - 36, 72, 72, c["couleur"] if bonne else "#FFFFFF", rayon=18)
            sl.texte(x + 48, yy + ht / 2 - 36, 72, 72, "ABCD"[k], 34, c["surCouleur"] if bonne else C["texte"], graisse=800,
                     align=CENTRE, ancre=MILIEU, interligne=1.2, forme=lettre, retour=False)
            sl.texte(x + 152, yy + ht / 2 - h_p / 2, wt, h_p + 4, p, 40, graisse=600)
            if bonne:
                sl.texte(x + 778 - l_bonne - 20, yy + ht / 2 - 19.5, l_bonne + 20, 39, "Bonne réponse", 26, c["encre"], graisse=800,
                         align=DROITE, retour=False)


def a_retenir(sl: Slide, s):
    c = sl.deck.c
    co = c["surCouleur"]
    sl.rect(0, 0, 720, BAS, c["couleur"])
    titre_s = s.get("titre") or "À retenir"
    sl.texte(120, 96, 520, hauteur(titre_s, 520, 96, 1.0, 700) + 8, titre_s, 96, co, graisse=700, interligne=1.0, espacement=-0.02)
    if s.get("sousTitre"):  # 30 px, graisse 600, interligne 1,4 ; bas à 904
        h_s = hauteur(s["sousTitre"], 520, 30, 1.4, 600)
        sl.texte(120, 904 - h_s, 520, h_s + 4, s["sousTitre"], 30, co, graisse=600, interligne=1.4)
    # Messages : numéro (72 px) aligné sur la ligne de base du texte (38 px, interligne 1,4), écart 56, centrés sur 500.
    dn = ligne_de_base(72, 1.0) - ligne_de_base(38, 1.4)  # le texte commence 21 px sous le haut du numéro
    messages = s["messages"]
    textes = [hauteur(m, 840, 38, 1.4) for m in messages]
    hauteurs = [max(72, dn + ht) for ht in textes]
    y = 500 - (sum(hauteurs) + 56 * (len(messages) - 1)) / 2
    for k, (m, ht, h) in enumerate(zip(messages, textes, hauteurs), 1):
        with sl.bloc():
            sl.texte(840, y, 100, 72, str(k), 72, c["encre"], graisse=800, interligne=1.0, retour=False)
            sl.texte(960, y + dn, 840, ht + 4, m, 38, C["texte2"], interligne=1.4, gras_coul=C["texte"])
        y += h + 56


def ressources(sl: Slide, s):
    c = sl.deck.c
    y = titre(sl, 120, 96, 1680, s.get("titre") or "Pour continuer") + 56
    items = s["ressources"]
    # Cartes : padding 44 48 ; titre, 14, description (qui s'étire), 14, lien ; rangées auto, la place libre partagée.
    contenus = []
    for r in items:
        h_t = hauteur(r.get("titre"), 730, 36, 1.5, 700)
        h_x = hauteur(r.get("texte"), 730, 28, 1.45)
        contenus.append((h_t, h_x, 44 + h_t + 14 + h_x + (14 + 42 if r.get("lien") else 0) + 44))
    rangs = [max(x[2] for x in contenus[i:i + 2]) for i in range(0, len(items), 2)]
    libre = max(0, (920 - y - 28 * (len(rangs) - 1)) - sum(rangs)) / max(1, len(rangs))
    for k, (r, (h_t, h_x, _)) in enumerate(zip(items, contenus)):
        x = 120 + (k % 2) * (826 + 28)
        yy = y + sum(h + libre + 28 for h in rangs[: k // 2])
        hc = rangs[k // 2] + libre
        lien = r.get("lien") or ""
        skazy = r.get("skazy") or "skazy.nc" in lien
        with sl.bloc():
            sl.rect(x, yy, 826, hc, c["teinte"] if skazy else C["fondCarte"], rayon=28)
            sl.texte(x + 48, yy + 44, 730, h_t + 4, r.get("titre"), 36, graisse=700)
            sl.texte(x + 48, yy + 44 + h_t + 14, 730, h_x + 4, r.get("texte"), 28, C["texte3"], interligne=1.45, gras_coul=C["texte"])
            if lien:
                affiche = re.sub(r"^https?://(www\.)?", "", lien).rstrip("/")
                sl.texte(x + 48, yy + hc - 44 - 42, 730, 42, affiche, 28, c["encre"], graisse=700, lien=lien, souligne=True)


def evaluation(sl: Slide, s):
    c = sl.deck.c
    y = titre(sl, 120, 96, 1680, s.get("titre") or "Votre avis compte") + 56
    message = s.get("message") or "Votre attestation de fin de formation vous parvient par e-mail."
    h_m = hauteur(message, 1680, 28, 1.5)
    bas = 920 - h_m - 56  # bas des cartes
    surtitre = {"coul": c["encre"], "graisse": 800, "majuscules": True, "espacement": 0.14}
    with sl.bloc():
        sl.rect(120, y, 824, bas - y, c["teinte"], rayon=32)
        sl.pile(176, y + 56, 712, [
            ("1 · Quiz final", 24, surtitre),
            ("Mesurez votre progression", 44, {"graisse": 700, "interligne": 1.15}),
            ("Quelques questions pour vérifier ce que vous avez retenu de la formation.", 30, {"coul": C["texte2"]}),
        ], 22)
        sl.pastille(176, bas - 56 - 73, s.get("quiz") or "Sur kahoot.it, code affiché à l'écran", 30, pad_x=32, pad_y=14,
                    majuscules=False)
    with sl.bloc():
        sl.rect(976, y, 824, bas - y, C["fondCarte"], rayon=32)
        elements = [
            ("2 · Satisfaction", 24, surtitre),
            ("Trois minutes, anonyme", 44, {"graisse": 700, "interligne": 1.15}),
            ("Vos retours font évoluer la formation. Scannez le QR code.", 30, {"coul": C["texte2"]}),
        ]
        total = sum(hauteur_opts(t, 392, tl, o) for t, tl, o in elements) + 22 * 2
        sl.pile(1032, (y + bas) / 2 - total / 2, 392, elements, 22)
        qx, qy = 1464, (y + bas) / 2 - 140
        sl.rect(qx, qy, 280, 280, "#FFFFFF", rayon=24)
        sl.visuel(s.get("qrCode"), qx + 16, qy + 16, 248, 248, "QR code", 12)
    with sl.bloc():
        sl.texte(120, 920 - h_m, 1680, h_m + 4, message, 28, C["texte4"])


def questions(sl: Slide, s):
    c = sl.deck.c
    sl.photo(ch.url_photo(s.get("photo") or ch.CHARTE["photoQuestions"]), 0, 0, 960, H, position="right", saturation=1.2,
             contraste=1.05)
    sl.rect(960, 0, 960, H, c["couleur"])
    logo_entete(sl)
    t = s.get("titre") or "Avez-vous des questions ?"
    h = hauteur(t, 720, 88, 1.05, 700, espacement=-0.02)
    sl.texte(1080, (216 + 898.8) / 2 - h / 2, 720, h + 10, equilibrer(t, 720, 88, 700, -0.02), 88, c["surCouleur"], graisse=700,
             interligne=1.05, espacement=-0.02)
    contact(sl)


def battle(sl: Slide, s):
    c = sl.deck.c
    # Badge : padding 28, icône 28, écart 12, « BATTLE » ; hauteur 56.
    lb = largeur("Battle", 24, 800, espacement=0.12, majuscules=True)
    sl.rect(120, 88, 28 + 28 + 12 + lb + 28, 56, C["texte"], rayon=28)
    sl.icone("zap", 148, 102, 28, "#FFFFFF")
    sl.texte(188, 88, lb + 8, 56, "Battle", 24, "#FFFFFF", graisse=800, majuscules=True, espacement=0.12, ancre=MILIEU,
             interligne=1.2, retour=False)
    titre_s = s.get("titre") or "Battle de questions"
    accroche = s.get("accroche") or "Deux équipes, un objectif : piéger l'autre."
    w_acc = largeur(accroche, 30) + 4
    bas_titre = titre(sl, 120, 168, 1680 - 40 - w_acc, titre_s)
    sl.texte(1800 - w_acc, bas_titre - 45, w_acc, 45, accroche, 30, C["texte4"], align=DROITE, retour=False)
    # Règles en bas (bas à 936) ; cartes entre le titre (+ 44) et les règles (− 44).
    regles = s["regles"]
    n = len(regles)
    lr = (1680 - 32 * (n - 1)) / n
    h_regles = max([48] + [hauteur(r, lr - 66, 26, 1.4) for r in regles])
    y_r = 936 - h_regles
    y0, y1 = bas_titre + 44, y_r - 44
    lc = (1680 - 140 - 80) / 2
    for k, (nom, fond, coul, trait) in enumerate((
        ("Équipe A", c["couleur"], c["surCouleur"], ch.melange(c["surCouleur"], c["couleur"], 0.4)),
        ("Équipe B", C["fondCarte"], C["texte"], ch.melange(C["texte"], C["fondCarte"], 0.4)),
    )):
        x = 120 + k * (lc + 220)
        b = 3 * k  # bordure de 3 px de l'équipe B
        xi, yi, wi = x + 52 + b, y0 + 48 + b, lc - 104 - 2 * b
        with sl.bloc():
            sl.rect(x, y0, lc, y1 - y0, fond, rayon=32, trait=C["bordure"] if k else None, epaisseur=3)
            sl.texte(xi, yi, wi, 36, nom, 24, coul, graisse=800, majuscules=True, espacement=0.14)
            sl.texte(xi, yi + 60, wi, 39, "Nom de l'équipe", 26, coul, graisse=700)
            sl.pointilles(xi, yi + 214, wi, trait, 3)
            sl.texte(xi, yi + 241, wi, 39, "Score", 26, coul, graisse=700)
            sl.rect(xi, yi + 292, wi, (y1 - 48 - b) - (yi + 292), None, rayon=20, trait=trait, epaisseur=3, tirets=True)
        if k == 0:
            with sl.bloc():
                vs = sl.rond(120 + lc + 40, (y0 + y1) / 2 - 70, 140, c["teinte"])
                sl.texte(120 + lc + 40, (y0 + y1) / 2 - 70, 140, 140, "VS", 52, c["encre"], graisse=800, align=CENTRE, ancre=MILIEU,
                         interligne=1.2, forme=vs, retour=False)
    for k, r in enumerate(regles):
        x = 120 + k * (lr + 32)
        with sl.bloc():
            sl.numero(x, y_r, 48, k + 1, 24, c["teinte"], c["encre"])
            sl.texte(x + 66, y_r, lr - 66, h_regles + 4, r, 26, C["texte3"], interligne=1.4, gras_coul=C["texte"])


def accompagnement_360(sl: Slide, s):
    from support_html import POLES

    y = titre(sl, 120, 96, 1200, "Skazy") + 6
    sl.texte(120, y, 1200, 66, "Accompagnement 360°", 44, C["texte2"], graisse=500)
    y += 66 + 64
    lc = (1680 - 28 * 3) / 4
    for k, (nom, coul, texte_) in enumerate(POLES):
        x = 120 + k * (lc + 28)
        with sl.bloc():
            sl.rect(x, y, lc, 920 - y, C["fondCarte"], rayon=28)
            # « skazy » (52 px, interligne 0,9) puis le pôle (30 px) dessous ; texte 28 px plus bas.
            sl.texte(x + 40, y + 48, lc - 80, 47, "skazy", 52, C["texte2"], graisse=800, interligne=0.9, espacement=-0.02)
            sl.texte(x + 40, y + 94.8, lc - 80, 27, nom, 30, coul, graisse=800, interligne=0.9)
            sl.texte(x + 40, y + 149.8, lc - 80, 400, texte_, 26, C["texte3"])
    sl.sticker(1920 - 120 - 150, 80, 150, 8, "rocket")


def autres_formations(sl: Slide, s):
    c_cat = sl.deck.c
    y = titre(sl, 120, 88, 1680, "Nos autres formations") + 10
    sl.texte(120, y, 1680, 51, "Le catalogue complet : **{}**".format(ch.CONTACT["site"]), 34, C["texte2"],
             gras_coul=c_cat["encre"], gras_souligne=True, lien=ch.CONTACT["url"])
    texte_ = s.get("texte") or "**L'IA au quotidien** : ChatGPT, Claude, Copilot, Mistral, Gemini. Et aussi Make, Canva, Figma, Power BI, Microsoft 365."
    h_x = hauteur(texte_, 1680, 28, 1.4)
    y0, y1 = y + 51 + 48, 936 - h_x - 48
    lc, hc = (1680 - 24 * 4) / 5, (y1 - y0 - 24) / 2
    with sl.bloc():
        for k, c in enumerate(ch.CHARTE["categories"] + [None]):
            x = 120 + (k % 5) * (lc + 24)
            y = y0 + (k // 5) * (hc + 24)
            if c is None:  # tuile « Tout le catalogue » : padding 32, textes en bas, écart 10
                sl.rect(x, y, lc, hc, C["texte"], rayon=22)
                h_tc = hauteur("Tout le catalogue", lc - 64, 30, 1.2, 700)
                sl.texte(x + 32, y + hc - 32 - 36 - 10 - h_tc, lc - 64, h_tc + 4, "Tout le catalogue", 30, "#FFFFFF", graisse=700,
                         interligne=1.2)
                sl.texte(x + 32, y + hc - 32 - 36, lc - 64, 36, ch.CONTACT["site"], 24, "#FFFFFF", graisse=700,
                         lien=ch.CONTACT["url"], souligne=True)
                continue
            sl.rect(x, y, lc, hc, c["couleur"], rayon=22)
            sl.polygone([(x, y), (x + lc * 0.62, y), (x + lc * 0.62 * 0.78, y + hc), (x, y + hc)], "#FFFFFF", alpha=0.24)
            # Icône centrée dans le triangle (padding droit 12 % de la tuile, bas 64) : x + 0,25 × largeur.
            sl.icone(c["icone"], x + lc * 0.25 - 42, y + (hc - 64) / 2 - 42, 84, "#FFFFFF" if blanc(c) else C["texte"])
            lt = min(lc - 40, largeur(c["nom"], 26, 700) + 40 + 2)
            hl = 31.2 * nb_lignes(c["nom"], lt - 40, 26, 700) + 20
            sl.rect(x + 20, y + hc - 20 - hl, lt, hl, "#FFFFFF", rayon=12)
            sl.texte(x + 40, y + hc - 20 - hl + 10, lt - 40, hl - 20, c["nom"], 26, graisse=700, interligne=1.2)
    with sl.bloc():
        sl.texte(120, 936 - h_x, 1680, h_x + 4, texte_, 28, C["texte3"], interligne=1.4, gras_coul=C["texte"])


RENDUS = {
    "couverture": couverture, "faisons-connaissance": faisons_connaissance, "objectifs": objectifs, "programme": programme,
    "chapitre": chapitre, "contenu": contenu, "chiffre-cle": chiffre_cle, "frise": frise, "graphique": graphique,
    "image": image, "texte": texte, "citation": citation, "avant-apres": avant_apres, "notion-cle": notion_cle,
    "code": code, "capture": capture, "pas-a-pas": pas_a_pas, "comparatif": comparatif, "raccourcis": raccourcis,
    "bon-a-savoir": bon_a_savoir, "exercice": exercice, "exercice-etapes": exercice_etapes, "quiz": quiz,
    "a-retenir": a_retenir, "ressources": ressources, "evaluation": evaluation, "questions": questions,
    "battle": battle, "accompagnement-360": accompagnement_360, "autres-formations": autres_formations,
}

SANS_PIED = {"couverture", "image", "questions"}
# Titre, intercalaires, pauses et questions restent fixes ; la slide réponse d'un quiz est elle-même la révélation.
SANS_ANIMATION = {"couverture", "chapitre", "image", "questions"}


def animee(slide: dict) -> bool:
    return slide["type"] not in SANS_ANIMATION and not slide.get("_reponse")


def rendre_pptx(s: dict) -> bytes:
    deck = Deck(s)
    for slide in s["_slides"]:
        t = slide["type"]
        sl = Slide(deck)
        RENDUS[t](sl, slide)
        if t not in SANS_PIED:
            sl.pied(couleur_cat=(t == "chapitre"))
        sl.notes(slide.get("notes"))
        if deck.f["animations"] and animee(slide):
            sl.animer()
    tampon = io.BytesIO()
    deck.prs.save(tampon)
    # Allégé (dispositions inutilisées, vignette, images à leur taille d'affichage) : voir pptx_google_slides.py.
    from pptx_google_slides import preparer_pptx

    return preparer_pptx(tampon.getvalue())[0]
