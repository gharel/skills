import io
import json

from conftest import SKILLS

import rendre_quiz as rq

EXEMPLE = SKILLS / "skazy-quiz" / "exemples" / "quiz-exemple.json"


def quiz():
    return json.loads(EXEMPLE.read_text(encoding="utf-8"))


def test_exemple_valide():
    q, rapport = rq.normaliser_quiz(quiz())
    assert rapport.erreurs == []
    assert len(q["questions"]) == 12
    assert {x["type"] for x in q["questions"]} == set(rq.TYPES)


def test_regles_de_l_editeur():
    donnees = {
        "title": "T",
        "category": "Marketing",
        "questions": [
            {"type": "single", "text": "Q", "answers": ["a", "b"], "correct": [0, 1]},
            {"type": "multiple", "text": "Q", "answers": ["a", "b", "c"], "correct": [0]},
            {"type": "text", "text": "Q", "accepted": []},
            {"type": "slider", "text": "Q", "slider": {"min": 5, "max": 1, "step": 1, "answer": 3}},
            {"type": "order", "text": "Q", "items": ["a", "b"]},
            {"type": "single", "text": "x" * 241, "answers": ["a", "b"], "correct": [0], "time": 25},
        ],
    }
    q, rapport = rq.normaliser_quiz(donnees)
    textes = " ".join(rapport.erreurs)
    for attendu in ("une bonne réponse", "deux bonnes réponses", "réponse acceptée", "maximale", "3 ou 4 éléments", "trop long"):
        assert attendu in textes
    assert q["category"] == "webmarketing"
    assert any("25 s ramené" in a for a in rapport.avertissements)


def test_lien_de_partage_aller_retour():
    q, _ = rq.normaliser_quiz(quiz())
    lien = rq.lien_partage(q)
    assert lien.startswith("https://gharel.github.io/quiz/#/partage?q=")
    code = lien.split("q=")[1]
    assert "+" not in code and "/" not in code
    donnees = rq.lire_lien(lien)
    assert donnees["title"] == "Excel : les bases"
    assert len(donnees["questions"]) == 12


def test_export_skazy():
    q, _ = rq.normaliser_quiz(quiz())
    e = rq.export_skazy(q)
    assert e["format"] == "skazy-quiz"
    assert e["version"] == 1
    assert {"id", "title", "description", "category", "questions", "createdAt", "updatedAt"} <= set(e["quiz"])


def test_kahoot_gabarit_officiel():
    from openpyxl import load_workbook

    q, rapport = rq.normaliser_quiz(quiz())
    ws = load_workbook(io.BytesIO(rq.export_kahoot(q, rapport))).active
    assert ws["B8"].value == "Question - max 120 characters"
    assert ws["H8"].value == "Correct answer(s) - choose at least one"
    lignes = [r for r in ws.iter_rows(min_row=9, values_only=True) if r[1]]
    assert len(lignes) == 10  # réponse libre et remise en ordre trop longue écartées
    for _, question, *reponses, temps, bonnes in lignes:
        assert len(question) <= 120
        assert all(len(a) <= 75 for a in reponses if a)
        assert temps in rq.TEMPS_KAHOOT
        assert all(1 <= int(k) <= 4 for k in str(bonnes).split(","))
    vrai_faux = [r for r in lignes if r[2] == "Vrai"]
    assert vrai_faux
    assert all(r[3] == "Faux" and r[4] is None for r in vrai_faux)


def test_google_forms():
    q, rapport = rq.normaliser_quiz(quiz())
    script = rq.export_forms(q, rapport)
    assert "FormApp.create(QUIZ.titre)" in script
    assert "setIsQuiz(true)" in script
    donnees = json.loads(script.split("const QUIZ = ", 1)[1].split(";\n\nfunction", 1)[0])
    assert len(donnees["questions"]) == 12
    assert any(x["type"] == "cases" for x in donnees["questions"])
    assert any(x["type"] == "texte" for x in donnees["questions"])


def test_conversion_curseur_et_ordre():
    curseur = {"type": "slider", "text": "Q", "slider": {"min": 1, "max": 10, "step": 1, "answer": 5, "tolerance": 0}}
    c = rq.en_choix(curseur, 3)
    assert len(c["choix"]) == 4
    assert c["choix"][c["bonnes"][0]] == "5"
    ordre = {"type": "order", "text": "Q", "items": ["a", "b", "c"]}
    c = rq.en_choix(ordre, 1)
    assert c["choix"][c["bonnes"][0]] == "a → b → c"
    assert len(set(c["choix"])) == len(c["choix"])
