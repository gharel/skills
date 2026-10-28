import base64
import io
import json
import re
import zipfile

from conftest import SKILLS

import docx_google_docs as g
import rendre_plan as r

EXEMPLE = SKILLS / "skazy-plan-formation" / "exemples" / "plan-exemple.json"
EXEMPLE_IA = SKILLS / "skazy-plan-formation" / "exemples" / "plan-ia-reportings.json"


def plan(**modifs):
    donnees = json.loads(EXEMPLE.read_text(encoding="utf-8"))
    donnees.update(modifs)
    return r.normaliser_plan(donnees)


def texte_tableaux(partie):
    return " ".join(c.text for t in partie.tables for rangee in t.rows for c in rangee.cells)


def test_exemple_valide_sans_avertissement():
    p, rapport = plan()
    assert rapport.erreurs == []
    assert rapport.avertissements == []
    assert p["_cat"]["nom"] == "Bureautique"
    assert p["_modalites"]["tarif"] == "Sur devis"


def test_erreurs_et_avertissements():
    _, rapport = plan(categorie="Cuisine")
    assert any("Catégorie inconnue" in e for e in rapport.erreurs)
    _, rapport = plan(description="Cette formation présente Excel.", objectifs=["Comprendre Excel"])
    textes = " ".join(rapport.avertissements)
    assert "Cette formation" in textes
    assert "peu observable" in textes
    assert "1 objectifs" in textes


def test_debordement_signale():
    point = "Un point assez long pour occuper toute une ligne du programme"
    longs = [{"titre": "Partie {}".format(i), "points": [point] * 6} for i in range(8)]
    _, rapport = plan(programme=longs)
    assert any("dépasse" in a for a in rapport.avertissements)


def test_html_reprend_le_modele():
    p, _ = plan()
    html = r.rendre_html(p)
    assert "--c:#A977DA" in html
    assert "Excel : les bases" in html.replace(" ", " ")
    assert 'class="page p2"' in html
    assert "RIDET" in html
    assert "Créer un classeur" in html and "Prenez en main Excel" in html


def test_html_une_seule_page_sans_modalites():
    p, _ = plan(modalites=False)
    assert 'class="page p2"' not in r.rendre_html(p)


def plan_ia(**modifs):
    donnees = json.loads(EXEMPLE_IA.read_text(encoding="utf-8"))
    donnees.update(modifs)
    return r.normaliser_plan(donnees)


def test_exemple_de_reference_valide():
    p, rapport = plan_ia()
    assert rapport.erreurs == [] and rapport.avertissements == []
    m = p["_modalites"]
    assert [f["icone"] for f in m["formats"]] == ["users", "video", "blend"]
    assert m["encadrement"].startswith("Formateur qualifié") and len(m["moyens"]) == 3


def test_docx_reprend_la_mise_en_page_de_reference():
    from docx import Document

    p, _ = plan_ia()
    donnees = r.rendre_docx(p)
    doc = Document(io.BytesIO(donnees))
    # Deux sections : la page 2 a sa marge basse et son pied (boîte de contact) ; aucun en-tête.
    assert len(doc.sections) == 2
    z = zipfile.ZipFile(io.BytesIO(donnees))
    assert not any(n.startswith("word/header") for n in z.namelist())
    assert "formation@skazy.nc" in " ".join(par.text for par in doc.sections[0].footer.paragraphs) + texte_tableaux(doc.sections[0].footer)
    pied2 = texte_tableaux(doc.sections[1].footer) + " ".join(par.text for par in doc.sections[1].footer.paragraphs)
    assert "Une question sur cette formation" in pied2 and "RIDET" in pied2
    corps = " ".join(par.text for par in doc.paragraphs) + texte_tableaux(doc)
    for attendu in ("L’IA pour vos reportings", "Intelligence Artificielle", "MODALITÉS ET INFOS PRATIQUES", "EFFECTIF",
                    "Modalités et lieu de la formation", "En présentiel", "Hybride", "Accompagnement personnalisé",
                    "Méthodes et moyens pédagogiques", "Encadrement", "Évaluation des acquis", "Mode de validation",
                    "01 Cadrer les usages et les données · 30 min"):
        assert attendu in corps.replace(" ", " "), attendu
    document = z.read("word/document.xml").decode("utf-8")
    # Bandeau et sticker ancrés à la page dans le titre ; fonds arrondis derrière le texte.
    titre = re.search(r"<w:body><w:p\b.*?</w:p>", document, re.S).group(0)
    assert titre.count('relativeFrom="page"') == 4 and "L’IA pour vos reportings" in titre
    # Ancrée à la page, une image devient dans Google Docs un caractère de la taille de son passage : 1 pt.
    for passage in re.findall(r"<w:r>(?:(?!</w:r>).)*?relativeFrom=\"page\".*?</w:r>", document, re.S):
        assert '<w:sz w:val="2"/>' in passage
    ancres = re.findall(r"<wp:anchor\b[^>]*>", document)
    assert len(ancres) == 2 + 1 + 1 + 3 + 3 + 1
    assert sum('behindDoc="1"' in a for a in ancres) == len(ancres) - 1  # le sticker est devant
    # Chaque fond relatif à un paragraphe est porté par un paragraphe vide suivi de son tableau.
    for porteur in re.findall(r'<w:p\b(?:(?!</w:p>).)*relativeFrom="paragraph".*?</w:p>(<w:tbl>)?', document, re.S):
        assert porteur == "<w:tbl>"


def test_docx_options_de_la_page_2():
    p, _ = plan_ia(modalites={"formats": False, "accompagnement": False, "moyens": [], "validation": False})
    document = zipfile.ZipFile(io.BytesIO(r.rendre_docx(p))).read("word/document.xml").decode("utf-8")
    assert "En présentiel" not in document and "Accompagnement" not in document and "Mode de validation" not in document
    assert len(re.findall(r"<wp:anchor\b", document)) == 2 + 1 + 1 + 3
    html = r.rendre_html(p)
    assert "Modalités et lieu" not in html and "accomp" not in html.split("<main>")[1]
    p, _ = plan_ia(modalites=False)
    assert zipfile.ZipFile(io.BytesIO(r.rendre_docx(p))).read("word/document.xml").decode("utf-8").count("<w:sectPr") == 1


def test_docx_pret_pour_google_docs():
    p, _ = plan_ia()
    donnees = r.rendre_docx(p)
    assert len(donnees) < 40_000
    z = zipfile.ZipFile(io.BytesIO(donnees))
    noms = z.namelist()
    assert not any(n.startswith(("customXml/", "word/fonts/")) or "stylesWithEffects" in n or "thumbnail" in n for n in noms)
    xml = "".join(z.read(n).decode("utf-8") for n in noms if n.endswith(".xml"))
    assert 'w:lineRule="exact"' not in xml and "<w:caps" not in xml and "latentStyles" not in xml
    defauts = re.search(r"<w:docDefaults>.*?</w:docDefaults>", z.read("word/styles.xml").decode("utf-8"), re.S).group(0)
    assert 'w:ascii="Raleway"' in defauts and 'w:val="fr-FR"' in defauts
    assert g.plus_long_repete(base64.b64encode(donnees).decode())[0] <= 40
    # Mises en page que Google Docs lit comme Word : voir references/google-docs.md.
    assert "<w:tcMar>" not in xml  # aucune marge de cellule : espaces de paragraphe à la place
    for marges in re.findall(r"<w:tblCellMar>.*?</w:tblCellMar>", z.read("word/document.xml").decode("utf-8"), re.S):
        assert set(re.findall(r'w:w="(\d+)"', marges)) == {"0"}
    for tblpr in re.findall(r"<w:tblPr>.*?</w:tblPr>", xml, re.S):
        assert tblpr.count("<w:tblBorders>") <= 1 and tblpr.count("<w:tblLayout") <= 1
    assert 'w:hRule="atLeast"' in xml and 'w:hRule="exact"' not in xml
    numerotation = z.read("word/numbering.xml").decode("utf-8")
    assert 'w:lvlText w:val="•"' in numerotation and 'w:ascii="Arial"' in numerotation
    # Propriétés redondantes retirées : la police par défaut est Raleway.
    parties = "".join(z.read(n).decode("utf-8") for n in noms if re.match(r"word/(document|header\d|footer\d)\.xml$", n))
    assert "<w:szCs" not in parties and 'w:ascii="Raleway"' not in parties
    assert 'w:ascii="Raleway ExtraBold"' in parties


def test_largeur_des_pastilles():
    import plan_docx as pd

    # Raleway Bold 10 pt, mesures de la police : « Intelligence Artificielle » 109,8 pt, « Intermédiaire » 65,1 pt.
    assert abs(pd.chasse_gras("Intelligence Artificielle", 10) - 109.8) < 0.5
    assert abs(pd.chasse_gras("Intermédiaire", 10) - 65.1) < 0.5
    assert pd.nombre_de_lignes("L’IA pour vos reportings", 31.5, pd.UTILE - 82.5) == 1
    assert pd.nombre_de_lignes("Piloter un projet agile avec Scrum et Kanban au quotidien", 26, pd.UTILE - 82.5) == 2


def test_comparer_recopie():
    assert g.comparer_recopie("QUJDREVG\nR0g=", "QUJDREVGR0g=") == []
    ecarts = g.comparer_recopie("QUJDREVGR0g=", "QUJDRXVGR0g=")
    assert len(ecarts) == 1 and ecarts[0][0] == 5


def test_interligne_google_docs():
    import interligne_google_docs as ig
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    titre = doc.add_paragraph()
    titre.add_run("Titre du plan").font.size = Pt(20)
    titre.paragraph_format.line_spacing = Pt(30)  # interligne exact
    doc.add_paragraph().paragraph_format.line_spacing = Pt(6)  # espacement vide
    doc.add_paragraph("Texte courant")
    tampon = io.BytesIO()
    doc.save(tampon)
    paragraphes = ig.paragraphes_docx(tampon.getvalue())
    assert [(p["texte"], p["line"], p["regle"]) for p in paragraphes] == [
        ("Titre du plan", 600, "exact"), ("", 120, "exact"), ("Texte courant", None, None)]

    def para(debut, texte, taille=None, interligne=None):
        style = {"fontSize": {"magnitude": taille, "unit": "PT"}} if taille else {}
        return {"startIndex": debut, "endIndex": debut + len(texte) + 1,
                "paragraph": {"elements": [{"textRun": {"content": texte + "\n", "textStyle": style}}],
                              "paragraphStyle": {"lineSpacing": interligne} if interligne else {}}}

    lecture = {"content": {"revisionId": "r1", "tabs": [{"documentTab": {"body": {"content": [
        {"endIndex": 1, "sectionBreak": {}},
        para(1, "Titre du plan", 20, 250), para(15, ""), para(16, "Texte courant", 10, 100),
        {"startIndex": 30, "endIndex": 31, "sectionBreak": {}}, para(31, ""),
        {"startIndex": 32, "endIndex": 40, "table": {"tableRows": []}},
    ]}}}]}}
    requetes, non_apparies = ig.requetes_interligne(lecture, paragraphes)
    styles = [r["updateParagraphStyle"] for r in requetes if "updateParagraphStyle" in r]
    assert styles[0]["range"] == {"startIndex": 1, "endIndex": 15}
    assert styles[0]["paragraphStyle"]["lineSpacing"] == ig.interligne_juste(600, 20) == 127.8
    assert styles[1]["paragraphStyle"]["lineSpacing"] == ig.interligne_juste(120, 10) == 51.1
    # Paragraphe vide ajouté par Google avant le tableau d'en-tête de la section 2 : 1 pt et 100 %.
    assert {"startIndex": 31, "endIndex": 32} == requetes[-1]["updateParagraphStyle"]["range"]
    assert requetes[-2]["updateTextStyle"]["textStyle"]["fontSize"]["magnitude"] == 1
    assert len(styles) == 3 and "Texte courant" not in json.dumps(non_apparies)


def test_comparer_rendus(tmp_path):
    import comparer_rendus as cr

    fitz = __import__("pytest").importorskip("fitz")
    for nom, decalage in (("word.pdf", 0), ("google.pdf", 0.5)):
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text((48, 100 + decalage), "Excel : les bases", fontsize=24)
        page.insert_text((48, 200), "Objectifs", fontsize=12)
        doc.save(tmp_path / nom)
    export = tmp_path / "export.json"
    export.write_text(json.dumps({"content": base64.b64encode((tmp_path / "google.pdf").read_bytes()).decode()}))
    rapport = cr.comparer(str(tmp_path / "word.pdf"), str(export), str(tmp_path / "images"), tolerance=0.5)
    assert rapport["pages"] == (1, 1)
    lignes = rapport["par_page"][0]["lignes"]
    assert [x["texte"] for x in lignes] == ["Excel : les bases", "Objectifs"]
    assert abs(lignes[0]["dy"] - 0.5 * 25.4 / 72) < 0.01 and not any(x["hors_tolerance"] for x in lignes)
    assert (tmp_path / "images" / "cmp-1.png").exists() and (tmp_path / "images" / "diff-1.png").exists()


def test_preparer_un_docx_venu_de_word():
    from docx import Document
    from docx.shared import Pt
    from PIL import Image

    doc = Document()
    par = doc.add_paragraph()
    run = par.add_run("Titre")
    run.font.size = Pt(20)
    par.paragraph_format.line_spacing = Pt(30)  # interligne exact
    image = io.BytesIO()
    Image.new("RGB", (1200, 600), "#A977DA").save(image, "PNG")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(image.getvalue()), width=Pt(72))
    tampon = io.BytesIO()
    doc.save(tampon)
    # Police embarquée, comme Word l'écrit.
    source = zipfile.ZipFile(io.BytesIO(tampon.getvalue()))
    fichiers = {n: source.read(n) for n in source.namelist()}
    fichiers["word/fonts/font1.odttf"] = b"0" * 5000
    fichiers["word/_rels/fontTable.xml.rels"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font1.odttf"/></Relationships>'
    )
    fichiers["[Content_Types].xml"] = fichiers["[Content_Types].xml"].replace(
        b"<Default ", b'<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ', 1
    )
    avec_police = io.BytesIO()
    with zipfile.ZipFile(avec_police, "w") as z:
        for n, v in fichiers.items():
            z.writestr(n, v)

    resultat, messages = g.preparer_docx(avec_police.getvalue())
    z = zipfile.ZipFile(io.BytesIO(resultat))
    assert not any(n.startswith("word/fonts/") for n in z.namelist())
    assert b"odttf" not in z.read("[Content_Types].xml") and b"odttf" not in z.read("word/_rels/fontTable.xml.rels")
    document = z.read("word/document.xml").decode("utf-8")
    attendu = round(240 * 30 / (20 * g.HAUTEUR_LIGNE))
    assert 'w:line="{}" w:lineRule="auto"'.format(attendu) in document or 'w:lineRule="auto" w:line="{}"'.format(attendu) in document
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert media == ["word/media/0.png"]
    assert Image.open(io.BytesIO(z.read(media[0]))).width <= 110 + 8
    assert any("interligne" in m for m in messages)
    Document(io.BytesIO(resultat))  # le fichier reste lisible
